from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_simulation_结算.xlsx")
OUTPUT = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_simulation_价格上限.xlsx")
MAX_PRICE = 0.8


def style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for index, column in enumerate(ws.columns, start=1):
        length = max((len(str(cell.value or "")) for cell in column), default=8)
        ws.column_dimensions[get_column_letter(index)].width = min(max(length + 2, 10), 45)


wb = load_workbook(SOURCE)
source = wb["sim结算_首单"]
headers = [cell.value for cell in source[1]]
rows = [row for row in source.iter_rows(min_row=2, values_only=True) if isinstance(row[0], int)]
price_index = headers.index("买入价格(PM ask)")
result_index = headers.index("胜负")
gross_index = headers.index("毛利润")
fee_index = headers.index("手续费")
net_index = headers.index("净利润")
kept = [row for row in rows if row[price_index] <= MAX_PRICE]
excluded = [row for row in rows if row[price_index] > MAX_PRICE]

for name in ("sim结算_价格≤0.8", "sim排除_价格>0.8"):
    if name in wb.sheetnames:
        del wb[name]

kept_ws = wb.create_sheet("sim结算_价格≤0.8")
kept_ws.append(headers)
for number, row in enumerate(kept, start=1):
    values = list(row)
    values[0] = number
    kept_ws.append(values)
wins = sum(row[result_index] == "赢" for row in kept)
losses = sum(row[result_index] == "输" for row in kept)
gross = sum(row[gross_index] for row in kept)
fee = sum(row[fee_index] for row in kept)
net = sum(row[net_index] for row in kept)
kept_ws.append([])
kept_ws.append(["汇总", None, None, None, None, None, None, len(kept), None, None, f"{wins}胜{losses}负", gross, fee, net])
style(kept_ws)

excluded_ws = wb.create_sheet("sim排除_价格>0.8")
excluded_ws.append(headers + ["排除原因"])
for number, row in enumerate(excluded, start=1):
    values = list(row)
    values[0] = number
    excluded_ws.append(values + [f"买入价格 {row[price_index]:.3f} > {MAX_PRICE:.1f}"])
style(excluded_ws)

meta = wb["sim口径_1008-09"]
meta.append(("价格上限", "排除买入价格 > 0.8；价格恰好等于0.8仍保留。"))
meta.append(("价格上限后结果", f"{len(kept)}场，{wins}胜{losses}负；毛利润={gross:.4f}，手续费={fee:.4f}，净利润={net:.4f}"))
wb.save(OUTPUT)
print(f"kept={len(kept)} excluded={len(excluded)} wins={wins} losses={losses}")
print(f"gross={gross:.6f} fee={fee:.6f} net={net:.6f}")
for row in excluded:
    print("excluded", row[1], row[price_index], row[result_index], row[net_index])
