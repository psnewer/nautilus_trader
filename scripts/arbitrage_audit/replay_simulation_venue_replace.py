"""回放 simulation 日志中的 VenueReplace 动态反转策略。

口径：convert=True、attitude=True、deviate_convert=True，不启用 tier_convert/tier_ignore；
最终 PM ask > 0.8 的机会跳过，同一 pair 取首个通过价格上限的机会。
"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path("/Users/miller/Desktop/PM/连续触发策略_全日志明细_20260929-20261009_simulation.xlsx")
OUTPUT = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_simulation_动态反转.xlsx")
MAX_PRICE = 0.8
FEE_RATE = 0.05

WINNERS = {
    "Tennis|Carlo Alberto Caniato|Francesco Passaro": "no",
    "Tennis|Daniel Rincon|Pedro Martinez": "no",
    "Tennis|Gabriele Piraino|Enrico Dalla Valle": "no",
    "Tennis|Gustavo Heide|Eduardo Ribeiro": "no",
    "Tennis|Joaquin Aguilar|Lautaro Midon": "no",
    "Tennis|Juan Carlos Prado|Pedro Boscardin Dias": "yes",
    "Tennis|Max Hans Rehberg|Hynek Barton": "no",
    "Tennis|Miguel Tobon|Tomas Barrios": "yes",
    "Tennis|Pablo Llamas Ruiz|Petr Nesterov": "yes",
    "Tennis|Vilius Gaubas|Duje Ajdukovic": "yes",
}

# 该行的 TrendGate 日志早于同次 evaluate-scheduled 日志落盘，旧解析器未关联到 venue。
# 日志只淘汰 PM YES；结合该 candidate 的正 rebate，保留腿是 OE NO。
SOURCE_VENUE_CORRECTIONS = {72981: "ORBITEXCH"}


def opposite(outcome: str) -> str:
    return "no" if outcome == "yes" else "yes"


def style(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for index, column in enumerate(ws.columns, start=1):
        width = max((len(str(cell.value or "")) for cell in column), default=8) + 2
        ws.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 44)


def row_value(row, columns, name):
    return row[columns[name]]


def replay(row, columns):
    source = row_value(row, columns, "原腿venue(推断)")
    line_no = row_value(row, columns, "日志行")
    source = SOURCE_VENUE_CORRECTIONS.get(line_no, source)
    original = row_value(row, columns, "连续触发原腿方向")
    pm_bid = row_value(row, columns, f"PM_{original.upper()}_bid")
    pm_ask = row_value(row, columns, f"PM_{original.upper()}_ask")
    target = opposite(original)
    opposite_ask = row_value(row, columns, f"PM_{target.upper()}_ask")
    start = row_value(row, columns, f"start_price_{original.upper()}")
    commission = row_value(row, columns, "PM_ask_sum")
    commission_ok = commission is not None and 0.98 <= commission <= 1.02

    flipped = False
    if source == "POLYMARKET":
        flipped = True
        action = "convert反转"
    elif source == "ORBITEXCH" and commission_ok and start is not None and pm_bid is not None:
        if pm_bid <= start:
            flipped = True
            action = "attitude反转"
        elif pm_bid >= 1.2 * start:
            flipped = True
            action = "deviate_convert反转"
        else:
            action = "默认同方向替换"
    elif source == "ORBITEXCH":
        if start is None:
            action = "默认同方向替换（缺start_price）"
        elif not commission_ok:
            action = "默认同方向替换（commission不合格）"
        elif pm_bid is None:
            action = "默认同方向替换（缺bid）"
        else:
            action = "默认同方向替换"
    else:
        return {
            "source": source,
            "original": original,
            "action": "无法确认原腿venue",
            "final": None,
            "price": None,
            "price_ok": False,
            "bid": pm_bid,
            "start": start,
            "commission": commission,
            "commission_ok": commission_ok,
        }

    final = target if flipped else original
    price = opposite_ask if flipped else pm_ask
    return {
        "source": source,
        "original": original,
        "action": action,
        "final": final,
        "price": price,
        "price_ok": price is not None and price <= MAX_PRICE,
        "bid": pm_bid,
        "start": start,
        "commission": commission,
        "commission_ok": commission_ok,
    }


wb = load_workbook(SOURCE)
raw = wb["sim原始_1008-09"]
columns = {cell.value: index for index, cell in enumerate(raw[1])}

for sheet_name in ("sim动态反转_全触发", "sim动态反转_首单", "sim动态反转_汇总"):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

all_headers = [
    "序号", "日志行", "时间(北京时间)", "pair_id", "比分", "原腿venue", "原腿方向",
    "原方向PM_bid", "原方向start_price", "PM_ask_sum", "commission合格", "动作",
    "回放方向", "回放选手", "回放PM_ask", "price≤0.8", "是否该场首个合格机会",
]
all_ws = wb.create_sheet("sim动态反转_全触发")
all_ws.append(all_headers)
selected = {}
all_results = []

for number, row in enumerate(raw.iter_rows(min_row=2, values_only=True), start=1):
    pair = row_value(row, columns, "pair_id")
    result = replay(row, columns)
    _, yes_name, no_name = pair.split("|", 2)
    final_name = yes_name if result["final"] == "yes" else no_name if result["final"] == "no" else None
    is_first = result["price_ok"] and pair not in selected
    item = {
        "row": row,
        "pair": pair,
        "result": result,
        "final_name": final_name,
        "is_first": is_first,
    }
    all_results.append(item)
    if is_first:
        selected[pair] = item
    all_ws.append([
        number,
        row_value(row, columns, "日志行"),
        row_value(row, columns, "simulation时间(北京时间)"),
        pair,
        row_value(row, columns, "触发比分"),
        result["source"],
        result["original"],
        result["bid"],
        result["start"],
        result["commission"],
        "是" if result["commission_ok"] else "否",
        result["action"],
        result["final"],
        final_name,
        result["price"],
        "是" if result["price_ok"] else "否",
        "是" if is_first else "否",
    ])
style(all_ws)

first_headers = [
    "序号", "时间(北京时间)", "pair_id", "比分", "原腿venue", "原腿方向", "动作",
    "买入方向", "买入选手", "价格", "数量", "胜者方向", "胜者", "胜负", "毛利润",
    "手续费", "净利润",
]
first_ws = wb.create_sheet("sim动态反转_首单")
first_ws.append(first_headers)
gross_total = fee_total = 0.0
wins = losses = 0
action_counts = {}
for number, (pair, item) in enumerate(selected.items(), start=1):
    row = item["row"]
    result = item["result"]
    direction = result["final"]
    price = result["price"]
    qty = float(row_value(row, columns, "数量"))
    winner = WINNERS[pair]
    _, yes_name, no_name = pair.split("|", 2)
    winner_name = yes_name if winner == "yes" else no_name
    won = direction == winner
    gross = qty * ((1 - price) if won else -price)
    fee = qty * price * (1 - price) * FEE_RATE
    net = gross - fee
    gross_total += gross
    fee_total += fee
    wins += int(won)
    losses += int(not won)
    action_counts[result["action"]] = action_counts.get(result["action"], 0) + 1
    first_ws.append([
        number,
        row_value(row, columns, "simulation时间(北京时间)"),
        pair,
        row_value(row, columns, "触发比分"),
        result["source"],
        result["original"],
        result["action"],
        direction,
        item["final_name"],
        price,
        qty,
        winner,
        winner_name,
        "赢" if won else "输",
        gross,
        fee,
        net,
    ])
first_ws.append([])
first_ws.append([
    "汇总", None, None, None, None, None, None, None, None, None, len(selected), None, None,
    f"{wins}胜{losses}负", gross_total, fee_total, gross_total - fee_total,
])
style(first_ws)

summary_ws = wb.create_sheet("sim动态反转_汇总")
summary_ws.append(["项目", "值"])
summary_ws.append(["策略参数", "convert=true, attitude=true, deviate_convert=true；不启用tier_convert/tier_ignore"])
summary_ws.append(["价格门控", "最终买入PM ask > 0.8 跳过；同一pair继续寻找后续首个合格机会"])
summary_ws.append(["原始simulation触发", len(all_results)])
summary_ws.append(["最终计入场次", len(selected)])
summary_ws.append(["胜负", f"{wins}胜{losses}负"])
summary_ws.append(["毛利润", gross_total])
summary_ws.append(["手续费", fee_total])
summary_ws.append(["净利润", gross_total - fee_total])
baseline = {}
baseline_gross = baseline_fee = 0.0
baseline_wins = 0
for item in all_results:
    pair = item["pair"]
    row = item["row"]
    price = row_value(row, columns, "下单价")
    if pair in baseline or price is None or price > MAX_PRICE:
        continue
    direction = row_value(row, columns, "连续触发原腿方向")
    qty = float(row_value(row, columns, "数量"))
    won = direction == WINNERS[pair]
    gross = qty * ((1 - price) if won else -price)
    fee = qty * price * (1 - price) * FEE_RATE
    baseline[pair] = True
    baseline_gross += gross
    baseline_fee += fee
    baseline_wins += int(won)
summary_ws.append(["同口径无动态反转场次", len(baseline)])
summary_ws.append(["同口径无动态反转胜负", f"{baseline_wins}胜{len(baseline)-baseline_wins}负"])
summary_ws.append(["同口径无动态反转毛利润", baseline_gross])
summary_ws.append(["同口径无动态反转净利润", baseline_gross - baseline_fee])
for action, count in sorted(action_counts.items()):
    summary_ws.append([f"动作：{action}", count])
all_action_counts = {}
for item in all_results:
    action = item["result"]["action"]
    all_action_counts[action] = all_action_counts.get(action, 0) + 1
for action, count in sorted(all_action_counts.items()):
    summary_ws.append([f"全270条动作：{action}", count])
summary_ws.append([
    "原腿venue未确认说明",
    "5条后续重复触发无法唯一还原原腿venue，但均发生在对应pair首个合格机会之后，不影响逐场首单和利润。",
])
style(summary_ws)

wb.save(OUTPUT)
print(f"output={OUTPUT}")
print(f"selected={len(selected)} wins={wins} losses={losses}")
print(f"gross={gross_total:.6f} fee={fee_total:.6f} net={gross_total-fee_total:.6f}")
print(f"actions={action_counts}")
for pair, item in selected.items():
    result = item["result"]
    winner = WINNERS[pair]
    print(pair, result["source"], result["original"], result["action"], result["final"], result["price"], "WIN" if result["final"] == winner else "LOSE")
