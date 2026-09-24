"""基于成交审计工作簿生成原生 PM 腿利润测算页。"""

from __future__ import annotations

import json
import re
import subprocess
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT = Path("/Users/miller/Desktop/成交订单比分_start_price_venue_stale_window_20260921-22.xlsx")
LOGS = (
    "/private/tmp/arb_remote_log_20260922.out",
    "/private/tmp/arb_remote_nohup_20260922.out",
)
ANSI = re.compile(r"\x1b\[[0-9;]*m")
ISO_TS = re.compile(r"(2026-\d\d-\d\dT\d\d:\d\d:\d\d\.\d+Z)")


def parse_ts(line: str) -> datetime | None:
    match = ISO_TS.search(line)
    return None if match is None else datetime.fromisoformat(match.group(1).replace("Z", "+00:00"))


orders = {}
fills = {}
for path in LOGS:
    with open(path, "r", errors="replace") as handle:
        for raw in handle:
            if "client_order_id=ARB-" not in raw:
                continue
            line = ANSI.sub("", raw.rstrip("\n"))
            ts = parse_ts(line)
            if ts is None:
                continue
            if "OrderInitialized(" in line:
                order_id = re.search(r"client_order_id=(ARB-[^,]+)", line)
                pair = re.search(r"arb:pair_id=([^']+)", line)
                instrument = re.search(r"instrument_id=([^,]+)", line)
                role = re.search(r"arb:leg_key=polymarket:([^:']+):", line)
                if order_id and pair and instrument and role:
                    orders[order_id.group(1)] = {
                        "pair": pair.group(1),
                        "instrument": instrument.group(1),
                        "actual_role": role.group(1),
                    }
            elif "OrderFilled(" in line:
                order_id = re.search(r"client_order_id=(ARB-[^,]+)", line)
                qty = re.search(r"last_qty=([0-9.]+)", line)
                price = re.search(r"last_px=([0-9.]+)", line)
                commission = re.search(r"commission=([0-9.]+)", line)
                if order_id and qty and price and commission:
                    fills[order_id.group(1)] = {
                        "ts": ts,
                        "qty": float(qty.group(1)),
                        "price": float(price.group(1)),
                        "commission": float(commission.group(1)),
                    }

filled_orders = []
fee_rate_by_condition = defaultdict(list)
for order_id, fill in fills.items():
    order = orders.get(order_id)
    if order is None:
        continue
    instrument = order["instrument"]
    condition_id, token_venue = instrument.split("-", 1)
    token_id = token_venue.rsplit(".", 1)[0]
    denominator = fill["qty"] * fill["price"] * (1.0 - fill["price"])
    fee_rate = fill["commission"] / denominator if denominator else 0.0
    fee_rate_by_condition[condition_id].append(fee_rate)
    filled_orders.append({
        **order,
        **fill,
        "order_id": order_id,
        "condition_id": condition_id,
        "actual_token_id": token_id,
    })

orders_by_pair = defaultdict(list)
for item in sorted(filled_orders, key=lambda value: value["ts"]):
    orders_by_pair[item["pair"]].append(item)


def fetch_market(condition_id: str) -> dict:
    url = f"https://clob.polymarket.com/markets/{condition_id}"
    completed = subprocess.run(
        ["curl", "-sS", "--retry", "5", "--retry-all-errors", "--max-time", "30", url],
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(completed.stdout)


conditions = sorted({item["condition_id"] for item in filled_orders})
with ThreadPoolExecutor(max_workers=4) as pool:
    markets = dict(zip(conditions, pool.map(fetch_market, conditions), strict=True))

wb = load_workbook(REPORT)
source = wb["成交明细"]
headers = {cell.value: cell.column for cell in source[1]}
pair_offsets = defaultdict(int)
results = []
for source_row in range(2, source.max_row + 1):
    sport = str(source.cell(source_row, headers["赛事"]).value)
    player_1 = str(source.cell(source_row, headers["选手1"]).value)
    player_2 = str(source.cell(source_row, headers["选手2"]).value)
    pair_id = f"{sport}|{player_1}|{player_2}"
    offset = pair_offsets[pair_id]
    pair_offsets[pair_id] += 1
    order = orders_by_pair[pair_id][offset]

    actual_role = str(source.cell(source_row, headers["成交方向"]).value)
    native_venue = str(source.cell(source_row, headers["原生腿venue"]).value)
    pm_role = actual_role if native_venue == "POLYMARKET" else ("no" if actual_role == "yes" else "yes")
    price = float(source.cell(source_row, headers[f"触发时_PM_{pm_role.upper()}"]).value)
    quantity = 5.05

    market = markets[order["condition_id"]]
    tokens = market["tokens"]
    actual_outcome = next(
        token["outcome"] for token in tokens if str(token["token_id"]) == order["actual_token_id"]
    )
    other_outcome = next(token["outcome"] for token in tokens if token["outcome"] != actual_outcome)
    outcome_by_role = {
        order["actual_role"]: actual_outcome,
        "no" if order["actual_role"] == "yes" else "yes": other_outcome,
    }
    selected_outcome = outcome_by_role[pm_role]
    winner = next((token["outcome"] for token in tokens if token.get("winner")), None)
    won = winner == selected_outcome
    fee_rate = sum(fee_rate_by_condition[order["condition_id"]]) / len(
        fee_rate_by_condition[order["condition_id"]],
    )
    commission = round(quantity * price * fee_rate * (1.0 - price), 5)
    principal = quantity * price
    payout = quantity if won else 0.0
    profit = payout - principal - commission
    results.append([
        source.cell(source_row, headers["成交时间(北京时间)"]).value,
        sport,
        player_1,
        player_2,
        pm_role,
        selected_outcome,
        winner,
        "赢" if won else "输",
        quantity,
        price,
        fee_rate,
        commission,
        principal,
        payout,
        profit,
        order["condition_id"],
        native_venue,
        actual_role,
    ])

for sheet_name in ("原生PM腿利润测算", "利润汇总"):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

summary = wb.create_sheet("利润汇总", 0)
total_principal = sum(row[12] for row in results)
total_commission = sum(row[11] for row in results)
total_payout = sum(row[13] for row in results)
total_profit = sum(row[14] for row in results)
wins = sum(row[7] == "赢" for row in results)
summary_rows = [
    ("假设", "每次成交机会只买 one_side_rebate 原始跨 venue candidate 中的原生 PM 腿"),
    ("订单数", len(results)),
    ("每单数量", 5.05),
    ("赢/输", f"{wins}/{len(results) - wins}"),
    ("本金合计", total_principal),
    ("手续费合计", total_commission),
    ("结算收入", total_payout),
    ("净利润", total_profit),
    ("本金收益率", total_profit / total_principal if total_principal else None),
    ("成交假设", "全部按触发快照 PM ask 足额成交，手续费率由同 condition 实际成交反推"),
]
for row in summary_rows:
    summary.append(row)
summary["A1"].font = summary["B1"].font = Font(bold=True, color="FFFFFF")
summary["A1"].fill = summary["B1"].fill = PatternFill("solid", fgColor="1F4E78")
summary.column_dimensions["A"].width = 18
summary.column_dimensions["B"].width = 85
for row_index in range(2, summary.max_row + 1):
    summary.cell(row_index, 2).alignment = Alignment(wrap_text=True)
for row_index in range(5, 9):
    summary.cell(row_index, 2).number_format = "0.0000"
summary.cell(9, 2).number_format = "0.00%"

detail = wb.create_sheet("原生PM腿利润测算", 1)
detail_headers = [
    "触发时间(北京时间)", "赛事", "选手1", "选手2", "PM角色", "假设买入选手", "最终胜者",
    "结果", "数量", "下单价", "手续费率", "手续费", "本金", "结算收入", "净利润",
    "condition_id", "实际成交腿原生venue", "实际成交方向",
]
detail.append(detail_headers)
for row in results:
    detail.append(row)
for cell in detail[1]:
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for row_index in range(2, detail.max_row + 1):
    detail.cell(row_index, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    detail.cell(row_index, 8).fill = PatternFill(
        "solid",
        fgColor="D9EAD3" if detail.cell(row_index, 8).value == "赢" else "F4CCCC",
    )
    for column in range(9, 16):
        detail.cell(row_index, column).number_format = "0.000000"
detail.freeze_panes = "A2"
detail.auto_filter.ref = detail.dimensions
for column, width in enumerate((21, 10, 22, 22, 10, 22, 22, 9, 10, 10, 12, 12, 12, 12, 12, 70, 20, 14), 1):
    detail.column_dimensions[get_column_letter(column)].width = width

wb.save(REPORT)
print(json.dumps({
    "orders": len(results),
    "wins": wins,
    "losses": len(results) - wins,
    "principal": round(total_principal, 6),
    "commission": round(total_commission, 6),
    "payout": round(total_payout, 6),
    "profit": round(total_profit, 6),
    "roi": round(total_profit / total_principal, 8),
    "fee_rates": sorted({round(rate, 8) for rates in fee_rate_by_condition.values() for rate in rates}),
}, ensure_ascii=False))
