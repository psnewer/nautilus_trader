"""用官方 Polymarket market 元数据重建历史订单的 tier 分类、策略方向与利润表。

脚本不请求网络：market 元数据与日志审计结果由参数传入。无 spread 价格取
目标方向 bid，缺 bid 时取 ``1 - 对向 ask``；spread=0.05 在 limit 价之后应用。
"""

from __future__ import annotations

import argparse
import json
import math
import re
import unicodedata
from collections import defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


LOWER_TIER_TOURNAMENTS = {
    "Adana",
    "Bari",
    "Columbus",
    "Curitiba",
    "Jingshan",
    "Mouilleron-Le-Captif",
    "Plovdiv 4",
    "Porto 2",
    "San Diego 2",
    "St. Tropez",
    "Tolentino",
}
LOWER_TIER_MARKERS = ("WTA", "CHALLENGER", "UTR", "ITF")
TIER_IGNORE = 0.3
SPREAD = 0.05


def opposite(side: str) -> str:
    return "no" if side == "yes" else "yes"


def number(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def classify_market(market: dict) -> tuple[bool, str, str]:
    question = str(market.get("question") or "")
    slug = str(market.get("market_slug") or market.get("slug") or "")
    tournament = question.split(":", 1)[0].strip()
    haystack = f"{question} {slug}".upper()
    if slug.lower().startswith("wta-"):
        return True, tournament, "official_market_slug:wta"
    marker = next((item for item in LOWER_TIER_MARKERS if item in haystack), None)
    if marker:
        return True, tournament, f"official_market_marker:{marker}"
    if tournament in LOWER_TIER_TOURNAMENTS:
        return True, tournament, "official_market_tournament:challenger"
    return False, tournament, "official_market_tournament:main_tour"


def normalize_player(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = re.sub(r"\bluis guto miguel\b", "luis miguel", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def official_winner(market: dict, match: str) -> str | None:
    """Tennis 2-way market 的第一/第二 token 分别映射 pair yes/no。"""
    players = match.split(" vs ", 1)
    tokens = market.get("tokens") or []
    if len(players) != 2 or len(tokens) != 2:
        raise ValueError(f"非两方 tennis market: match={match!r} tokens={tokens!r}")
    token_players = [str(token.get("outcome") or "") for token in tokens]
    if [normalize_player(value) for value in players] != [
        normalize_player(value) for value in token_players
    ]:
        raise ValueError(
            f"market token 顺序与 pair 不一致: match={match!r} tokens={token_players!r}",
        )
    winners = [idx for idx, token in enumerate(tokens) if token.get("winner") is True]
    if winners == [0]:
        return "yes"
    if winners == [1]:
        return "no"
    if not winners:
        return None
    raise ValueError(f"market 结算方向不唯一: match={match!r} winners={winners!r}")


def base_action(record: dict) -> tuple[str, str]:
    source = record["source"]
    if record["native_venue"] == "POLYMARKET":
        return opposite(source), "venue_replace convert反买"
    commission = record["ask_sum"]
    start = record["start"]
    bid = record["bid"]
    if commission is not None and 0.98 <= commission <= 1.02 and start is not None and bid is not None:
        if bid <= start:
            return opposite(source), "attitude反买"
        if bid >= 1.2 * start:
            return opposite(source), "deviate_convert反买"
        return source, "买原方向"

    # 38 笔早期订单已无原始触发帧，从旧表中已审计的基础分支恢复，
    # 只重算 tier pre/post；不用 start_yes+start_no 冒充当前 commission。
    old_actions = f"{record['old_pre_action']} | {record['old_post_action']}"
    if "venue_replace convert" in old_actions:
        return opposite(source), "venue_replace convert反买"
    if "attitude" in old_actions:
        return opposite(source), "attitude反买"
    if "deviate_convert" in old_actions:
        return opposite(source), "deviate_convert反买"
    return source, "买原方向"


def replay(record: dict, tier_mode: str) -> tuple[str, str]:
    source = record["source"]
    start = record["start"]
    tier_eligible = record["lower"] and start is not None and start >= TIER_IGNORE
    if tier_mode == "pre" and tier_eligible:
        return opposite(source), "tier_convert pre反买"

    side, action = base_action(record)
    ignored = None
    if record["lower"] and not tier_eligible:
        ignored = "缺start_price" if start is None else f"start_price={start:.2f} < {TIER_IGNORE:.2f}"
    if tier_mode == "post" and tier_eligible:
        final = opposite(side)
        if side == source:
            return final, "tier_convert post反买"
        return final, f"{action} + tier_convert post再反转（最终原方向）"
    if ignored:
        return side, f"{action}（tier_convert {tier_mode}未命中：{ignored}）"
    return side, action


def target_bid(record: dict, target: str) -> tuple[float, str]:
    if target == record["logged_side"] and record["logged_bid"] is not None:
        return record["logged_bid"], "logged_bid"
    if target == record["source"] and record["source_bid"] is not None:
        return record["source_bid"], "source_bid"
    if target != record["source"] and record["source_ask"] is not None:
        return 1.0 - record["source_ask"], "1-source_ask"
    if target != record["logged_side"] and record["logged_ask"] is not None:
        return 1.0 - record["logged_ask"], "1-logged_ask"
    raise ValueError(f"{record['id']} 缺少 {target} 方向 bid 及对向 ask")


def pnl(side: str, winner: str, price: float, quantity: float) -> float:
    return quantity * ((1.0 if side == winner else 0.0) - price)


def canonical_orders(records: list[dict]) -> tuple[list[dict], dict[str, str]]:
    """总盘只保留非风控订单，并按比赛去重。

    同场优先级：已成交 > 部分成交 > 其它状态；同一优先级保留时间顺序中首笔。
    """
    included = []
    exclusion = {}
    grouped: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        if "风控拒绝" in str(record["status"]):
            exclusion[record["id"]] = "排除：风控拒绝"
            continue
        grouped[record["match"]].append(record)
    priority = {"已成交": 0, "部分成交": 1}
    for rows in grouped.values():
        selected = min(rows, key=lambda row: priority.get(str(row["status"]), 2))
        included.append(selected)
        for row in rows:
            if row is not selected:
                exclusion[row["id"]] = f"排除：同场重复，保留{selected['id']}"
    included_ids = {record["id"] for record in included}
    for record in records:
        if record["id"] in included_ids:
            exclusion[record["id"]] = "纳入"
    return included, exclusion


def replace_rows(ws, headers: list[str], rows: list[list[object]]) -> None:
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("market_metadata", type=Path)
    parser.add_argument("combined_audits", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    metadata = json.loads(args.market_metadata.read_text(encoding="utf-8"))
    audits = json.loads(args.combined_audits.read_text(encoding="utf-8"))["rows"]
    wb = load_workbook(args.workbook)
    raw = wb["nohup成交_统一规则"]
    raw_headers = [cell.value for cell in raw[1]]
    raw_by_id = {}
    for values in raw.iter_rows(min_row=2, values_only=True):
        row = dict(zip(raw_headers, values, strict=True))
        if row.get("client_order_id"):
            raw_by_id[str(row["client_order_id"])] = row

    old_direction = wb["pre_post方向重算"]
    direction_headers = [cell.value for cell in old_direction[1]]
    direction_by_id = {
        str(values[0]): dict(zip(direction_headers, values, strict=True))
        for values in old_direction.iter_rows(min_row=2, values_only=True)
        if values[0]
    }

    records = []
    for order_id, row in raw_by_id.items():
        old = direction_by_id.get(order_id, {})
        order_meta = metadata["orders"][order_id]
        market = metadata["markets"][order_meta["condition_id"]]
        lower, competition, lower_source = classify_market(market)
        audit = audits.get(order_id, {})
        yes_ask = number(audit.get("触发时_PM_YES"))
        no_ask = number(audit.get("触发时_PM_NO"))
        ask_sum = yes_ask + no_ask if yes_ask is not None and no_ask is not None else None
        native = str(audit.get("原生腿venue") or old.get("原生腿venue") or "EXTERNAL").upper()
        record = {
            "id": order_id,
            "match": row["比赛"],
            "status": row["订单状态"],
            "source": str(row["规则输入方向(venue_replace前)"]).lower(),
            "native_venue": native,
            "lower": lower,
            "competition": competition,
            "lower_source": lower_source,
            "start": number(row["start_price"]),
            "bid": number(row["current_PM_bid(规则判断价)"]),
            "ask_sum": ask_sum,
            "source_bid": number(row["current_PM_bid(规则判断价)"]),
            "source_ask": number(row["current_PM_ask(参考)"]),
            "logged_side": str(row["历史OrderInitialized最终方向"] or "").lower(),
            "logged_bid": number(row["历史最终方向PM_bid(下单时)"]),
            "logged_ask": number(row["历史最终方向PM_ask(下单时)"]),
            "historical_side": str(row["历史OrderInitialized最终方向"] or "").lower(),
            "winner": official_winner(market, str(row["比赛"])),
            "quantity": number(row["原始下单数量"]) or 10.0,
            "old_pre_action": str(old.get("pre口径动作") or row.get("pre口径策略动作") or ""),
            "old_post_action": str(old.get("post口径动作") or row.get("post口径策略动作") or ""),
        }
        if record["source"] not in {"yes", "no"}:
            raise ValueError(f"{order_id} 方向无效")
        for mode in ("pre", "post"):
            side, action = replay(record, mode)
            price, price_source = target_bid(record, side)
            record[mode] = {
                "side": side,
                "action": action,
                "price": price,
                "price_source": price_source,
                "won": None if record["winner"] is None else side == record["winner"],
                "pnl": None if record["winner"] is None else pnl(
                    side, record["winner"], price, record["quantity"]
                ),
                "spread_price": max(0.01, price - SPREAD),
            }
        records.append(record)

    canonical, inclusion = canonical_orders(records)

    # 原始明细表中的策略列也同步，历史实际成交列不动。
    raw_index = {name: idx + 1 for idx, name in enumerate(raw_headers)}
    for name in ("官方最终胜方", "历史实际方向官方输赢"):
        if name not in raw_index:
            column = raw.max_column + 1
            raw.cell(1, column, name)
            raw_index[name] = column
    row_index = {str(raw.cell(r, raw_index["client_order_id"]).value): r for r in range(2, raw.max_row + 1)}
    for record in records:
        r = row_index[record["id"]]
        raw.cell(r, raw_index["pre口径策略动作"], record["pre"]["action"])
        raw.cell(r, raw_index["pre口径买入方向"], record["pre"]["side"])
        raw.cell(r, raw_index["post口径策略动作"], record["post"]["action"])
        raw.cell(r, raw_index["post口径买入方向"], record["post"]["side"])
        raw.cell(r, raw_index["官方最终胜方"], record["winner"])
        actual_side = str(record["historical_side"] or "").lower()
        raw.cell(
            r,
            raw_index["历史实际方向官方输赢"],
            "未出" if record["winner"] is None else (
                "赢" if actual_side == record["winner"] else "输" if actual_side in {"yes", "no"} else None
            ),
        )

    direction_rows = []
    profit_rows = []
    spread_rows = []
    for record in records:
        pre, post = record["pre"], record["post"]
        direction_rows.append([
            record["id"], record["match"], record["status"], record["source"],
            record["native_venue"], "是" if record["lower"] else "否", record["start"],
            record["source_bid"], pre["action"], pre["side"], post["action"], post["side"],
            "是" if pre["side"] != post["side"] else "否", record["winner"],
            record["competition"], record["lower_source"],
        ])
        profit_rows.append([
            record["id"], record["match"], record["status"], record["winner"], record["source"],
            pre["action"], pre["side"], pre["price"], pre["price_source"],
            "未出" if pre["won"] is None else "赢" if pre["won"] else "输", pre["pnl"], post["action"], post["side"],
            post["price"], post["price_source"], "未出" if post["won"] is None else "赢" if post["won"] else "输", post["pnl"],
            record["quantity"], "是" if inclusion[record["id"]] == "纳入" else "否", inclusion[record["id"]],
        ])
        spread_values = []
        for mode in ("pre", "post"):
            result = record[mode]
            unfilled = "撤单" in str(record["status"]) and result["side"] == record["historical_side"]
            spread_pnl = 0.0 if unfilled else None if record["winner"] is None else pnl(
                result["side"], record["winner"], result["spread_price"], record["quantity"]
            )
            spread_values.extend([
                result["action"], result["side"], result["price"], result["spread_price"],
                result["price_source"], "未成交" if unfilled else "按成交测算",
                "-" if unfilled else "未出" if result["won"] is None else ("赢" if result["won"] else "输"), spread_pnl,
            ])
        spread_rows.append([
            record["id"], record["match"], record["status"], record["historical_side"],
            record["winner"], *spread_values[:8], *spread_values[8:], record["quantity"],
            "是" if inclusion[record["id"]] == "纳入" else "否", inclusion[record["id"]],
        ])

    replace_rows(wb["pre_post方向重算"], [
        "client_order_id", "比赛", "订单状态", "原始腿方向", "原生腿venue", "低级别赛事",
        "start_price", "原方向PM bid", "pre口径动作", "pre口径方向", "post口径动作",
        "post口径方向", "pre/post方向不同", "最终胜方", "官方competition", "赛事分类来源",
    ], direction_rows)
    replace_rows(wb["pre_post利润重算"], [
        "client_order_id", "比赛", "订单状态", "最终胜方", "原始腿方向", "pre口径动作",
        "pre买入方向", "pre价格", "pre价格来源", "pre输赢", "pre利润", "post口径动作",
        "post买入方向", "post价格", "post价格来源", "post输赢", "post利润", "数量",
        "总盘纳入", "总盘纳入说明",
    ], profit_rows)
    spread_headers = [
        "client_order_id", "比赛", "历史订单状态", "历史下单方向", "最终胜方", "pre动作", "pre方向",
        "pre_limit基础价", "pre_spread后价格", "pre价格来源", "pre测算成交状态", "pre输赢", "pre利润",
        "post动作", "post方向", "post_limit基础价", "post_spread后价格", "post价格来源",
        "post测算成交状态", "post输赢", "post利润", "数量", "总盘纳入", "总盘纳入说明",
    ]
    replace_rows(wb["spread005利润重算"], spread_headers, spread_rows)

    summary_rows = []
    for mode in ("pre", "post"):
        settled = [record for record in canonical if record["winner"] is not None]
        values = [record[mode] for record in settled]
        summary_rows.append([
            f"tier_convert={mode}", len(values), sum(item["won"] for item in values),
            sum(not item["won"] for item in values), sum(item["pnl"] for item in values),
            "排除风控拒绝；同场只计1笔且优先已成交/部分成交；无spread；目标bid，缺失用1-对向ask；不计手续费",
        ])
    summary_rows.append([
        "pre/post方向不同", sum(r["pre"]["side"] != r["post"]["side"] for r in canonical),
        None, None, None, "tier分类全部来自官方market元数据",
    ])
    summary_rows.append([
        "官方未结算", sum(r["winner"] is None for r in canonical), None, None, None,
        "方向明细已写入，不纳入当前利润汇总",
    ])
    replace_rows(wb["重算汇总"], ["口径", "订单数", "赢", "输", "利润", "说明"], summary_rows)

    spread_summary = []
    event_summary = []
    for mode in ("pre", "post"):
        for scope_name, selected in [("风控排除；同场优先成交且仅一笔", canonical)]:
            calculated = []
            unfilled = []
            unresolved = []
            for r in selected:
                result = r[mode]
                canceled = "撤单" in str(r["status"]) and result["side"] == r["historical_side"]
                (unfilled if canceled else unresolved if r["winner"] is None else calculated).append(r)
            def spread_value(r):
                x = r[mode]
                return pnl(x["side"], r["winner"], x["spread_price"], r["quantity"])
            spread_summary.append([
                f"tier_convert={mode}", scope_name, len(selected), len(unfilled), len(unresolved), len(calculated),
                sum(r[mode]["won"] for r in calculated), sum(not r[mode]["won"] for r in calculated),
                sum(spread_value(r) for r in calculated),
            ])
            for lower in (True, False):
                tier_selected = [r for r in selected if r["lower"] == lower]
                tier_unfilled = [
                    r for r in tier_selected
                    if "撤单" in str(r["status"]) and r[mode]["side"] == r["historical_side"]
                ]
                tier_unresolved = [
                    r for r in tier_selected if r not in tier_unfilled and r["winner"] is None
                ]
                tier_calculated = [
                    r for r in tier_selected if r not in tier_unfilled and r not in tier_unresolved
                ]
                event_summary.append([
                    f"tier_convert={mode}", scope_name, "低级别" if lower else "非低级别",
                    len(tier_selected), len(tier_unfilled), len(tier_unresolved), len(tier_calculated),
                    sum(r[mode]["won"] for r in tier_calculated),
                    sum(not r[mode]["won"] for r in tier_calculated),
                    sum(spread_value(r) for r in tier_calculated),
                ])
    replace_rows(wb["spread005汇总"], [
        "tier口径", "风控拒绝口径", "纳入订单行", "撤单同方向未成交", "官方未结算", "计算成交笔数", "赢", "输", "利润",
    ], spread_summary)
    replace_rows(wb["spread005赛事级别汇总"], [
        "tier口径", "风控拒绝口径", "赛事级别", "纳入订单行", "撤单同方向未成交",
        "官方未结算", "计算成交笔数", "赢", "输", "利润",
    ], event_summary)

    if "重算说明_20261004" in wb.sheetnames:
        del wb["重算说明_20261004"]
    notes = wb.create_sheet("重算说明_20261004")
    note_rows = [
        ["项目", "内容"],
        ["订单覆盖", f"{len(records)}/{len(raw_by_id)}"],
        ["总盘纳入", f"{len(canonical)}笔；排除风控拒绝，同场去重并优先已成交>部分成交>撤单"],
        ["赛事分类", "Polymarket 官方 market question/slug；所有 WTA 及 Challenger/UTR/ITF 视为低级别"],
        ["最终胜方", "Polymarket 官方结算 token winner；第一/第二选手映射 yes/no，并校验 token outcome 与 pair 名称"],
        ["tier_ignore", TIER_IGNORE],
        ["无spread价格", "目标方向 bid；缺失时 1-对向 ask；不计手续费"],
        ["spread价格", "max(0.01, limit_bid-0.05)"],
        ["撤单口径", "历史已撤单且历史下单方向=重算策略方向时，spread表记为未成交"],
        ["早期日志", "38笔缺原始触发ask向量，保留旧表已审计的convert/attitude/deviate基础分支，重算tier pre/post"],
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 100
    notes.freeze_panes = "A2"

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)

    old_lower = {oid: str(row["低级别赛事"]) == "是" for oid, row in direction_by_id.items()}
    report = {
        "orders": len(records),
        "canonical_orders": len(canonical),
        "risk_excluded": sum("风控拒绝" in str(r["status"]) for r in records),
        "duplicates_excluded": sum(value.startswith("排除：同场重复") for value in inclusion.values()),
        "lower_tier": sum(r["lower"] for r in records),
        "new_orders": [r["id"] for r in records if r["id"] not in direction_by_id],
        "classification_changed": [
            r["id"] for r in records
            if r["id"] in old_lower and old_lower[r["id"]] != r["lower"]
        ],
        "pre_direction_changed": [
            r["id"] for r in records
            if r["id"] in direction_by_id
            and direction_by_id[r["id"]]["pre口径方向"] != r["pre"]["side"]
        ],
        "post_direction_changed": [
            r["id"] for r in records
            if r["id"] in direction_by_id
            and direction_by_id[r["id"]]["post口径方向"] != r["post"]["side"]
        ],
        "direct": {
            mode: {
                "wins": sum(r[mode]["won"] is True for r in canonical),
                "losses": sum(r[mode]["won"] is False for r in canonical),
                "unresolved": sum(r[mode]["won"] is None for r in canonical),
                "profit": sum(r[mode]["pnl"] or 0.0 for r in canonical),
            }
            for mode in ("pre", "post")
        },
        "spread_summary": spread_summary,
    }
    if args.report:
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
