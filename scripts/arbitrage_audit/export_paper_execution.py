"""把 Polymarket paper SQLite 账本导出为可复算的 Excel 明细。"""

from __future__ import annotations

import argparse
import json
import sqlite3
from collections import defaultdict
from datetime import datetime
from datetime import timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


LOCAL_TZ = ZoneInfo("Asia/Shanghai")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _local_time(*, ns: int | None = None, seconds: int | None = None):
    value = (ns / 1_000_000_000) if ns else seconds
    if not value:
        return None
    return datetime.fromtimestamp(value, tz=timezone.utc).astimezone(LOCAL_TZ).replace(tzinfo=None)


def _pair_parts(pair_id: str | None) -> tuple[str, str, str]:
    parts = str(pair_id or "").split("|")
    return (
        parts[0] if len(parts) > 0 else "",
        parts[1] if len(parts) > 1 else "",
        parts[2] if len(parts) > 2 else "",
    )


def _strategy_type(origin: str | None, intent: str | None) -> str:
    return f"{origin or 'unknown'}/{intent or 'unknown'}"


def _original_or_reverse(outcome: str | None, source_outcome: str | None) -> str | None:
    actual = str(outcome or "").strip().lower()
    source = str(source_outcome or "").strip().lower()
    if actual not in {"yes", "no"} or source not in {"yes", "no"}:
        return None
    return "原单" if actual == source else "反买单"


def _load_audit_rows(audit_json: Path | None) -> dict[str, dict]:
    if audit_json is None:
        return {}
    payload = json.loads(audit_json.read_text(encoding="utf-8"))
    return {
        str(row["client_order_id"]): row
        for row in payload.get("rows", [])
        if row.get("client_order_id")
    }


def _load_rows(database: Path, audit_json: Path | None = None) -> list[dict]:
    audit_rows = _load_audit_rows(audit_json)
    db = sqlite3.connect(f"file:{database}?mode=ro", uri=True)
    db.row_factory = sqlite3.Row
    try:
        orders = [dict(row) for row in db.execute("SELECT * FROM orders")]
        trades = {}
        for row in db.execute("SELECT payload FROM trades"):
            payload = json.loads(row["payload"])
            if payload.get("trader_side") == "MAKER":
                maker_orders = payload.get("maker_orders") or []
                order_id = maker_orders[0].get("order_id") if maker_orders else ""
            else:
                order_id = payload.get("taker_order_id")
            if order_id:
                trades[str(order_id)] = payload
    finally:
        db.close()

    result = []
    for order in orders:
        trade = trades.get(str(order["id"]))
        audit = audit_rows.get(str(order.get("client_order_id") or ""), {})
        pair_id = order.get("pair_id")
        sport, participant_1, participant_2 = _pair_parts(pair_id)
        fill_price = float(trade["price"]) if trade else None
        fill_qty = float(trade["size"]) if trade else 0.0
        result.append({
            **order,
            "sport": sport,
            "participant_1": participant_1,
            "participant_2": participant_2,
            "strategy_type": _strategy_type(order.get("origin"), order.get("intent")),
            "submitted_time": _local_time(
                ns=order.get("ts_submitted_ns"),
                seconds=order.get("created_at"),
            ),
            "fill_time": _local_time(seconds=int(trade["match_time"])) if trade else None,
            "fill_price": fill_price,
            "fill_qty": fill_qty,
            "fill_notional": fill_price * fill_qty if fill_price is not None else 0.0,
            "liquidity": trade.get("trader_side") if trade else None,
            "trade_id": trade.get("id") if trade else None,
            "audit": audit,
        })
    return sorted(
        result,
        key=lambda row: (str(row.get("pair_id") or "~"), row["submitted_time"] or datetime.max),
    )


def _style_sheet(sheet, *, widths: dict[str, int] | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, cell in enumerate(sheet[1], 1):
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(
            str(cell.value),
            15,
        )


def _write_detail(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.active
    sheet.title = "订单明细"
    headers = [
        "提交时间(北京时间)", "首次成交时间(北京时间)", "首次成交比分", "首次成交盘/局",
        "pair_id", "赛事", "选手1", "选手2",
        "策略订单类型(origin/intent)", "origin", "intent", "client_order_id", "venue_order_id",
        "opportunity_id", "leg_key", "instrument_id", "token_id", "outcome", "selection_role",
        "venue_replace前方向", "原单/反买单", "side", "TIF", "post_only", "limit_price", "order_qty",
        "status", "撤单时间(北京时间)",
        "撤单比分", "撤单盘/局", "撤单方向PM_bid", "撤单方向PM_ask",
        "撤单时_PM_YES_bid", "撤单时_PM_YES_ask", "撤单时_PM_NO_bid", "撤单时_PM_NO_ask",
        "撤单时_OE_YES_bid", "撤单时_OE_YES_ask", "撤单时_OE_NO_bid", "撤单时_OE_NO_ask",
        "撤单时_SE_YES_bid", "撤单时_SE_YES_ask", "撤单时_SE_NO_bid", "撤单时_SE_NO_ask",
        "matched_qty", "fill_price", "fill_qty", "fill_notional", "liquidity", "下单比分",
        "strategy_id", "trade_id",
    ]
    sheet.append(headers)
    for row in rows:
        audit = row["audit"]
        sheet.append([
            row["submitted_time"], row["fill_time"], audit.get("首次成交比分"),
            audit.get("首次成交盘/局"), row.get("pair_id"), row["sport"],
            row["participant_1"], row["participant_2"], row["strategy_type"], row.get("origin"),
            row.get("intent"), row.get("client_order_id"), row["id"], row.get("opportunity_id"),
            row.get("leg_key"), row.get("instrument_id"), row["asset_id"], row.get("outcome"),
            row.get("selection_role"), audit.get("venue_replace前方向"),
            _original_or_reverse(
                row.get("outcome") or audit.get("下单方向"),
                audit.get("venue_replace前方向"),
            ),
            row["side"], row["order_type"], bool(row["post_only"]), row["price"],
            row["original_size"], row["status"], audit.get("撤单时间(北京时间)"),
            audit.get("撤单比分"), audit.get("撤单盘/局"), audit.get("撤单方向PM_bid"),
            audit.get("撤单方向PM_ask"),
            *[
                audit.get(f"撤单时_{venue}_{role}_{side}")
                for venue in ("PM", "OE", "SE")
                for role in ("YES", "NO")
                for side in ("bid", "ask")
            ],
            row["size_matched"],
            row["fill_price"], row["fill_qty"], row["fill_notional"], row["liquidity"],
            row.get("order_score"), row.get("strategy_id"), row["trade_id"],
        ])
    header_columns = {str(cell.value): cell.column for cell in sheet[1]}
    datetime_headers = ("提交时间(北京时间)", "首次成交时间(北京时间)", "撤单时间(北京时间)")
    numeric_headers = (
        "limit_price", "order_qty", "撤单方向PM_bid", "撤单方向PM_ask",
        "撤单时_PM_YES_bid", "撤单时_PM_YES_ask", "撤单时_PM_NO_bid", "撤单时_PM_NO_ask",
        "撤单时_OE_YES_bid", "撤单时_OE_YES_ask", "撤单时_OE_NO_bid", "撤单时_OE_NO_ask",
        "撤单时_SE_YES_bid", "撤单时_SE_YES_ask", "撤单时_SE_NO_bid", "撤单时_SE_NO_ask",
        "matched_qty", "fill_price", "fill_qty", "fill_notional",
    )
    for row_index in range(2, sheet.max_row + 1):
        for header in datetime_headers:
            sheet.cell(row_index, header_columns[header]).number_format = "yyyy-mm-dd hh:mm:ss.000"
        for header in numeric_headers:
            sheet.cell(row_index, header_columns[header]).number_format = "0.000000"
    _style_sheet(sheet, widths={
        "提交时间(北京时间)": 24,
        "首次成交时间(北京时间)": 24,
        "首次成交比分": 22,
        "首次成交盘/局": 14,
        "撤单时间(北京时间)": 24,
        "撤单比分": 20,
        "pair_id": 42,
        "选手1": 23,
        "选手2": 23,
        "策略订单类型(origin/intent)": 28,
        "venue_replace前方向": 20,
        "原单/反买单": 16,
        "client_order_id": 20,
        "venue_order_id": 38,
        "opportunity_id": 24,
        "instrument_id": 48,
        "token_id": 30,
        "下单比分": 24,
        "trade_id": 30,
    })


def _write_pair_summary(workbook: Workbook, rows: list[dict]) -> None:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("pair_id") or "(metadata缺失)")].append(row)
    sheet = workbook.create_sheet("Pair汇总")
    sheet.append([
        "pair_id", "赛事", "选手1", "选手2", "订单数", "成交数", "LIVE数", "撤单数",
        "BUY数", "SELL数", "成交数量", "成交名义额", "origin", "intent", "订单类型(origin/intent)",
        "首次提交", "最后提交",
    ])
    for pair_id, items in sorted(grouped.items()):
        sheet.append([
            pair_id, items[0]["sport"], items[0]["participant_1"], items[0]["participant_2"],
            len(items), sum(row["status"] == "MATCHED" for row in items),
            sum(row["status"] == "LIVE" for row in items),
            sum(row["status"] == "CANCELED" for row in items),
            sum(row["side"] == "BUY" for row in items), sum(row["side"] == "SELL" for row in items),
            sum(row["fill_qty"] for row in items), sum(row["fill_notional"] for row in items),
            ", ".join(sorted({str(row.get("origin") or "unknown") for row in items})),
            ", ".join(sorted({str(row.get("intent") or "unknown") for row in items})),
            ", ".join(sorted({row["strategy_type"] for row in items})),
            min(row["submitted_time"] for row in items if row["submitted_time"] is not None),
            max(row["submitted_time"] for row in items if row["submitted_time"] is not None),
        ])
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, 12).number_format = "0.000000"
        sheet.cell(row_index, 16).number_format = "yyyy-mm-dd hh:mm:ss.000"
        sheet.cell(row_index, 17).number_format = "yyyy-mm-dd hh:mm:ss.000"
    _style_sheet(sheet, widths={
        "pair_id": 42,
        "选手1": 23,
        "选手2": 23,
        "订单类型(origin/intent)": 36,
        "首次提交": 24,
        "最后提交": 24,
    })


def _write_type_summary(workbook: Workbook, rows: list[dict]) -> None:
    grouped: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in rows:
        grouped[(str(row.get("origin") or "unknown"), str(row.get("intent") or "unknown"))].append(row)
    sheet = workbook.create_sheet("类型汇总")
    sheet.append([
        "origin", "intent", "订单类型(origin/intent)", "pair数", "订单数", "成交数", "LIVE数",
        "撤单数", "成交数量", "成交名义额",
    ])
    for (origin, intent), items in sorted(grouped.items()):
        sheet.append([
            origin, intent, _strategy_type(origin, intent),
            len({str(row.get("pair_id") or "") for row in items}), len(items),
            sum(row["status"] == "MATCHED" for row in items),
            sum(row["status"] == "LIVE" for row in items),
            sum(row["status"] == "CANCELED" for row in items),
            sum(row["fill_qty"] for row in items), sum(row["fill_notional"] for row in items),
        ])
    _style_sheet(sheet, widths={"订单类型(origin/intent)": 28})


def export_workbook(
    database: Path,
    output: Path,
    audit_json: Path | None = None,
) -> tuple[int, int]:
    rows = _load_rows(database, audit_json)
    workbook = Workbook()
    _write_detail(workbook, rows)
    _write_pair_summary(workbook, rows)
    _write_type_summary(workbook, rows)
    notes = workbook.create_sheet("说明")
    notes.append(["项目", "说明"])
    notes.append(["源数据库", str(database.resolve())])
    notes.append(["导出时间", datetime.now(LOCAL_TZ).replace(tzinfo=None)])
    notes.append(["订单类型", "按 NT Order metadata 的 origin/intent 组合；旧数据库未采集时显示 unknown。"])
    notes.append(["成交定价", "taker 使用提交时对手盘报价；后续 maker 使用订单自身 limit。"])
    notes.append(["归集", "订单明细按 pair_id/提交时间排序；Pair汇总每个 pair 一行。"])
    notes.append(["成交比分", "取首次 OrderFilled 发生前最近一条比分；未成交或日志没有比分时留空。"])
    notes.append(["撤单快照", "取 OrderCanceled 之前最近一次策略可见盘口和最近一条比分；日志没有的数据留空。"])
    _style_sheet(notes, widths={"项目": 20, "说明": 100})
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return len(rows), len({str(row.get("pair_id") or "") for row in rows})


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", type=Path, default=Path("var/polymarket-paper.sqlite3"))
    parser.add_argument("--output", type=Path)
    parser.add_argument("--audit-json", type=Path, help="audit_nohup_orders.py 生成的 JSON")
    args = parser.parse_args(argv)
    output = args.output or (
        Path.home() / "Desktop" / "PM" /
        f"paper订单明细_{datetime.now(LOCAL_TZ):%Y%m%d_%H%M%S}.xlsx"
    )
    order_count, pair_count = export_workbook(args.database, output, args.audit_json)
    print(f"saved={output} orders={order_count} pairs={pair_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
