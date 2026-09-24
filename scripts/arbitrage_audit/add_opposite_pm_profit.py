"""基于成交审计工作簿生成对手 PM 腿利润测算页。"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT = Path("/Users/miller/Desktop/成交订单比分_start_price_venue_stale_window_20260921-22.xlsx")
wb = load_workbook(REPORT)
source = wb["成交明细"]
native = wb["原生PM腿利润测算"]
headers = {cell.value: cell.column for cell in source[1]}

rows = []
for row_index, native_row in enumerate(native.iter_rows(min_row=2, values_only=True), 2):
    (
        trigger_time,
        sport,
        player_1,
        player_2,
        native_role,
        native_outcome,
        winner,
        _,
        quantity,
        _,
        fee_rate,
        *_,
    ) = native_row
    opponent_role = "no" if native_role == "yes" else "yes"
    opponent_outcome = player_2 if native_outcome == player_1 else player_1
    price = float(source.cell(row_index, headers[f"触发时_PM_{opponent_role.upper()}"]).value)
    commission = round(quantity * price * fee_rate * (1.0 - price), 5)
    principal = quantity * price
    won = winner == opponent_outcome
    payout = quantity if won else 0.0
    profit = payout - principal - commission
    rows.append([
        trigger_time,
        sport,
        player_1,
        player_2,
        opponent_role,
        opponent_outcome,
        winner,
        "赢" if won else "输",
        quantity,
        price,
        fee_rate,
        commission,
        principal,
        payout,
        profit,
        native_outcome,
    ])

for name in ("对手PM腿利润汇总", "对手PM腿利润测算"):
    if name in wb.sheetnames:
        del wb[name]

summary = wb.create_sheet("对手PM腿利润汇总", 0)
principal = sum(row[12] for row in rows)
commission = sum(row[11] for row in rows)
payout = sum(row[13] for row in rows)
profit = sum(row[14] for row in rows)
wins = sum(row[7] == "赢" for row in rows)
for item in (
    ("假设", "每次改买 one_side_rebate 原生 PM 腿的相反 outcome，使用同一触发时刻的 PM ask"),
    ("订单数", len(rows)),
    ("每单数量", 5.05),
    ("赢/输", f"{wins}/{len(rows) - wins}"),
    ("本金合计", principal),
    ("手续费合计", commission),
    ("结算收入", payout),
    ("手续费前利润", payout - principal),
    ("净利润", profit),
    ("本金收益率", profit / principal),
    ("成交假设", "全部按触发快照的对手盘 PM ask 足额成交；手续费率沿用同 condition 实际成交反推值"),
):
    summary.append(item)
for cell in summary[1]:
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)
summary.column_dimensions["A"].width = 18
summary.column_dimensions["B"].width = 88
for row_index in range(2, summary.max_row + 1):
    summary.cell(row_index, 2).alignment = Alignment(wrap_text=True)
for row_index in range(5, 10):
    summary.cell(row_index, 2).number_format = "0.000000"
summary.cell(10, 2).number_format = "0.00%"

detail = wb.create_sheet("对手PM腿利润测算", 1)
detail_headers = [
    "触发时间(北京时间)", "赛事", "选手1", "选手2", "对手PM角色", "假设买入选手", "最终胜者",
    "结果", "数量", "对手盘PM ask", "手续费率", "手续费", "本金", "结算收入", "净利润",
    "原生PM腿买入选手",
]
detail.append(detail_headers)
for row in rows:
    detail.append(row)
for cell in detail[1]:
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for row_index in range(2, detail.max_row + 1):
    detail.cell(row_index, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    detail.cell(row_index, 8).fill = PatternFill(
        "solid",
        fgColor="D9EAD3" if detail.cell(row_index, 8).value == "赢" else "F4CCCC",
    )
    for column in range(9, 16):
        detail.cell(row_index, column).number_format = "0.000000"
detail.freeze_panes = "A2"
detail.auto_filter.ref = detail.dimensions
for column, width in enumerate((21, 10, 22, 22, 12, 22, 22, 9, 10, 14, 12, 12, 12, 12, 12, 22), 1):
    detail.column_dimensions[get_column_letter(column)].width = width

wb.save(REPORT)
print(
    f"orders={len(rows)} wins={wins} losses={len(rows) - wins} "
    f"principal={principal:.6f} commission={commission:.6f} payout={payout:.6f} "
    f"profit={profit:.6f} roi={profit / principal:.8f}",
)
