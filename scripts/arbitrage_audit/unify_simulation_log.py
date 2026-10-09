from __future__ import annotations

import ast
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LOG = Path("/private/tmp/remote_nohup_simulation_20261009.out")
SOURCE_BOOK = Path("/Users/miller/Desktop/PM/连续触发策略_全日志明细_20260929-20261007.xlsx")
OUTPUT_BOOK = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_simulation.xlsx")
LOCAL_TZ = timezone(timedelta(hours=8))
ANSI = re.compile(r"\x1b\[[0-9;]*m")
ISO_TS = re.compile(r"(2026-\d\d-\d\dT\d\d:\d\d:\d\d\.\d+Z)")
LOCAL_TS = re.compile(r"^(2026-\d\d-\d\d \d\d:\d\d:\d\d,\d+)")


def timestamp(line: str) -> datetime | None:
    match = ISO_TS.search(line)
    if match:
        return datetime.fromisoformat(match.group(1).replace("Z", "+00:00"))
    match = LOCAL_TS.search(line)
    if match:
        value = datetime.strptime(match.group(1), "%Y-%m-%d %H:%M:%S,%f")
        return value.replace(tzinfo=LOCAL_TZ).astimezone(timezone.utc)
    return None


def clean_score(value: str | None) -> str | None:
    if value is None:
        return None
    return value.split("->")[-1].strip().strip("'")


def winner_from_score(score: str | None, status: str | None) -> str | None:
    if status != "finished":
        return None
    yes_sets = no_sets = 0
    for token in (score or "").split(","):
        match = re.match(r"\s*(\d+)\s*-\s*(\d+)", token)
        if not match:
            continue
        left, right = map(int, match.groups())
        if left > right:
            yes_sets += 1
        elif right > left:
            no_sets += 1
    if yes_sets > no_sets:
        return "yes"
    if no_sets > yes_sets:
        return "no"
    return None


def books_by_venue(snapshot: dict | None) -> dict[tuple[str, str], dict]:
    return {
        (book.get("venue"), book.get("outcome")): book
        for book in (snapshot or {}).get("books", [])
    }


def quote(books: dict, venue: str, outcome: str, side: str):
    return books.get((venue, outcome), {}).get(side)


def format_sheet(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss.000"
    for index, column in enumerate(ws.columns, start=1):
        values = [str(cell.value) if cell.value is not None else "" for cell in column[:200]]
        width = min(max(max((len(value) for value in values), default=8) + 2, 10), 42)
        ws.column_dimensions[get_column_letter(index)].width = width


lines = [ANSI.sub("", line) for line in LOG.read_text(errors="replace").splitlines()]
score_re = re.compile(
    r"game=\d+ (.+?) vs (.+?) score=(.*?) period=([^ ]+) elapsed=.*? status=([^ ]+)$"
)
schedule_re = re.compile(r"pair_id=(.+?), sport=.*?, event=([^,]+), order_books=(\[.*\])$")
start_re = re.compile(r"pair=(.+?) game=\d+ source=([^ ]+) prices=(\{.*\})$")
gate_re = re.compile(
    r"ConsecutiveTriggerGate: pair=(.*?) history=([^ ]+) outcome=(yes|no) score=(.*?) "
    r"count=(\d+) required=(\d+) passed=True"
)
trend_re = re.compile(r"TrendGate: pair=(.*?) drop leg=.*?\.(POLYMARKET|ORBITEXCH) outcome=(yes|no)")
sim_re = re.compile(
    r"PlaceBets\[simulation\]: pair=(.*?) strategy=([^ ]+) rate=([^ ]+) instrument=([^ ]+) "
    r"venue=([^ ]+) role=(yes|no) side=([^ ]+) price=([^ ]+) qty=([^ ]+) "
    r"intent=([^ ]+) opportunity_id=([^ ]+)"
)

latest_score: dict[str, dict] = {}
latest_snapshot: dict[str, dict] = {}
latest_start: dict[str, dict] = {}
latest_gate: dict[str, dict] = {}
latest_trend_drops: dict[str, list[dict]] = defaultdict(list)
all_scores: dict[str, list[dict]] = defaultdict(list)
records: list[dict] = []

for line_no, line in enumerate(lines, start=1):
    ts = timestamp(line)
    if ts is None:
        continue

    if "Sports score updated: game=" in line:
        match = score_re.search(line)
        if match:
            pair = f"Tennis|{match.group(1)}|{match.group(2)}"
            item = {
                "ts": ts,
                "score": clean_score(match.group(3)),
                "period": match.group(4),
                "status": match.group(5),
            }
            latest_score[pair] = item
            all_scores[pair].append(item)
        continue

    if "Strategy evaluate scheduled: pair_id=" in line:
        match = schedule_re.search(line)
        if match:
            pair = match.group(1)
            try:
                books = ast.literal_eval(match.group(3))
            except (SyntaxError, ValueError):
                books = []
            latest_snapshot[pair] = {
                "ts": ts,
                "event": match.group(2),
                "books": books,
            }
            latest_trend_drops[pair] = []
        continue

    if "Start price captured: pair=" in line:
        match = start_re.search(line)
        if match:
            try:
                prices = ast.literal_eval(match.group(3))
            except (SyntaxError, ValueError):
                prices = {}
            latest_start[match.group(1)] = {
                "ts": ts,
                "source": match.group(2),
                **prices,
            }
        continue

    match = trend_re.search(line)
    if match:
        latest_trend_drops[match.group(1)].append(
            {"venue": match.group(2), "outcome": match.group(3)}
        )
        continue

    match = gate_re.search(line)
    if match:
        latest_gate[match.group(1)] = {
            "ts": ts,
            "history": match.group(2),
            "outcome": match.group(3),
            "score": clean_score(match.group(4)),
            "count": int(match.group(5)),
            "required": int(match.group(6)),
        }
        continue

    match = sim_re.search(line)
    if not match:
        continue
    pair = match.group(1)
    snapshot = latest_snapshot.get(pair)
    books = books_by_venue(snapshot)
    gate = latest_gate.get(pair, {})
    score = latest_score.get(pair, {})
    start = latest_start.get(pair, {})
    dropped = latest_trend_drops.get(pair, [])
    dropped_venues = {item["venue"] for item in dropped}
    if dropped_venues == {"POLYMARKET"}:
        source_venue = "ORBITEXCH"
    elif dropped_venues == {"ORBITEXCH"}:
        source_venue = "POLYMARKET"
    else:
        source_venue = None
    original = gate.get("outcome")
    final = match.group(6)
    snapshot_delta_ms = None
    if snapshot:
        snapshot_delta_ms = (ts - snapshot["ts"]).total_seconds() * 1000
    records.append(
        {
            "line_no": line_no,
            "ts": ts,
            "pair": pair,
            "strategy": match.group(2),
            "rate": float(match.group(3)),
            "instrument": match.group(4),
            "venue": match.group(5),
            "final": final,
            "side": match.group(7),
            "price": float(match.group(8)),
            "qty": float(match.group(9)),
            "intent": match.group(10),
            "opportunity_id": match.group(11),
            "original": original,
            "reversed": original is not None and original != final,
            "gate_history": gate.get("history"),
            "gate_score": gate.get("score"),
            "gate_count": gate.get("count"),
            "score": score.get("score"),
            "period": score.get("period"),
            "status": score.get("status"),
            "start_yes": start.get("yes"),
            "start_no": start.get("no"),
            "start_source": start.get("source"),
            "source_venue": source_venue,
            "event": (snapshot or {}).get("event"),
            "snapshot_delta_ms": snapshot_delta_ms,
            "pm_yes_bid": quote(books, "POLYMARKET", "yes", "best_bid"),
            "pm_yes_ask": quote(books, "POLYMARKET", "yes", "best_ask"),
            "pm_no_bid": quote(books, "POLYMARKET", "no", "best_bid"),
            "pm_no_ask": quote(books, "POLYMARKET", "no", "best_ask"),
            "oe_yes_bid": quote(books, "ORBITEXCH", "yes", "best_bid"),
            "oe_yes_ask": quote(books, "ORBITEXCH", "yes", "best_ask"),
            "oe_no_bid": quote(books, "ORBITEXCH", "no", "best_bid"),
            "oe_no_ask": quote(books, "ORBITEXCH", "no", "best_ask"),
        }
    )

if not records:
    raise SystemExit("未解析到 PlaceBets[simulation] 记录")

for record in records:
    _, yes_name, no_name = record["pair"].split("|", 2)
    record["yes_name"] = yes_name
    record["no_name"] = no_name
    record["original_name"] = yes_name if record["original"] == "yes" else no_name if record["original"] == "no" else None
    record["final_name"] = yes_name if record["final"] == "yes" else no_name
    asks = (record["pm_yes_ask"], record["pm_no_ask"])
    record["pm_ask_sum"] = sum(asks) if None not in asks else None

by_pair: dict[str, list[dict]] = defaultdict(list)
for record in records:
    by_pair[record["pair"]].append(record)

wb = load_workbook(SOURCE_BOOK)
for sheet_name in ("sim原始_1008-09", "sim逐场_1008-09", "sim口径_1008-09"):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

raw_headers = [
    "序号", "日志行", "simulation时间(北京时间)", "pair_id", "选手YES", "选手NO",
    "策略", "rate", "连续触发原腿方向", "原腿选手", "原腿venue(推断)",
    "最终模拟方向", "最终模拟选手", "是否反转", "下单价", "数量", "side", "intent",
    "opportunity_id", "触发比分", "比分帧", "盘", "状态", "start_price_YES",
    "start_price_NO", "start来源", "PM_YES_bid", "PM_YES_ask", "PM_NO_bid", "PM_NO_ask",
    "PM_ask_sum", "OE_YES_bid", "OE_YES_ask", "OE_NO_bid", "OE_NO_ask", "触发事件",
    "OBD到simulation间隔_ms", "instrument", "日志来源",
]
raw_ws = wb.create_sheet("sim原始_1008-09")
raw_ws.append(raw_headers)
for index, record in enumerate(records, start=1):
    raw_ws.append([
        index,
        record["line_no"],
        record["ts"].astimezone(LOCAL_TZ).replace(tzinfo=None),
        record["pair"],
        record["yes_name"],
        record["no_name"],
        record["strategy"],
        record["rate"],
        record["original"],
        record["original_name"],
        record["source_venue"],
        record["final"],
        record["final_name"],
        "是" if record["reversed"] else "否",
        record["price"],
        record["qty"],
        record["side"],
        record["intent"],
        record["opportunity_id"],
        record["gate_score"],
        record["score"],
        record["period"],
        record["status"],
        record["start_yes"],
        record["start_no"],
        record["start_source"],
        record["pm_yes_bid"],
        record["pm_yes_ask"],
        record["pm_no_bid"],
        record["pm_no_ask"],
        record["pm_ask_sum"],
        record["oe_yes_bid"],
        record["oe_yes_ask"],
        record["oe_no_bid"],
        record["oe_no_ask"],
        record["event"],
        record["snapshot_delta_ms"],
        record["instrument"],
        LOG.name,
    ])
format_sheet(raw_ws)

pair_headers = [
    "序号", "pair_id", "选手YES", "选手NO", "原始触发数", "方向计数", "原腿方向计数",
    "反转触发数", "首次simulation时间(北京时间)", "末次simulation时间(北京时间)",
    "首次触发比分", "末次触发比分", "首次最终方向", "首次最终选手", "首次价格",
    "最低价格", "最高价格", "start_price_YES", "start_price_NO", "末尾比分", "末尾状态",
    "推断胜者方向", "推断胜者", "真实OrderInitialized数", "归并口径说明",
]
pair_ws = wb.create_sheet("sim逐场_1008-09")
pair_ws.append(pair_headers)
for index, (pair, items) in enumerate(sorted(by_pair.items()), start=1):
    first, last = items[0], items[-1]
    final_state = all_scores.get(pair, [{}])[-1]
    winner = winner_from_score(final_state.get("score"), final_state.get("status"))
    _, yes_name, no_name = pair.split("|", 2)
    pair_ws.append([
        index,
        pair,
        yes_name,
        no_name,
        len(items),
        ", ".join(f"{key}:{value}" for key, value in Counter(item["final"] for item in items).items()),
        ", ".join(f"{key}:{value}" for key, value in Counter(item["original"] for item in items).items()),
        sum(item["reversed"] for item in items),
        first["ts"].astimezone(LOCAL_TZ).replace(tzinfo=None),
        last["ts"].astimezone(LOCAL_TZ).replace(tzinfo=None),
        first["gate_score"],
        last["gate_score"],
        first["final"],
        first["final_name"],
        first["price"],
        min(item["price"] for item in items),
        max(item["price"] for item in items),
        next((item["start_yes"] for item in items if item["start_yes"] is not None), None),
        next((item["start_no"] for item in items if item["start_no"] is not None), None),
        final_state.get("score"),
        final_state.get("status"),
        winner,
        yes_name if winner == "yes" else no_name if winner == "no" else None,
        0,
        "simulation不写订单/仓位；本行按pair归并，首条仅作一场一单分析入口",
    ])
format_sheet(pair_ws)

meta_ws = wb.create_sheet("sim口径_1008-09")
meta_rows = [
    ("项目", "值/说明"),
    ("日志文件", str(LOG)),
    ("日志时间范围(北京时间)", f"{records[0]['ts'].astimezone(LOCAL_TZ)} 至 {records[-1]['ts'].astimezone(LOCAL_TZ)}"),
    ("PlaceBets[simulation]条数", len(records)),
    ("唯一pair数", len(by_pair)),
    ("唯一opportunity_id数", len({record['opportunity_id'] for record in records})),
    ("真实OrderInitialized条数", sum("OrderInitialized" in line for line in lines)),
    ("原始明细口径", "每一条PlaceBets[simulation]日志保留一行；关联此前最近的同pair比分、start_price和order_books帧。"),
    ("逐场口径", "按pair归并；simulation不会写入订单/持仓，因此重复触发不能直接当作真实可累计订单。"),
    ("原腿venue口径", "根据同次评估TrendGate丢弃的venue反推；无法唯一反推时留空。"),
    ("胜者口径", "仅当日志状态为finished时按终局盘分推断；未finished留空，不擅自补结果。"),
    ("数据完整性", "OBD到simulation间隔用于核对关联帧；原始日志行号可回查。"),
]
for row in meta_rows:
    meta_ws.append(row)
format_sheet(meta_ws)
meta_ws.auto_filter.ref = "A1:B1"

wb.save(OUTPUT_BOOK)
print(f"output={OUTPUT_BOOK}")
print(f"records={len(records)} pairs={len(by_pair)} unique_opportunities={len({r['opportunity_id'] for r in records})}")
print(f"with_start={sum(r['start_yes'] is not None and r['start_no'] is not None for r in records)}")
print(f"with_gate={sum(r['original'] is not None for r in records)} with_score={sum(bool(r['gate_score']) for r in records)}")
print(f"snapshot_delta_max_ms={max(r['snapshot_delta_ms'] or 0 for r in records):.3f}")
for pair, items in sorted(by_pair.items()):
    final_state = all_scores.get(pair, [{}])[-1]
    print(
        pair,
        len(items),
        Counter(item['final'] for item in items),
        f"price={min(item['price'] for item in items):.3f}-{max(item['price'] for item in items):.3f}",
        f"start={items[-1]['start_yes']}/{items[-1]['start_no']}",
        f"final={final_state.get('score')} {final_state.get('status')}",
    )
