"""把 nohup 订单审计结果追加为工作簿中的独立真实下单明细 sheet。"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ANSI = re.compile(r"\x1b\[[0-9;]*m")
ISO_TS = re.compile(r"(\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d\.\d+Z)")
LOCAL_TS = re.compile(r"(\d{4}-\d\d-\d\d \d\d:\d\d:\d\d,\d+)")
BEIJING = timezone(timedelta(hours=8))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("audit_json", type=Path)
    parser.add_argument("log", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--sheet", default="远端下单明细_1009")
    parser.add_argument(
        "--append",
        action="store_true",
        help="保留目标 sheet 现有行，并按 client_order_id 去重追加",
    )
    return parser.parse_args()


def line_ts(line: str) -> datetime | None:
    match = ISO_TS.search(line)
    if match:
        return datetime.fromisoformat(match.group(1).replace("Z", "+00:00"))
    match = LOCAL_TS.search(line)
    if match:
        local = datetime.strptime(match.group(1), "%Y-%m-%d %H:%M:%S,%f")
        return local.replace(tzinfo=BEIJING).astimezone(timezone.utc)
    return None


def event_ts(line: str, fallback: datetime) -> datetime:
    match = re.search(r"ts_event=(\d+)", line)
    if not match:
        return fallback
    return datetime.fromtimestamp(int(match.group(1)) / 1_000_000_000, timezone.utc)


def prior(items: list[dict], target: datetime) -> dict | None:
    found = None
    for item in items:
        if item["ts"] <= target:
            found = item
        else:
            break
    return found


def winner_from_score(score: str | None, status: str | None) -> str | None:
    if status != "finished":
        return None
    yes_sets = no_sets = 0
    for token in (score or "").split(","):
        match = re.match(r"\s*(\d+)\s*-\s*(\d+)", token)
        if not match:
            continue
        left, right = map(int, match.groups())
        if left > right:
            yes_sets += 1
        elif right > left:
            no_sets += 1
    if yes_sets > no_sets:
        return "yes"
    if no_sets > yes_sets:
        return "no"
    return None


def consecutive_scores_before(
    observations: list[dict],
    target: datetime,
) -> tuple[str | None, str | None]:
    current_index = None
    for index in range(len(observations) - 1, -1, -1):
        observation = observations[index]
        if observation["ts"] <= target and observation["passed"]:
            if (target - observation["ts"]).total_seconds() <= 5:
                current_index = index
            break
    if current_index is None:
        return None, None

    current = observations[current_index]
    for previous in reversed(observations[:current_index]):
        if previous["outcome"] != current["outcome"]:
            break
        if previous["score"] != current["score"]:
            return previous["score"], current["score"]
    return None, current["score"]


def load_log_context(
    path: Path,
) -> tuple[dict[str, list[dict]], dict[str, datetime], dict[str, list[dict]]]:
    scores: dict[str, list[dict]] = defaultdict(list)
    fill_times: dict[str, datetime] = {}
    consecutive: dict[str, list[dict]] = defaultdict(list)
    score_re = re.compile(
        r"game=\d+ (.+?) vs (.+?) score=(.*?) period=([^ ]+) elapsed=.*? status=([^ ]+)$",
    )
    consecutive_re = re.compile(
        r"ConsecutiveTriggerGate: pair=(.+?) history=.*? outcome=(yes|no) "
        r"score='(.*?)' count=\d+ required=\d+ passed=(True|False)",
    )
    with path.open(errors="replace") as stream:
        for raw in stream:
            line = ANSI.sub("", raw.rstrip())
            ts = line_ts(line)
            if ts is None:
                continue
            match = score_re.search(line)
            if match:
                pair = f"Tennis|{match.group(1)}|{match.group(2)}"
                scores[pair].append(
                    {
                        "ts": ts,
                        "score": match.group(3).split("->")[-1],
                        "period": match.group(4),
                        "status": match.group(5),
                    },
                )
            match = consecutive_re.search(line)
            if match:
                consecutive[match.group(1)].append({
                    "ts": ts,
                    "outcome": match.group(2),
                    "score": match.group(3),
                    "passed": match.group(4) == "True",
                })
            if "OrderFilled(" in line and "client_order_id=ARB-" in line:
                order_match = re.search(r"client_order_id=(ARB-[^,]+)", line)
                if order_match:
                    fill_times.setdefault(order_match.group(1), event_ts(line, ts))
    for collection in (scores, consecutive):
        for items in collection.values():
            items.sort(key=lambda item: item["ts"])
    return scores, fill_times, consecutive


def collapse_risk_rejections(rows: list[dict]) -> list[dict]:
    """风控拒绝未到 venue，同一赛事只留首条，次数和 ID 写进备注。"""
    rejected: dict[tuple[str, str], list[dict]] = defaultdict(list)
    kept: list[dict] = []
    for row in rows:
        if row["状态"] == "未成交-风控拒绝":
            rejected[(row["选手1"], row["选手2"])].append(row)
        else:
            kept.append(row)
    for group in rejected.values():
        first = dict(group[0])
        ids = ", ".join(item["client_order_id"] for item in group)
        first["合并备注"] = f"同场风控拒绝初始化 {len(group)} 次；订单ID：{ids}"
        kept.append(first)
    return sorted(kept, key=lambda row: row["下单时间(北京时间)"])


def format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss.000"
            elif isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for column_number, column in enumerate(sheet.columns, 1):
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column_number)].width = min(max(width + 2, 11), 42)


def main() -> None:
    args = parse_args()
    audit = json.loads(args.audit_json.read_text(encoding="utf-8"))
    rows = collapse_risk_rejections(audit["rows"])
    scores, fill_times, consecutive = load_log_context(args.log)

    headers = [
        "下单时间(北京时间)", "client_order_id", "赛事", "状态", "买入方向", "买入选手",
        "限价", "下单数量", "已成交量", "成交均价", "首次成交时间(北京时间)",
        "连续触发比分1", "连续触发比分2", "下单比分", "下单盘数",
        "首次成交比分", "首次成交盘数", "最终比分", "最终状态",
        "胜者方向", "胜者", "start_price_YES", "start_price_NO", "start来源", "触发类型",
        "原生腿venue", "venue_replace前方向", "是否反买", "触发venue",
        "变化前_PM_YES", "变化前_PM_NO", "变化前_OE_YES", "变化前_OE_NO",
        "触发时_PM_YES", "触发时_PM_NO", "触发时_OE_YES", "触发时_OE_NO",
        "另一venue变化间隔(秒)", "后来变化venue", "合并备注",
    ]

    workbook = load_workbook(args.workbook)
    existing_ids: set[str] = set()
    if args.append and args.sheet in workbook.sheetnames:
        sheet = workbook[args.sheet]
        existing_headers = [cell.value for cell in sheet[1]]
        if existing_headers != headers:
            raise ValueError(f"目标 sheet 列结构不一致: {args.sheet}")
        order_id_column = headers.index("client_order_id") + 1
        existing_ids = {
            str(sheet.cell(row, order_id_column).value)
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, order_id_column).value
        }
    else:
        if args.sheet in workbook.sheetnames:
            del workbook[args.sheet]
        sheet = workbook.create_sheet(args.sheet, 1)
        sheet.append(headers)

    added_rows = 0
    for row in rows:
        if row["client_order_id"] in existing_ids:
            continue
        pair = f"Tennis|{row['选手1']}|{row['选手2']}"
        pair_scores = scores.get(pair, [])
        order_time = datetime.strptime(
            row["下单时间(北京时间)"],
            "%Y-%m-%d %H:%M:%S.%f",
        )
        order_time_utc = order_time.replace(tzinfo=BEIJING).astimezone(timezone.utc)
        trigger_score_1, trigger_score_2 = consecutive_scores_before(
            consecutive.get(pair, []),
            order_time_utc,
        )
        fill_time = fill_times.get(row["client_order_id"])
        fill_score = prior(pair_scores, fill_time) if fill_time else None
        final = pair_scores[-1] if pair_scores else None
        winner = winner_from_score(
            final.get("score") if final else None,
            final.get("status") if final else None,
        )
        direction = row["下单方向"]
        winner_name = row["选手1"] if winner == "yes" else row["选手2"] if winner == "no" else None
        native_direction = row.get("venue_replace前方向")
        values = {
            "下单时间(北京时间)": order_time,
            "client_order_id": row["client_order_id"],
            "赛事": f"{row['选手1']} vs {row['选手2']}",
            "状态": row["状态"],
            "买入方向": direction,
            "买入选手": row["选手1"] if direction == "yes" else row["选手2"],
            "限价": row["限价"],
            "下单数量": row["下单数量"],
            "已成交量": row["已成交量"],
            "成交均价": row["成交均价"],
            "首次成交时间(北京时间)": (
                fill_time.astimezone(BEIJING).replace(tzinfo=None) if fill_time else None
            ),
            "连续触发比分1": trigger_score_1,
            "连续触发比分2": trigger_score_2,
            "下单比分": row["下单比分"],
            "下单盘数": row["盘/局"],
            "首次成交比分": fill_score.get("score") if fill_score else None,
            "首次成交盘数": fill_score.get("period") if fill_score else None,
            "最终比分": final.get("score") if final else None,
            "最终状态": final.get("status") if final else None,
            "胜者方向": winner,
            "胜者": winner_name,
            "start_price_YES": row["start_yes"],
            "start_price_NO": row["start_no"],
            "start来源": row["start来源"],
            "触发类型": row["触发类型"] or "连续触发",
            "原生腿venue": row["原生腿venue"],
            "venue_replace前方向": native_direction,
            "是否反买": (
                native_direction != direction if native_direction in {"yes", "no"} else None
            ),
            "触发venue": row["触发venue"],
            "变化前_PM_YES": row["变化前_PM_YES"],
            "变化前_PM_NO": row["变化前_PM_NO"],
            "变化前_OE_YES": row["变化前_OE_YES"],
            "变化前_OE_NO": row["变化前_OE_NO"],
            "触发时_PM_YES": row["触发时_PM_YES"],
            "触发时_PM_NO": row["触发时_PM_NO"],
            "触发时_OE_YES": row["触发时_OE_YES"],
            "触发时_OE_NO": row["触发时_OE_NO"],
            "另一venue变化间隔(秒)": row["另一venue变化间隔(秒)"],
            "后来变化venue": row["后来变化venue"],
            "合并备注": row.get("合并备注"),
        }
        sheet.append([values.get(header) for header in headers])
        added_rows += 1

    format_sheet(sheet)
    workbook.save(args.output)
    print(json.dumps({
        "source_orders": len(audit["rows"]),
        "detail_rows": len(rows),
        "risk_rejections_collapsed": len(audit["rows"]) - len(rows),
        "added_rows": added_rows,
        "total_sheet_rows": sheet.max_row - 1,
        "sheet": args.sheet,
        "output": str(args.output),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
