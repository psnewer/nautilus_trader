"""按当前 venue_replace 配置重算历史下单方向。

该脚本只重放决定方向的分支，不重放 score/share/risk 等前置门控。历史订单已经通过
这些门控；输入方向取 Excel 的 ``venue_replace前`` 列，盘口取下单时保存值。
"""

from __future__ import annotations

import argparse
import gzip
import json
import math
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


LOWER_TIER_NAME_MARKERS = ("CHALLENGER", "WTA", "UTR", "ITF")
LOWER_TIER_TOURNAMENTS = {
    "Adana",
    "Bari",
    "Columbus",
    "Curitiba",
    "Jingshan",
    "Mouilleron-Le-Captif",
    "Plovdiv 4",
    "Porto 2",
    "San Diego 2",
    "St. Tropez",
    "Tolentino",
}
ORDER_PATTERN = re.compile(
    r"OrderInitialized\(instrument_id=(0x[0-9a-f]+)-[^,]+, client_order_id=([^,]+),",
)
TIER_PATTERN = re.compile(
    r"VenueReplace: pair=(.+?) tier_convert(?:=(?:pre|post))? "
    r"(?:hit|eligible) oe_competition=(.+)$",
)


def opposite(side: str) -> str:
    return "no" if side == "yes" else "yes"


def positive_float(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) and result > 0 else None


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = text.rsplit(":", 1)[-1]
    text = re.sub(r"\bluis guto miguel\b", "luis miguel", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def load_audits(paths: list[Path]) -> dict[str, dict]:
    result = {}
    for path in paths:
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("rows", []) if isinstance(payload, dict) else payload
        for row in rows:
            if isinstance(row, dict) and row.get("client_order_id"):
                result[str(row["client_order_id"])] = row
    return result


def scan_logs(paths: list[Path]) -> tuple[dict[str, str], dict[str, str]]:
    condition_by_order = {}
    lower_competition_by_pair = {}
    for path in paths:
        opener = gzip.open if path.suffix == ".gz" else open
        with opener(path, "rt", encoding="utf-8", errors="replace") as stream:
            for line in stream:
                order_match = ORDER_PATTERN.search(line)
                if order_match:
                    condition_by_order[order_match.group(2)] = order_match.group(1)
                tier_match = TIER_PATTERN.search(line)
                if tier_match:
                    lower_competition_by_pair[tier_match.group(1)] = tier_match.group(2)
    return condition_by_order, lower_competition_by_pair


def position_metadata(path: Path) -> dict[str, list[dict]]:
    result: dict[str, list[dict]] = defaultdict(list)
    for row in json.loads(path.read_text(encoding="utf-8")):
        result[normalize(row.get("title", ""))].append(row)
    return result


def position_is_lower(rows: list[dict]) -> bool:
    for row in rows:
        slug = str(row.get("slug", "")).upper()
        title = str(row.get("title", ""))
        tournament = title.split(":", 1)[0]
        if any(marker in slug for marker in LOWER_TIER_NAME_MARKERS):
            return True
        if tournament in LOWER_TIER_TOURNAMENTS:
            return True
    return False


def infer_native_venue(order_id: str, action: str, audit: dict[str, dict]) -> tuple[str | None, str]:
    row = audit.get(order_id, {})
    venue = str(row.get("原生腿venue") or "").upper()
    if venue:
        return venue, "audit_json"
    # 旧日志已不存在时，只把明确标注的 convert 当作原生 PM；其余不臆测具体外部 venue。
    if "venue_replace convert" in action:
        return "POLYMARKET", "legacy_action"
    return "EXTERNAL", "legacy_action"


def is_lower_tier(
    match_name: str,
    pair_id: str,
    action: str,
    position_rows: list[dict],
    lower_competitions: dict[str, str],
    explicit_lower_pairs: set[str],
) -> tuple[bool, str]:
    if pair_id in explicit_lower_pairs:
        return True, "explicit_log_extract"
    if pair_id in lower_competitions:
        return True, f"log:{lower_competitions[pair_id]}"
    if "tier_convert" in action:
        return True, "legacy_action"
    if position_is_lower(position_rows):
        return True, "position_metadata"
    return False, "not_identified"


def replay_direction(order: dict) -> tuple[str, str]:
    source = order["source_side"]
    start = order["start"]
    bid = order["bid"]
    ask_sum = order["ask_sum"]
    if order["lower_tier"] and start is not None and start >= 0.3:
        return opposite(source), "tier_convert pre"
    if order["native_venue"] == "POLYMARKET":
        return opposite(source), "venue_replace convert"
    commission_ok = ask_sum is not None and 0.98 <= ask_sum <= 1.02
    if commission_ok and start is not None and bid is not None:
        if bid <= start:
            return opposite(source), "attitude"
        if bid >= 1.2 * start:
            return opposite(source), "deviate_convert"
    return source, "same outcome"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("positions", type=Path)
    parser.add_argument("--audit", type=Path, action="append", default=[])
    parser.add_argument("--log", type=Path, action="append", default=[])
    parser.add_argument(
        "--lower-pair",
        action="append",
        default=[],
        help="远端日志只提取未下载时，显式传入已证实的低级别 pair_id",
    )
    parser.add_argument("--details", type=Path)
    args = parser.parse_args()

    audits = load_audits(args.audit)
    _, lower_competitions = scan_logs(args.log)
    positions = position_metadata(args.positions)
    sheet = load_workbook(args.workbook, read_only=True, data_only=True)[
        "nohup成交_统一规则"
    ]
    headers = [str(value) for value in next(sheet.iter_rows(values_only=True))]
    records = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=True))
        order_id = str(row["client_order_id"] or "")
        if not order_id or row["比赛"] == "合计":
            continue
        source = str(row["规则输入方向(venue_replace前)"] or "").lower()
        if source not in {"yes", "no"}:
            continue
        match_name = str(row["比赛"])
        player_1, player_2 = match_name.split(" vs ", 1)
        pair_id = f"Tennis|{player_1}|{player_2}"
        action = str(row["策略动作"] or "")
        native_venue, native_source = infer_native_venue(order_id, action, audits)
        lower, lower_source = is_lower_tier(
            match_name,
            pair_id,
            action,
            positions.get(normalize(match_name), []),
            lower_competitions,
            set(args.lower_pair),
        )
        audit_row = audits.get(order_id, {})
        trigger_yes_ask = positive_float(audit_row.get("触发时_PM_YES"))
        trigger_no_ask = positive_float(audit_row.get("触发时_PM_NO"))
        source_ask = positive_float(row["current_PM_ask(参考)"])
        final_side = str(row["OrderInitialized最终方向"] or "").lower()
        final_ask = positive_float(row["最终方向PM_ask(下单时)"])
        # 两列方向相反时可恢复完整 PM ask 向量；同向时历史表没有保存另一侧 ask。
        ask_sum = (
            source_ask + final_ask
            if source_ask is not None
            and final_ask is not None
            and final_side in {"yes", "no"}
            and final_side != source
            else None
        )
        # 旧报表把 start 向量之和误写为 commission 依据。缺另一侧当前 ask 时，只有该值
        # 位于合法区间才允许动态分支；结果标成 fallback，便于结果审计。
        ask_sum_source = "current_asks"
        if trigger_yes_ask is not None and trigger_no_ask is not None:
            ask_sum = trigger_yes_ask + trigger_no_ask
            ask_sum_source = "audit_trigger_asks"
        elif ask_sum is None:
            ask_sum = positive_float(row["start_yes+start_no"])
            ask_sum_source = "start_sum_fallback"
        record = {
            "id": order_id,
            "match": match_name,
            "status": row["订单状态"],
            "source_side": source,
            "start": positive_float(row["start_price"]),
            "bid": positive_float(row["current_PM_bid(规则判断价)"]),
            "ask_sum": ask_sum,
            "ask_sum_source": ask_sum_source,
            "native_venue": native_venue,
            "native_source": native_source,
            "lower_tier": lower,
            "lower_source": lower_source,
            "logged_side": final_side,
        }
        record["replayed_side"], record["reason"] = replay_direction(record)
        record["changed"] = record["replayed_side"] != final_side
        records.append(record)

    summary = {
        "orders": len(records),
        "replayed_yes": sum(row["replayed_side"] == "yes" for row in records),
        "replayed_no": sum(row["replayed_side"] == "no" for row in records),
        "different_from_logged": sum(row["changed"] for row in records),
        "lower_tier": sum(row["lower_tier"] for row in records),
        "native_from_legacy_action": sum(row["native_source"] == "legacy_action" for row in records),
        "commission_start_sum_fallback": sum(
            row["ask_sum_source"] == "start_sum_fallback" for row in records
        ),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if args.details:
        args.details.write_text(
            json.dumps({"summary": summary, "rows": records}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


if __name__ == "__main__":
    main()
