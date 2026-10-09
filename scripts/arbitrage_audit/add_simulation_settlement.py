from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_simulation.xlsx")
OUTPUT = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_simulation_结算.xlsx")

# 9 场来自日志终局比分；Miguel Tobon 一场日志在赛点前结束，外部赛果为 YES 2-1。
WINNERS = {
    "Tennis|Carlo Alberto Caniato|Francesco Passaro": ("no", "日志终局 6-4, 3-6, 4-6"),
    "Tennis|Daniel Rincon|Pedro Martinez": ("no", "日志终局 5-7, 4-6"),
    "Tennis|Gabriele Piraino|Enrico Dalla Valle": ("no", "日志终局 4-6, 3-6"),
    "Tennis|Gustavo Heide|Eduardo Ribeiro": ("no", "日志终局 6-3, 3-6, 4-6"),
    "Tennis|Joaquin Aguilar|Lautaro Midon": ("no", "日志终局 6-7(1-7), 6-2, 6-7(9-11)"),
    "Tennis|Juan Carlos Prado|Pedro Boscardin Dias": ("yes", "日志末尾 6-2, 6-3（胜负已确定）"),
    "Tennis|Max Hans Rehberg|Hynek Barton": ("no", "日志终局 1-6, 7-6(7-1), 6-7(3-7)"),
    "Tennis|Miguel Tobon|Tomas Barrios": (
        "yes",
        "外部赛果 Miguel Tobon 2-1 Tomas Barrios；日志停在决胜盘抢七",
    ),
    "Tennis|Pablo Llamas Ruiz|Petr Nesterov": ("yes", "日志终局 3-6, 7-6(7-5), 6-3"),
    "Tennis|Vilius Gaubas|Duje Ajdukovic": ("yes", "日志末尾 5-7, 7-6(7-4), 6-2（胜负已确定）"),
}


def style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for i, column in enumerate(ws.columns, start=1):
        width = min(max(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 10), 45)
        ws.column_dimensions[get_column_letter(i)].width = width


wb = load_workbook(SOURCE)
if "sim结算_首单" in wb.sheetnames:
    del wb["sim结算_首单"]
ws = wb.create_sheet("sim结算_首单")
headers = [
    "序号", "pair_id", "买入方向", "买入选手", "simulation时间(北京时间)", "触发比分",
    "买入价格(PM ask)", "数量", "胜者方向", "胜者", "胜负", "毛利润", "手续费",
    "净利润", "赛果依据", "计数口径",
]
ws.append(headers)

raw = wb["sim原始_1008-09"]
index = {cell.value: i for i, cell in enumerate(raw[1])}
first_by_pair = {}
for row in raw.iter_rows(min_row=2, values_only=True):
    first_by_pair.setdefault(row[index["pair_id"]], row)

gross_total = fee_total = 0.0
wins = losses = 0
for number, (pair, row) in enumerate(first_by_pair.items(), start=1):
    direction = row[index["最终模拟方向"]]
    price = float(row[index["下单价"]])
    qty = float(row[index["数量"]])
    winner, basis = WINNERS[pair]
    _, yes_name, no_name = pair.split("|", 2)
    winner_name = yes_name if winner == "yes" else no_name
    won = direction == winner
    gross = qty * ((1 - price) if won else -price)
    fee = qty * price * (1 - price) * 0.05
    net = gross - fee
    gross_total += gross
    fee_total += fee
    wins += int(won)
    losses += int(not won)
    ws.append([
        number,
        pair,
        direction,
        row[index["最终模拟选手"]],
        row[index["simulation时间(北京时间)"]],
        row[index["触发比分"]],
        price,
        qty,
        winner,
        winner_name,
        "赢" if won else "输",
        gross,
        fee,
        net,
        basis,
        "每场仅取首条PlaceBets[simulation]；重复触发不重复计单",
    ])

ws.append([])
ws.append(["汇总", None, None, None, None, None, None, len(first_by_pair), None, None, f"{wins}胜{losses}负", gross_total, fee_total, gross_total - fee_total])
style(ws)

meta = wb["sim口径_1008-09"]
meta.append(("胜负与利润口径", "每场只取首条simulation；日志价格等于该方向PM ask，按立即成交处理；share=10。"))
meta.append(("毛利润", gross_total))
meta.append(("手续费模型", "qty × price × (1-price) × 5%；另列毛利润，便于不收手续费时使用。"))
meta.append(("净利润", gross_total - fee_total))

wb.save(OUTPUT)
print(f"output={OUTPUT}")
print(f"matches={len(first_by_pair)} wins={wins} losses={losses}")
print(f"gross={gross_total:.6f} fee={fee_total:.6f} net={gross_total-fee_total:.6f}")
