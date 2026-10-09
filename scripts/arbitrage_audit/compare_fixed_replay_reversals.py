"""固定 draw|win 真实回放信号，对比无反转与 VenueReplace 动态反转。

本脚本不重新选择连续触发机会。输入行固定为 ``replay_draw_win_enriched.json``
中的既有真实回放结果，只从原日志补齐同一时刻的完整报价、原腿 venue 和
信号发生前最近的 start_price。
"""

from __future__ import annotations

import ast
import gzip
import json
import re
import subprocess
from collections import defaultdict
from datetime import datetime
from datetime import timedelta
from datetime import timezone
from itertools import product
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


SOURCE_BOOK = Path(
    "/Users/miller/Desktop/PM/连续触发策略_全日志明细_20260929-20261009_simulation.xlsx",
)
OUTPUT_BOOK = Path(
    "/private/tmp/连续触发策略_全日志明细_20260929-20261009_固定回放反转对照.xlsx",
)
REPLAY_ROWS = Path("/private/tmp/replay_draw_win_enriched.json")
EVALS = Path("/private/tmp/tennis_evals.txt.gz")
LOG_PATHS = (
    Path("/Users/miller/Downloads/nohup.out"),
    Path("/private/tmp/remote_log2_20261007.out"),
    Path("/private/tmp/remote_log_20261007_latest.out"),
    Path("/private/tmp/remote_log1_20261007.out"),
    Path("/private/tmp/remote_nohup_20261007_1601.out"),
    Path("/private/tmp/remote_log3_20261008.out"),
    Path("/private/tmp/remote_nohup_20261008_latest.out"),
)
MIN_RATE = 0.08
MAX_PRICE = 0.8
FEE_RATE = 0.05
QTY = 10.0
LOCAL_TZ = timezone(timedelta(hours=8))

ANSI = re.compile(r"\x1b\[[0-9;]*m")
ISO = re.compile(r"(2026-\d\d-\d\dT\d\d:\d\d:\d\d\.\d+)Z")
LOCAL = re.compile(r"^(2026-\d\d-\d\d \d\d:\d\d:\d\d,\d+)")
SCHEDULED = re.compile(
    r"evaluate scheduled: pair_id=(Tennis\|.+?), sport=.*?, "
    r"event=([^,]+), order_books=(\[.*\])$",
)


def timestamp_ns(line: str) -> int | None:
    match = ISO.search(line)
    if match:
        return int(datetime.fromisoformat(match.group(1) + "+00:00").timestamp() * 1e9)
    match = LOCAL.search(line)
    if match:
        value = datetime.strptime(match.group(1), "%Y-%m-%d %H:%M:%S,%f")
        return int(value.replace(tzinfo=LOCAL_TZ).timestamp() * 1e9)
    return None


def quote_map(books: list[dict]) -> dict[tuple[str, str], dict]:
    return {
        (str(book.get("venue")), str(book.get("outcome"))): book
        for book in books
    }


def exact_snapshots(rows: list[dict]) -> dict[tuple[str, int], list[dict]]:
    targets = {(row["pair"], int(row["ts"])) for row in rows}
    snapshots = {}
    with gzip.open(EVALS, "rt", errors="replace") as stream:
        for line in stream:
            if "evaluate scheduled: pair_id=Tennis|" not in line:
                continue
            match = SCHEDULED.search(line.rstrip())
            if match is None:
                continue
            key = (match.group(1), timestamp_ns(line))
            if key not in targets:
                continue
            try:
                snapshots[key] = ast.literal_eval(match.group(3))
            except (SyntaxError, ValueError):
                continue
    missing = sorted(pair for pair, ts in targets if (pair, ts) not in snapshots)
    if missing:
        raise RuntimeError("固定回放行缺少精确行情帧:\n" + "\n".join(missing))
    return snapshots


def pm_quote_series(target_pairs: set[str]) -> dict[str, list[tuple[int, dict[str, dict]]]]:
    result: dict[str, list[tuple[int, dict[str, dict]]]] = defaultdict(list)
    with gzip.open(EVALS, "rt", errors="replace") as stream:
        for line in stream:
            if "evaluate scheduled: pair_id=Tennis|" not in line:
                continue
            match = SCHEDULED.search(line.rstrip())
            ts = timestamp_ns(line)
            if match is None or ts is None or match.group(1) not in target_pairs:
                continue
            try:
                books = ast.literal_eval(match.group(3))
            except (SyntaxError, ValueError):
                continue
            quotes = {
                str(book.get("outcome")): book
                for book in books
                if book.get("venue") == "POLYMARKET"
            }
            if quotes:
                result[match.group(1)].append((ts, quotes))
    for pair, values in result.items():
        unique = {
            (
                ts,
                quotes.get("yes", {}).get("best_bid"),
                quotes.get("yes", {}).get("best_ask"),
                quotes.get("no", {}).get("best_bid"),
                quotes.get("no", {}).get("best_ask"),
            ): (ts, quotes)
            for ts, quotes in values
        }
        result[pair] = sorted(unique.values(), key=lambda item: item[0])
    return result


def start_prices(target_pairs: set[str]) -> dict[str, list[tuple[int, dict]]]:
    paths = [str(path) for path in LOG_PATHS if path.exists()]
    process = subprocess.run(
        ["rg", "--no-filename", "Start price captured: pair=", *paths],
        capture_output=True,
        text=True,
        errors="replace",
        check=False,
    )
    if process.returncode not in (0, 1):
        raise RuntimeError(process.stderr)
    result: dict[str, list[tuple[int, dict]]] = defaultdict(list)
    for raw in process.stdout.splitlines():
        line = ANSI.sub("", raw)
        match = re.search(
            r"Start price captured: pair=(.+?) game=\d+ source=[^ ]+ prices=(\{.*\})$",
            line,
        )
        ts = timestamp_ns(line)
        if match is None or ts is None or match.group(1) not in target_pairs:
            continue
        try:
            prices = ast.literal_eval(match.group(2))
        except (SyntaxError, ValueError):
            continue
        result[match.group(1)].append((ts, prices))
    for pair, values in result.items():
        # 多份日志有重叠，同一时间只保留一份。
        result[pair] = sorted({ts: prices for ts, prices in values}.items())
    return result


def latest_start(series: list[tuple[int, dict]], ts: int) -> dict:
    current = {}
    for observed_at, prices in series:
        if observed_at > ts:
            break
        current = prices
    return current


def source_venue(books: list[dict], outcome: str) -> str:
    """复现既有回放 harness 的候选顺序，找出最终候选的原生 venue。

    harness 的 instrument 顺序固定为 PM yes/no、OE yes/no；ShareLimit 把候选
    share 统一压到10，CandiSelect 因此选择首个满足 min_rate 的候选。
    """
    quotes = quote_map(books)
    legs = {
        role: [
            ("PM", quotes.get(("POLYMARKET", role), {}).get("best_ask")),
            ("OE", quotes.get(("ORBITEXCH", role), {}).get("best_ask")),
        ]
        for role in ("yes", "no")
    }
    for yes_leg, no_leg in product(legs["yes"], legs["no"]):
        if yes_leg[1] is None or no_leg[1] is None:
            continue
        total = float(yes_leg[1]) + float(no_leg[1])
        for target_role, target_leg in (("yes", yes_leg), ("no", no_leg)):
            rate = (1.0 - total) / float(target_leg[1])
            if rate >= MIN_RATE:
                return yes_leg[0] if outcome == "yes" else no_leg[0]
    raise RuntimeError(f"无法从固定回放行情还原原腿 venue: outcome={outcome} books={books}")


def result(direction: str, price: float | None, winner: str | None) -> dict:
    excluded = None
    if price is None:
        excluded = "目标方向PM ask缺失"
    elif price > MAX_PRICE:
        excluded = "最终PM ask>0.8"
    elif winner not in ("yes", "no"):
        excluded = "赛果缺失"
    won = None if excluded else direction == winner
    gross = None if excluded else QTY * ((1.0 - price) if won else -price)
    fee = None if excluded else QTY * price * (1.0 - price) * FEE_RATE
    return {
        "excluded": excluded,
        "won": won,
        "gross": gross,
        "fee": fee,
        "net": None if excluded else gross - fee,
    }


def style(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for index, column in enumerate(ws.columns, 1):
        width = max((len(str(cell.value or "")) for cell in column), default=8) + 2
        ws.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 44)


rows = json.loads(REPLAY_ROWS.read_text())
snapshots = exact_snapshots(rows)
starts = start_prices({row["pair"] for row in rows})
future_quotes = pm_quote_series({row["pair"] for row in rows})
details = []
for row in rows:
    pair = row["pair"]
    ts = int(row["ts"])
    original = row["role"]
    opposite = "no" if original == "yes" else "yes"
    books = snapshots[(pair, ts)]
    quotes = quote_map(books)
    pm = {
        role: quotes.get(("POLYMARKET", role), {})
        for role in ("yes", "no")
    }
    start = latest_start(starts.get(pair, []), ts)
    venue = source_venue(books, original)
    bid = pm[original].get("best_bid")
    original_ask = pm[original].get("best_ask")
    opposite_ask = pm[opposite].get("best_ask")
    ask_values = [pm[role].get("best_ask") for role in ("yes", "no")]
    commission = sum(ask_values) if all(value is not None for value in ask_values) else None
    commission_ok = commission is not None and 0.98 <= commission <= 1.02
    start_price = start.get(original)

    action = "默认同方向替换"
    final = original
    if venue == "PM":
        action = "convert反转"
        final = opposite
    elif commission_ok and bid is not None and start_price is not None:
        if float(bid) <= float(start_price):
            action = "attitude反转"
            final = opposite
        elif 1.2 * float(start_price) <= float(bid) <= 1.3 * float(start_price):
            action = "deviate_convert反转（1.2-1.3）"
            final = opposite
        elif float(bid) > 1.3 * float(start_price):
            action = "默认同方向替换（deviate>1.3）"
    if venue == "OE" and start_price is None:
        action = "默认同方向替换（缺start_price）"
    elif venue == "OE" and not commission_ok:
        action = "默认同方向替换（commission不合格）"

    final_ask = pm[final].get("best_ask")
    final_bid = pm[final].get("best_bid")
    bid_minus_price = (
        round(float(final_bid) - 0.05, 3)
        if final_bid is not None and float(final_bid) > 0.05
        else None
    )
    bid_minus_fill_ts = None
    bid_fill_ts = None
    if final_bid is not None or bid_minus_price is not None:
        for later_ts, later_quotes in future_quotes.get(pair, []):
            later_ask = later_quotes.get(final, {}).get("best_ask")
            if later_ts <= ts or later_ask is None:
                continue
            if (
                bid_fill_ts is None
                and final_bid is not None
                and float(later_ask) <= float(final_bid) + 1e-12
            ):
                bid_fill_ts = later_ts
            if (
                bid_minus_fill_ts is None
                and bid_minus_price is not None
                and float(later_ask) <= bid_minus_price + 1e-12
            ):
                bid_minus_fill_ts = later_ts
            if bid_fill_ts is not None and (
                bid_minus_price is None or bid_minus_fill_ts is not None
            ):
                break
    baseline = result(original, original_ask, row.get("winner"))
    reversed_result = result(final, final_ask, row.get("winner"))
    bid_minus_excluded = None
    if bid_minus_price is None:
        bid_minus_excluded = "目标方向PM bid缺失或≤0.05"
    elif bid_minus_price > MAX_PRICE:
        bid_minus_excluded = "bid-0.05>0.8"
    elif bid_minus_fill_ts is None:
        bid_minus_excluded = "后续PM ask未触及限价"
    elif row.get("winner") not in ("yes", "no"):
        bid_minus_excluded = "赛果缺失"
    bid_minus_won = (
        None if bid_minus_excluded else final == row.get("winner")
    )
    bid_minus_gross = (
        None
        if bid_minus_excluded
        else QTY * ((1.0 - bid_minus_price) if bid_minus_won else -bid_minus_price)
    )
    bid_excluded = None
    if final_bid is None:
        bid_excluded = "目标方向PM bid缺失"
    elif float(final_bid) > MAX_PRICE:
        bid_excluded = "bid>0.8"
    elif bid_fill_ts is None:
        bid_excluded = "后续PM ask未触及限价"
    elif row.get("winner") not in ("yes", "no"):
        bid_excluded = "赛果缺失"
    bid_won = None if bid_excluded else final == row.get("winner")
    bid_gross = (
        None
        if bid_excluded
        else QTY * ((1.0 - float(final_bid)) if bid_won else -float(final_bid))
    )
    details.append(
        {
            **row,
            "original_venue": venue,
            "start_price": start_price,
            "original_bid": bid,
            "original_ask": original_ask,
            "opposite_ask": opposite_ask,
            "commission": commission,
            "commission_ok": commission_ok,
            "action": action,
            "final": final,
            "final_ask": final_ask,
            "final_bid": final_bid,
            "baseline": baseline,
            "strategy": reversed_result,
            "bid_minus_price": bid_minus_price,
            "bid_minus_fill_ts": bid_minus_fill_ts,
            "bid_minus_excluded": bid_minus_excluded,
            "bid_minus_won": bid_minus_won,
            "bid_minus_gross": bid_minus_gross,
            "bid_fill_ts": bid_fill_ts,
            "bid_excluded": bid_excluded,
            "bid_won": bid_won,
            "bid_gross": bid_gross,
        },
    )

workbook = load_workbook(SOURCE_BOOK)
for sheet_name in (
    "全量动态反转_明细",
    "全量动态反转_汇总",
    "全量动态反转_1.2-1.3",
    "1.2-1.3汇总",
    "固定回放_反转对照",
    "固定回放_反转汇总",
):
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

detail = workbook.create_sheet("固定回放_反转对照")
detail.append(
    [
        "序号", "pair_id", "等级", "信号时间(北京时间)", "比分", "盘", "原腿venue",
        "原腿方向", "原方向PM_bid", "原方向PM_ask", "对手方向PM_ask", "原方向start_price",
        "PM_ask_sum", "commission合格", "应用反转后的动作", "应用反转后的方向",
        "应用反转后的PM_ask", "应用反转后的PM_bid", "胜者方向", "不反转是否计入", "不反转胜负",
        "不反转毛利润", "不反转手续费", "不反转净利润", "应用反转是否计入",
        "应用反转胜负", "应用反转毛利润", "应用反转手续费", "应用反转净利润",
        "应用反转-不反转净差", "bid是否成交并结算", "bid首次可成交时间",
        "bid胜负", "bid毛利润", "bid排除原因", "bid-0.05限价", "bid-0.05是否成交并结算",
        "bid-0.05首次可成交时间", "bid-0.05胜负", "bid-0.05毛利润",
        "bid-0.05排除原因", "不反转排除原因", "应用反转排除原因",
    ],
)
for index, row in enumerate(sorted(details, key=lambda value: (value["ts"], value["pair"])), 1):
    baseline = row["baseline"]
    strategy = row["strategy"]
    detail.append(
        [
            index,
            row["pair"],
            row.get("tier"),
            datetime.fromtimestamp(row["ts"] / 1e9, tz=timezone.utc).astimezone(LOCAL_TZ).replace(tzinfo=None),
            row.get("score"),
            row.get("period"),
            row["original_venue"],
            row["role"],
            row["original_bid"],
            row["original_ask"],
            row["opposite_ask"],
            row["start_price"],
            row["commission"],
            "是" if row["commission_ok"] else "否",
            row["action"],
            row["final"],
            row["final_ask"],
            row["final_bid"],
            row.get("winner"),
            "是" if baseline["excluded"] is None else "否",
            "赢" if baseline["won"] is True else "输" if baseline["won"] is False else None,
            baseline["gross"],
            baseline["fee"],
            baseline["net"],
            "是" if strategy["excluded"] is None else "否",
            "赢" if strategy["won"] is True else "输" if strategy["won"] is False else None,
            strategy["gross"],
            strategy["fee"],
            strategy["net"],
            strategy["net"] - baseline["net"]
            if strategy["net"] is not None and baseline["net"] is not None
            else None,
            "是" if row["bid_excluded"] is None else "否",
            datetime.fromtimestamp(row["bid_fill_ts"] / 1e9, tz=timezone.utc)
            .astimezone(LOCAL_TZ)
            .replace(tzinfo=None)
            if row["bid_fill_ts"] is not None
            else None,
            "赢" if row["bid_won"] is True else "输" if row["bid_won"] is False else None,
            row["bid_gross"],
            row["bid_excluded"],
            row["bid_minus_price"],
            "是" if row["bid_minus_excluded"] is None else "否",
            datetime.fromtimestamp(row["bid_minus_fill_ts"] / 1e9, tz=timezone.utc)
            .astimezone(LOCAL_TZ)
            .replace(tzinfo=None)
            if row["bid_minus_fill_ts"] is not None
            else None,
            "赢" if row["bid_minus_won"] is True else "输" if row["bid_minus_won"] is False else None,
            row["bid_minus_gross"],
            row["bid_minus_excluded"],
            baseline["excluded"],
            strategy["excluded"],
        ],
    )
style(detail)

summary = workbook.create_sheet("固定回放_反转汇总")
summary.append(["分组", "口径", "候选", "计入", "胜", "负", "毛利润", "手续费", "净利润"])


def append_summary(group_name: str, label: str, group: list[dict], result_key: str) -> None:
    included = [row for row in group if row[result_key]["excluded"] is None]
    summary.append(
        [
            group_name,
            label,
            len(group),
            len(included),
            sum(row[result_key]["won"] is True for row in included),
            sum(row[result_key]["won"] is False for row in included),
            sum(row[result_key]["gross"] for row in included),
            sum(row[result_key]["fee"] for row in included),
            sum(row[result_key]["net"] for row in included),
        ],
    )


groups = (
    ("全部", details),
    ("低级别", [row for row in details if row.get("tier") == "低级别"]),
    ("非低级别", [row for row in details if row.get("tier") == "非低级别"]),
)
for group_name, group in groups:
    append_summary(group_name, "不应用反转", group, "baseline")
    append_summary(group_name, "应用convert/attitude/deviate 1.2-1.3", group, "strategy")

summary.append([])
summary.append(
    [
        "分组", "动作", "候选", "不反转计入", "不反转胜", "不反转负", "不反转净利润",
        "应用反转计入", "应用反转胜", "应用反转负", "应用反转净利润", "应用-不反转净差",
    ],
)
for group_name, group in groups:
    for action in sorted({row["action"] for row in group}):
        action_rows = [row for row in group if row["action"] == action]
        baseline_included = [
            row for row in action_rows if row["baseline"]["excluded"] is None
        ]
        strategy_included = [
            row for row in action_rows if row["strategy"]["excluded"] is None
        ]
        baseline_net = sum(row["baseline"]["net"] for row in baseline_included)
        strategy_net = sum(row["strategy"]["net"] for row in strategy_included)
        summary.append(
            [
                group_name,
                action,
                len(action_rows),
                len(baseline_included),
                sum(row["baseline"]["won"] is True for row in baseline_included),
                sum(row["baseline"]["won"] is False for row in baseline_included),
                baseline_net,
                len(strategy_included),
                sum(row["strategy"]["won"] is True for row in strategy_included),
                sum(row["strategy"]["won"] is False for row in strategy_included),
                strategy_net,
                strategy_net - baseline_net,
            ],
        )

summary.append([])
summary.append(["口径", "固定draw_win真实回放的256条信号，不重新选择触发时间；每场最多一条"])
summary.append(["动态反转", "convert优先；OE腿再判attitude；deviate仅1.2≤bid/start≤1.3"])
summary.append(["价格门控", "两种口径分别按各自最终方向的同帧PM ask≤0.8计入"])
summary.append(["手续费", "qty(10) × price × (1-price) × 5%"])
summary.append(["start_price", f"256条中{sum(row['start_price'] is not None for row in details)}条原方向可用"])
summary.append([])
reverse_actions = {"convert反转", "attitude反转", "deviate_convert反转（1.2-1.3）"}
summary.append(["反转动作", "bid候选", "确认成交", "有赛果", "胜", "负", "maker毛利润"])
for action in ("全部反转", "convert反转", "attitude反转", "deviate_convert反转（1.2-1.3）"):
    action_rows = [
        row
        for row in details
        if row["action"] in reverse_actions
        and (action == "全部反转" or row["action"] == action)
        and row["final_bid"] is not None
        and float(row["final_bid"]) <= MAX_PRICE
    ]
    filled = [row for row in action_rows if row["bid_fill_ts"] is not None]
    settled_fills = [row for row in filled if row.get("winner") in ("yes", "no")]
    summary.append(
        [
            action,
            len(action_rows),
            len(filled),
            len(settled_fills),
            sum(row["bid_won"] is True for row in settled_fills),
            sum(row["bid_won"] is False for row in settled_fills),
            sum(row["bid_gross"] for row in settled_fills),
        ],
    )
summary.append([])
summary.append(["反转动作", "bid-0.05候选", "确认成交", "有赛果", "胜", "负", "maker毛利润"])
for action in ("全部反转", "convert反转", "attitude反转", "deviate_convert反转（1.2-1.3）"):
    action_rows = [
        row
        for row in details
        if row["action"] in reverse_actions
        and (action == "全部反转" or row["action"] == action)
        and row["bid_minus_price"] is not None
        and row["bid_minus_price"] <= MAX_PRICE
    ]
    filled = [row for row in action_rows if row["bid_minus_fill_ts"] is not None]
    settled_fills = [row for row in filled if row.get("winner") in ("yes", "no")]
    summary.append(
        [
            action,
            len(action_rows),
            len(filled),
            len(settled_fills),
            sum(row["bid_minus_won"] is True for row in settled_fills),
            sum(row["bid_minus_won"] is False for row in settled_fills),
            sum(row["bid_minus_gross"] for row in settled_fills),
        ],
    )
style(summary)
section_fill = PatternFill("solid", fgColor="5B9BD5")
for row in summary.iter_rows(min_row=2):
    if row[0].value == "分组" and len(row) > 1 and row[1].value == "动作":
        for cell in row:
            cell.fill = section_fill
            cell.font = Font(color="FFFFFF", bold=True)

workbook.save(OUTPUT_BOOK)


def totals(key: str) -> tuple[int, int, int, float]:
    included = [row for row in details if row[key]["excluded"] is None]
    return (
        len(included),
        sum(row[key]["won"] is True for row in included),
        sum(row[key]["won"] is False for row in included),
        sum(row[key]["net"] for row in included),
    )


print(f"output={OUTPUT_BOOK}")
print(f"rows={len(details)} snapshots={len(snapshots)} starts={sum(row['start_price'] is not None for row in details)}")
print(f"source_venues={dict((venue, sum(row['original_venue'] == venue for row in details)) for venue in ('PM', 'OE'))}")
print(f"baseline={totals('baseline')}")
print(f"strategy={totals('strategy')}")
