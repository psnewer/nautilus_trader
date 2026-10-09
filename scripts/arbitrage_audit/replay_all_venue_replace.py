"""对连续触发工作簿中的全部标准化信号回放 VenueReplace 动态反转。"""

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path("/Users/miller/Desktop/PM/连续触发策略_全日志明细_20260929-20261009_simulation.xlsx")
OUTPUT = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261009_全量动态反转.xlsx")
MAX_PRICE = 0.8
FEE_RATE = 0.05

# 原工作簿缺失的已完赛结果，以及 2026-10-08/09 simulation 新增比赛。
KNOWN_WINNERS = {
    "Tennis|Clement Chidekh|Ugo Blanchet": ("yes", "外部赛果：Chidekh 2-1"),
    "Tennis|Bryce Nakashima|Keegan Smith": ("no", "外部赛果：Smith 2-1"),
    "Tennis|Polina Iatcenko|Dalma Galfi": ("yes", "外部赛果：Iatcenko 2-0"),
    "Tennis|Carlo Alberto Caniato|Francesco Passaro": ("no", "日志终局"),
    "Tennis|Daniel Rincon|Pedro Martinez": ("no", "日志终局"),
    "Tennis|Gabriele Piraino|Enrico Dalla Valle": ("no", "日志终局"),
    "Tennis|Gustavo Heide|Eduardo Ribeiro": ("no", "日志终局"),
    "Tennis|Joaquin Aguilar|Lautaro Midon": ("no", "日志终局"),
    "Tennis|Juan Carlos Prado|Pedro Boscardin Dias": ("yes", "日志终局"),
    "Tennis|Max Hans Rehberg|Hynek Barton": ("no", "日志终局"),
    "Tennis|Miguel Tobon|Tomas Barrios": ("yes", "外部赛果：Tobon 2-1"),
    "Tennis|Pablo Llamas Ruiz|Petr Nesterov": ("yes", "日志终局"),
    "Tennis|Vilius Gaubas|Duje Ajdukovic": ("yes", "日志终局"),
}


def opposite(outcome: str) -> str:
    return "no" if outcome == "yes" else "yes"


def style(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for index, column in enumerate(ws.columns, start=1):
        width = max((len(str(cell.value or "")) for cell in column), default=8) + 2
        ws.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 44)


def winner_map(workbook):
    result = {}
    sources = (
        ("逐场明细", "pair_id", "胜者方向"),
        ("连续前置重算_draw全量", "pair_id", "胜者方向"),
        ("draw_win真实回放", "pair_id", "胜者方向"),
    )
    for sheet_name, pair_column, winner_column in sources:
        ws = workbook[sheet_name]
        columns = {cell.value: index for index, cell in enumerate(ws[1])}
        for row in ws.iter_rows(min_row=2, values_only=True):
            pair = row[columns[pair_column]]
            winner = row[columns[winner_column]]
            if isinstance(pair, str) and winner in ("yes", "no"):
                result[pair] = (winner, f"工作簿:{sheet_name}")
    result.update(KNOWN_WINNERS)
    return result


def old_rows(workbook):
    ws = workbook["连续前置重算_draw全量"]
    columns = {cell.value: index for index, cell in enumerate(ws[1])}
    seen = set()
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        pair = row[columns["pair_id"]]
        if not isinstance(pair, str) or not pair.startswith("Tennis|") or pair in seen:
            continue
        seen.add(pair)
        rows.append({
            "source": "历史draw|win全量",
            "pair": pair,
            "tier": row[columns["等级分组"]],
            "time": row[columns["确认时间(北京时间)"]],
            "score": row[columns["确认比分"]],
            "period": row[columns["盘数"]],
            "venue": row[columns["触发源venue"]],
            "original": row[columns["原单方向"]],
            "start_yes": row[columns["start_price_YES"]],
            "start_no": row[columns["start_price_NO"]],
            "yes_bid": row[columns["PM_YES_bid"]],
            "yes_ask": row[columns["PM_YES_ask"]],
            "no_bid": row[columns["PM_NO_bid"]],
            "no_ask": row[columns["PM_NO_ask"]],
            "commission": row[columns["commission_ask_sum"]],
        })
    return rows, seen


def simulation_rows(workbook, existing):
    ws = workbook["sim原始_1008-09"]
    columns = {cell.value: index for index, cell in enumerate(ws[1])}
    seen = set()
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        pair = row[columns["pair_id"]]
        if pair in existing or pair in seen:
            continue
        seen.add(pair)
        venue = {"POLYMARKET": "PM", "ORBITEXCH": "OE"}.get(
            row[columns["原腿venue(推断)"]],
        )
        # 同次 TrendGate 日志先于 scheduled 行落盘；结合正 rebate 可确认保留腿为 OE NO。
        if row[columns["日志行"]] == 72981:
            venue = "OE"
        rows.append({
            "source": "2026-10-08/09 simulation",
            "pair": pair,
            "tier": None,
            "time": row[columns["simulation时间(北京时间)"]],
            "score": row[columns["触发比分"]],
            "period": row[columns["盘"]],
            "venue": venue,
            "original": row[columns["连续触发原腿方向"]],
            "start_yes": row[columns["start_price_YES"]],
            "start_no": row[columns["start_price_NO"]],
            "yes_bid": row[columns["PM_YES_bid"]],
            "yes_ask": row[columns["PM_YES_ask"]],
            "no_bid": row[columns["PM_NO_bid"]],
            "no_ask": row[columns["PM_NO_ask"]],
            "commission": row[columns["PM_ask_sum"]],
        })
    return rows


def replay(item, winners, deviate_upper: float | None = None):
    original = item["original"]
    target = opposite(original)
    bid = item[f"{original}_bid"]
    start = item[f"start_{original}"]
    original_ask = item[f"{original}_ask"]
    target_ask = item[f"{target}_ask"]
    commission = item["commission"]
    commission_ok = commission is not None and 0.98 <= commission <= 1.02
    action = "默认同方向替换"
    final = original
    price = original_ask
    if item["venue"] == "PM":
        action = "convert反转"
        final = target
        price = target_ask
    elif item["venue"] == "OE" and commission_ok and bid is not None and start is not None:
        if bid <= start:
            action = "attitude反转"
            final = target
            price = target_ask
        elif bid >= 1.2 * start and (deviate_upper is None or bid <= deviate_upper * start):
            action = "deviate_convert反转"
            final = target
            price = target_ask
        elif deviate_upper is not None and bid > deviate_upper * start:
            action = f"默认同方向替换（deviate>{deviate_upper:g}）"
    if item["venue"] == "OE" and start is None:
        action = "默认同方向替换（缺start_price）"
    elif item["venue"] == "OE" and not commission_ok:
        action = "默认同方向替换（commission不合格）"

    winner, winner_basis = winners.get(item["pair"], (None, None))
    excluded = None
    if item["venue"] not in ("PM", "OE"):
        excluded = "原腿venue无法确认"
    elif price is None:
        excluded = "最终PM ask缺失"
    elif price > MAX_PRICE:
        excluded = "最终PM ask>0.8"
    elif winner not in ("yes", "no"):
        excluded = "赛果缺失"
    won = None if excluded else final == winner
    gross = None if excluded else 10 * ((1 - price) if won else -price)
    fee = None if excluded else 10 * price * (1 - price) * FEE_RATE
    net = None if excluded else gross - fee
    return {
        **item,
        "bid": bid,
        "start": start,
        "commission_ok": commission_ok,
        "action": action,
        "final": final,
        "price": price,
        "winner": winner,
        "winner_basis": winner_basis,
        "excluded": excluded,
        "won": won,
        "gross": gross,
        "fee": fee,
        "net": net,
    }


workbook = load_workbook(SOURCE)
winners = winner_map(workbook)
rows, old_pairs = old_rows(workbook)
rows.extend(simulation_rows(workbook, old_pairs))
results = [replay(row, winners) for row in rows]
bounded_results = [replay(row, winners, deviate_upper=1.3) for row in rows]

for sheet_name in (
    "全量动态反转_明细",
    "全量动态反转_汇总",
    "全量动态反转_1.2-1.3",
    "1.2-1.3汇总",
):
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

detail = workbook.create_sheet("全量动态反转_明细")
headers = [
    "序号", "数据源", "时间(北京时间)", "pair_id", "等级", "比分", "盘", "原腿venue",
    "原腿方向", "原方向PM_bid", "原方向start_price", "PM_ask_sum", "commission合格",
    "动作", "最终方向", "最终价格(PM ask)", "价格门控", "胜者方向", "胜负", "毛利润",
    "手续费", "净利润", "赛果依据", "排除原因",
]
detail.append(headers)
for index, row in enumerate(results, start=1):
    detail.append([
        index, row["source"], row["time"], row["pair"], row["tier"], row["score"], row["period"],
        row["venue"], row["original"], row["bid"], row["start"], row["commission"],
        "是" if row["commission_ok"] else "否", row["action"], row["final"], row["price"],
        "通过" if row["excluded"] is None else "排除", row["winner"],
        "赢" if row["won"] is True else "输" if row["won"] is False else None,
        row["gross"], row["fee"], row["net"], row["winner_basis"], row["excluded"],
    ])
style(detail)

summary = workbook.create_sheet("全量动态反转_汇总")
summary.append(["动作", "候选场次", "计入场次", "胜", "负", "毛利润", "手续费", "净利润"])
for action in ("convert反转", "attitude反转", "deviate_convert反转", "默认同方向替换", "默认其它"):
    if action == "默认同方向替换":
        group = [row for row in results if row["action"] == action]
    elif action == "默认其它":
        group = [row for row in results if row["action"].startswith("默认") and row["action"] != "默认同方向替换"]
    else:
        group = [row for row in results if row["action"] == action]
    settled = [row for row in group if row["excluded"] is None]
    summary.append([
        action, len(group), len(settled), sum(row["won"] is True for row in settled),
        sum(row["won"] is False for row in settled), sum(row["gross"] for row in settled),
        sum(row["fee"] for row in settled), sum(row["net"] for row in settled),
    ])
settled = [row for row in results if row["excluded"] is None]
summary.append([
    "合计", len(results), len(settled), sum(row["won"] is True for row in settled),
    sum(row["won"] is False for row in settled), sum(row["gross"] for row in settled),
    sum(row["fee"] for row in settled), sum(row["net"] for row in settled),
])
summary.append([])
summary.append(["排除原因", "场次"])
for reason, count in Counter(row["excluded"] for row in results if row["excluded"]).items():
    summary.append([reason, count])
summary.append([])
summary.append(["口径", "说明"])
summary.append(["全量范围", f"历史draw|win标准化信号{len(old_pairs)}场 + simulation新增{len(results)-len(old_pairs)}场"])
summary.append(["VenueReplace", "convert=true, attitude=true, deviate_convert=true；不启用tier转换/tier_ignore/set_exempt"])
summary.append(["价格门控", "VenueReplace完成后，最终PM ask>0.8排除；每个标准化pair信号只计一次"])
summary.append(["手续费", "qty(10) × price × (1-price) × 5%"])
style(summary)

bounded = workbook.create_sheet("全量动态反转_1.2-1.3")
bounded.append(headers)
for index, row in enumerate(bounded_results, start=1):
    bounded.append([
        index, row["source"], row["time"], row["pair"], row["tier"], row["score"], row["period"],
        row["venue"], row["original"], row["bid"], row["start"], row["commission"],
        "是" if row["commission_ok"] else "否", row["action"], row["final"], row["price"],
        "通过" if row["excluded"] is None else "排除", row["winner"],
        "赢" if row["won"] is True else "输" if row["won"] is False else None,
        row["gross"], row["fee"], row["net"], row["winner_basis"], row["excluded"],
    ])
bounded_settled = [row for row in bounded_results if row["excluded"] is None]
bounded.append([])
bounded.append([
    "汇总", None, None, None, None, None, None, None, None, None, None, None, None,
    "deviate_convert仅在1.2≤bid/start≤1.3时反转", None, None, None, None,
    f"{sum(row['won'] is True for row in bounded_settled)}胜{sum(row['won'] is False for row in bounded_settled)}负",
    sum(row["gross"] for row in bounded_settled),
    sum(row["fee"] for row in bounded_settled),
    sum(row["net"] for row in bounded_settled),
])
for action in sorted({row["action"] for row in bounded_results}):
    group = [row for row in bounded_results if row["action"] == action and row["excluded"] is None]
    bounded.append([
        f"动作汇总:{action}", None, None, None, None, None, None, None, None, None, None, None,
        None, action, None, None, len(group), None,
        f"{sum(row['won'] is True for row in group)}胜{sum(row['won'] is False for row in group)}负",
        sum(row["gross"] for row in group), sum(row["fee"] for row in group),
        sum(row["net"] for row in group),
    ])
style(bounded)

bounded_summary = workbook.create_sheet("1.2-1.3汇总")
summary_headers = ["分组", "动作", "候选场次", "计入场次", "胜", "负", "毛利润", "手续费", "净利润"]
bounded_summary.append(summary_headers)


def append_summary_row(ws, group_name, action_name, group):
    included = [row for row in group if row["excluded"] is None]
    ws.append([
        group_name,
        action_name,
        len(group),
        len(included),
        sum(row["won"] is True for row in included),
        sum(row["won"] is False for row in included),
        sum(row["gross"] for row in included),
        sum(row["fee"] for row in included),
        sum(row["net"] for row in included),
    ])


tier_groups = (
    ("全部", bounded_results),
    ("低级别", [row for row in bounded_results if row["tier"] == "低级别"]),
    ("非低级别", [row for row in bounded_results if row["tier"] == "非低级别"]),
    ("等级未分类", [row for row in bounded_results if row["tier"] not in ("低级别", "非低级别")]),
)
for tier_name, tier_rows in tier_groups:
    append_summary_row(bounded_summary, tier_name, "全部动作", tier_rows)

bounded_summary.append([])
for tier_name, tier_rows in tier_groups:
    append_summary_row(
        bounded_summary,
        tier_name,
        "deviate_convert反转（1.2≤bid/start≤1.3）",
        [row for row in tier_rows if row["action"] == "deviate_convert反转"],
    )

bounded_summary.append([])
for tier_name, tier_rows in tier_groups:
    for action in sorted({row["action"] for row in tier_rows}):
        append_summary_row(
            bounded_summary,
            tier_name,
            action,
            [row for row in tier_rows if row["action"] == action],
        )

bounded_summary.append([])
bounded_summary.append(["口径", "deviate_convert仅在1.2≤bid/start≤1.3时反转；bid/start>1.3保持原方向"])
bounded_summary.append(["等级未分类", "simulation新增10场的源表没有competition/等级字段，未擅自归入低级别或非低级别"])
bounded_summary.append(["价格门控", "VenueReplace完成后，最终PM ask>0.8排除"])
bounded_summary.append(["手续费", "qty(10) × price × (1-price) × 5%"])
style(bounded_summary)

workbook.save(OUTPUT)
print(f"output={OUTPUT}")
print(f"candidates={len(results)} included={len(settled)}")
print(f"wins={sum(row['won'] is True for row in settled)} losses={sum(row['won'] is False for row in settled)}")
print(f"gross={sum(row['gross'] for row in settled):.6f} fee={sum(row['fee'] for row in settled):.6f} net={sum(row['net'] for row in settled):.6f}")
print(Counter(row["action"] for row in settled))
print(Counter(row["excluded"] for row in results if row["excluded"]))
print(
    "bounded_1.2_1.3",
    f"included={len(bounded_settled)}",
    f"wins={sum(row['won'] is True for row in bounded_settled)}",
    f"losses={sum(row['won'] is False for row in bounded_settled)}",
    f"gross={sum(row['gross'] for row in bounded_settled):.6f}",
    f"fee={sum(row['fee'] for row in bounded_settled):.6f}",
    f"net={sum(row['net'] for row in bounded_settled):.6f}",
)
print(Counter(row["action"] for row in bounded_settled))
