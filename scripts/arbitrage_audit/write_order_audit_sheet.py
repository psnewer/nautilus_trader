"""把 nohup.out 的 OrderInitialized 审计结果写入既有 Excel。"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DETAIL_SHEET = "下单明细"
SUMMARY_SHEET = "下单汇总"
TIME_COLUMNS = {
    "下单时间(北京时间)",
    "首次成交时间(北京时间)",
    "最接近限价OBD时间(北京时间)",
}
CLOSEST_OBD_COLUMNS = {
    "最接近限价OBD时间(北京时间)",
    "最接近OBD距下单(秒)",
    "最接近OBD_方向PM_bid",
    "最接近OBD_方向PM_ask",
    "最接近OBD_ask减限价",
    "最接近OBD_PM_YES_ask",
    "最接近OBD_PM_NO_ask",
    "最接近OBD_OE_YES_ask",
    "最接近OBD_OE_NO_ask",
    "最接近OBD_SE_YES_ask",
    "最接近OBD_SE_NO_ask",
}

# 这 5 行是旧脚本按首次成交时间追加到“成交明细”的错误记录。
# 新口径统一放入“下单明细”，时间锚点为 OrderInitialized。
INCORRECT_LEGACY_ROWS = {
    ("Anna Bondar", "Alina Charaeva", 0.46),
    ("Yeon-Woo Ku", "Gabriela Ruse", 0.22),
    ("Sho Shimabukuro", "Hugo Gaston", 0.44),
    ("Tallon Griekspoor", "Denis Shapovalov", 0.44),
    ("Jaime Faria", "Terence Atmane", 0.56),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("audit_json", type=Path)
    parser.add_argument("output", type=Path)
    return parser.parse_args()


def as_datetime(value: object) -> object:
    if not isinstance(value, str) or not value:
        return value
    try:
        return datetime.strptime(value, "%Y-%m-%d %H:%M:%S.%f")
    except ValueError:
        return value


def remove_incorrect_legacy_rows(workbook) -> int:
    sheet = workbook["成交明细"]
    deleted = 0
    for row_number in range(sheet.max_row, 1, -1):
        player1 = sheet.cell(row_number, 3).value
        player2 = sheet.cell(row_number, 4).value
        price = sheet.cell(row_number, 6).value
        try:
            key = (player1, player2, round(float(price), 8))
        except (TypeError, ValueError):
            continue
        if key in INCORRECT_LEGACY_ROWS:
            sheet.delete_rows(row_number)
            deleted += 1
    return deleted


def replace_sheet(workbook, title: str, index: int):
    if title in workbook.sheetnames:
        del workbook[title]
    return workbook.create_sheet(title, index)


def write_detail(workbook, rows: list[dict]) -> None:
    sheet = replace_sheet(workbook, DETAIL_SHEET, 1)
    headers = list(rows[0])
    sheet.append(headers)

    for source_row in rows:
        row = dict(source_row)
        # 风控拒绝发生在订单送达 venue 前，不属于盘口上的未成交挂单。
        if row.get("状态") == "未成交-风控拒绝":
            for column in CLOSEST_OBD_COLUMNS:
                row[column] = None
        values = []
        for header in headers:
            value = row.get(header)
            values.append(as_datetime(value) if header in TIME_COLUMNS else value)
        sheet.append(values)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    status_column = headers.index("状态") + 1
    status_fills = {
        "已成交": PatternFill("solid", fgColor="C6EFCE"),
        "部分成交": PatternFill("solid", fgColor="FFEB9C"),
        "未成交-已撤单": PatternFill("solid", fgColor="F4CCCC"),
        "未成交-风控拒绝": PatternFill("solid", fgColor="E7E6E6"),
        "未成交-日志结束时无终态": PatternFill("solid", fgColor="FCE5CD"),
    }
    for row_number in range(2, sheet.max_row + 1):
        status_cell = sheet.cell(row_number, status_column)
        status_cell.fill = status_fills.get(status_cell.value, PatternFill())

    for column_number, header in enumerate(headers, 1):
        letter = get_column_letter(column_number)
        if header in TIME_COLUMNS:
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row_number, column_number).number_format = "yyyy-mm-dd hh:mm:ss.000"
        elif header in {"限价", "成交均价"} or "bid" in header or "ask" in header:
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row_number, column_number).number_format = "0.0000"
        maximum = max(len(str(sheet.cell(row, column_number).value or "")) for row in range(1, sheet.max_row + 1))
        sheet.column_dimensions[letter].width = min(max(maximum + 2, 11), 34)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30


def write_summary(workbook, rows: list[dict], source: str, deleted: int) -> None:
    sheet = replace_sheet(workbook, SUMMARY_SHEET, 0)
    counts = Counter(row["状态"] for row in rows)
    source_order = [
        "已成交",
        "部分成交",
        "未成交-已撤单",
        "未成交-日志结束时无终态",
        "未成交-风控拒绝",
    ]
    data = [
        ("统计口径", "nohup.out 中全部 OrderInitialized 事件"),
        ("时间锚点", "OrderInitialized（北京时间），不是 OrderAccepted 或成交时间"),
        ("来源", source),
        ("订单总数", len(rows)),
        ("未取到 start_price", sum(row.get("start_yes") is None or row.get("start_no") is None for row in rows)),
        ("从旧成交明细移除的错误追加行", deleted),
        ("说明", "风控拒绝订单未送达 venue，因此不填写最接近成交 OBD"),
        (None, None),
        ("状态", "数量"),
        *[(status, counts.get(status, 0)) for status in source_order],
    ]
    for row in data:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for cell in sheet[9]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 72
    sheet.freeze_panes = "A2"


def main() -> None:
    args = parse_args()
    with args.audit_json.open(encoding="utf-8") as stream:
        audit = json.load(stream)
    rows = audit["rows"]
    if not rows:
        raise ValueError("审计 JSON 没有订单")

    workbook = load_workbook(args.workbook)
    deleted = remove_incorrect_legacy_rows(workbook)
    if deleted != len(INCORRECT_LEGACY_ROWS):
        raise ValueError(f"预期移除 5 行旧错误记录，实际移除 {deleted} 行")
    write_detail(workbook, rows)
    write_summary(workbook, rows, audit.get("source", ""), deleted)
    workbook.save(args.output)
    print(json.dumps({"orders": len(rows), "deleted_legacy_rows": deleted, "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
