from __future__ import annotations

import ast
import re
import subprocess
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PATHS = [
    Path("/Users/miller/Downloads/nohup.out"),
    Path("/private/tmp/remote_log2_20261007.out"),
    Path("/private/tmp/remote_log_20261007_latest.out"),
    Path("/private/tmp/remote_log1_20261007.out"),
    Path("/private/tmp/remote_nohup_20261007_1601.out"),
    Path("/private/tmp/remote_log3_20261008.out"),
    Path("/private/tmp/remote_nohup_20261008_latest.out"),
]
SOURCE_BOOK = Path("/Users/miller/Desktop/PM/连续触发策略_全日志明细_20260929-20261007.xlsx")
OUTPUT_BOOK = Path("/private/tmp/连续触发策略_全日志明细_20260929-20261007_连续前置重算.xlsx")
START = datetime(2026, 9, 29, tzinfo=timezone(timedelta(hours=8))).astimezone(timezone.utc)
LOCAL_TZ = timezone(timedelta(hours=8))
ANSI = re.compile(r"\x1b\[[0-9;]*m")
ISO = re.compile(r"(2026-10-\d\dT\d\d:\d\d:\d\d\.\d+Z)")
LOCAL = re.compile(r"^(2026-10-\d\d \d\d:\d\d:\d\d,\d+)")
SCORE_PART = re.compile(r"^(\d+)\s*-\s*(\d+)(?:\((\d+)\s*-\s*(\d+)\))?$")
FILTER = (
    "Strategy evaluate scheduled: pair_id=|Sports score updated: game=|"
    "Start price captured: pair=|Strategy evaluate result: pair_id=.*arb_hit=True.*VenueReplaceAction|"
    "TrendGate: pair=|tier_ignore hit"
)
KNOWN_WINNERS = {
    "Tennis|Yuta Shimizu|Bernard Tomic": "no",
    "Tennis|Jaume Munar|Kyrian Jacquet": "yes",
    "Tennis|Cagla Buyukakcay|Linda Klimovicova": "no",
    "Tennis|Philip Henning|Maks Kasnikowski": "yes",
    "Tennis|Francesco Maestrelli|Oliver Tarvet": "yes",
    "Tennis|Coco Gauff|Elise Mertens": "no",
}


def timestamp(line: str) -> datetime | None:
    match = ISO.search(line)
    if match:
        return datetime.fromisoformat(match.group(1).replace("Z", "+00:00"))
    match = LOCAL.search(line)
    if match:
        return datetime.strptime(match.group(1), "%Y-%m-%d %H:%M:%S,%f").replace(
            tzinfo=LOCAL_TZ,
        ).astimezone(timezone.utc)
    return None


def filtered_lines(path: Path, pattern: str):
    process = subprocess.Popen(
        ["rg", "--no-filename", pattern, str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        errors="replace",
    )
    assert process.stdout is not None
    for raw in process.stdout:
        yield ANSI.sub("", raw.rstrip())
    stderr = process.stderr.read() if process.stderr is not None else ""
    code = process.wait()
    if code not in (0, 1):
        raise RuntimeError(f"rg failed for {path}: {stderr}")


def pair_from_action(line: str) -> str | None:
    marker = "pair="
    if marker not in line:
        return None
    value = line.split(marker, 1)[1]
    for delimiter in (" drop leg=", " tier_ignore"):
        if delimiter in value:
            return value.split(delimiter, 1)[0]
    return None


def prior(values: list[dict], ts: datetime) -> dict | None:
    found = None
    for value in values:
        if value["ts"] <= ts:
            found = value
        else:
            break
    return found


def parse_books(raw: str) -> list[dict]:
    try:
        value = ast.literal_eval(raw)
    except Exception:
        return []
    return value if isinstance(value, list) else []


def pm_quotes(snapshot: dict | None) -> dict[str, dict]:
    return {
        str(book.get("outcome")): book
        for book in (snapshot or {}).get("books", [])
        if book.get("venue") == "POLYMARKET"
    }


def score_parts(score: str) -> list[tuple[int, int, int | None, int | None]] | None:
    result = []
    for token in [part.strip() for part in str(score or "").split(",") if part.strip()]:
        match = SCORE_PART.match(token)
        if match is None:
            return None
        left, right, tie_left, tie_right = match.groups()
        result.append(
            (int(left), int(right), int(tie_left) if tie_left else None, int(tie_right) if tie_right else None),
        )
    return result or None


def unit_complete(unit: tuple[int, int, int | None, int | None]) -> bool:
    left, right, tie_left, tie_right = unit
    if max(left, right) < 6:
        return False
    if max(left, right) == 6:
        return abs(left - right) >= 2
    if left == right == 6:
        return tie_left is not None and tie_right is not None and tie_left != tie_right
    return abs(left - right) >= 2 or {left, right} == {6, 7}


def match_comparison(score: str) -> int | None:
    parsed = score_parts(score)
    if parsed is None:
        return None
    if parsed[-1][0] == parsed[-1][1] == 6:
        return None
    completed = parsed if len(parsed) > 1 and unit_complete(parsed[-1]) else parsed[:-1]
    home_sets = sum(left > right for left, right, _, _ in completed)
    away_sets = sum(left < right for left, right, _, _ in completed)
    if home_sets != away_sets:
        return 1 if home_sets > away_sets else -1
    if len(completed) == len(parsed):
        return 0
    left, right, _, _ = parsed[-1]
    return (left > right) - (left < right)


def passes_score_filters(direction: str, score: str) -> bool:
    parsed = score_parts(score)
    comparison = match_comparison(score)
    if parsed is None or comparison is None:
        return False
    oriented_match = comparison if direction == "yes" else -comparison
    left, right, _, _ = parsed[-1]
    oriented_set = (left > right) - (left < right)
    if direction == "no":
        oriented_set = -oriented_set
    return oriented_match >= 0 and oriented_set > 0


def winner_from_score(score: str) -> str | None:
    yes_sets = no_sets = 0
    for token in str(score or "").split(","):
        match = re.match(r"\s*(\d+)\s*-\s*(\d+)", token)
        if match is None:
            continue
        left, right = map(int, match.groups())
        yes_sets += left > right
        no_sets += right > left
    if yes_sets >= 2 and yes_sets > no_sets:
        return "yes"
    if no_sets >= 2 and no_sets > yes_sets:
        return "no"
    return None


scores: dict[str, list[dict]] = defaultdict(list)
starts: dict[str, list[dict]] = defaultdict(list)
events: list[dict] = []
low_pairs: set[str] = set()

for path in PATHS:
    scheduled: dict[str, dict] = {}
    active: dict[str, dict] = {}
    for line in filtered_lines(path, FILTER):
        ts = timestamp(line)
        if ts is None or ts < START:
            continue
        if "Strategy evaluate scheduled: pair_id=" in line:
            match = re.search(r"pair_id=(.+?), sport=.*?, event=([^,]+), order_books=(\[.*\])$", line)
            if match:
                scheduled[match.group(1)] = {
                    "ts": ts,
                    "event": match.group(2),
                    "books": parse_books(match.group(3)),
                }
            continue
        if "Sports score updated: game=" in line:
            match = re.search(
                r"game=\d+ (.+?) vs (.+?) score=(.*?) period=([^ ]+) elapsed=.*? status=([^ ]+)$",
                line,
            )
            if match:
                pair = f"Tennis|{match.group(1)}|{match.group(2)}"
                scores[pair].append(
                    {
                        "ts": ts,
                        "score": match.group(3).split("->")[-1],
                        "period": match.group(4),
                        "status": match.group(5),
                    },
                )
            continue
        if "Start price captured: pair=" in line:
            match = re.search(r"pair=(.+?) game=\d+ source=([^ ]+) prices=(\{.*\})$", line)
            if match:
                try:
                    prices = ast.literal_eval(match.group(3))
                except Exception:
                    prices = {}
                starts[match.group(1)].append({"ts": ts, **prices})
            continue
        if "Strategy evaluate result: pair_id=" in line:
            match = re.search(r"pair_id=(.+?), arb_hit=True", line)
            if match:
                pair = match.group(1)
                event = {"pair": pair, "ts": ts, "snapshot": scheduled.get(pair), "trend": []}
                events.append(event)
                active[pair] = event
            continue
        pair = pair_from_action(line)
        if pair is None:
            continue
        if "tier_ignore hit" in line:
            low_pairs.add(pair)
        event = active.get(pair)
        if event is not None and 0 <= (ts - event["ts"]).total_seconds() <= 2 and "TrendGate:" in line:
            event["trend"].append(line)

for collection in (scores, starts):
    for values in collection.values():
        values.sort(key=lambda value: value["ts"])

# 延续昨天工作簿中已经人工复核过的等级分组，并吸收日志里的 tier_ignore 证据。
old_book = load_workbook(SOURCE_BOOK, read_only=True, data_only=True)
old_detail = old_book["逐场明细"]
header = {value: index for index, value in enumerate(next(old_detail.iter_rows(values_only=True)))}
for row in old_detail.iter_rows(min_row=2, values_only=True):
    if row[header["等级分组"]] == "低级别":
        low_pairs.add(row[header["pair_id"]])
old_book.close()


def source_direction(event: dict) -> str | None:
    dropped = set()
    for line in event["trend"]:
        match = re.search(r" outcome=(yes|no) ", line)
        if match:
            dropped.add(match.group(1))
    if dropped == {"yes"}:
        return "no"
    if dropped == {"no"}:
        return "yes"
    return None


def source_venue(event: dict) -> str | None:
    venues = set()
    for line in event["trend"]:
        venues.add("PM" if ".POLYMARKET" in line else "OE" if ".ORBITEXCH" in line else "?")
    if venues == {"PM"}:
        return "OE"
    if venues == {"OE"}:
        return "PM"
    return None


deduped: dict[tuple, dict] = {}
for event in events:
    direction = source_direction(event)
    state = prior(scores[event["pair"]], event["ts"])
    if direction is None or state is None or score_parts(state["score"]) is None:
        continue
    event["direction"] = direction
    event["score"] = state["score"]
    event["period"] = state["period"]
    key = (event["pair"], event["ts"], direction, state["score"])
    previous = deduped.get(key)
    if previous is None or len(event["trend"]) > len(previous["trend"]):
        deduped[key] = event

by_pair: dict[str, list[dict]] = defaultdict(list)
for event in sorted(deduped.values(), key=lambda value: value["ts"]):
    by_pair[event["pair"]].append(event)

chosen: list[tuple[str, dict, dict]] = []
for pair, pair_events in by_pair.items():
    history: list[dict] = []
    for current in pair_events:
        count = 1
        previous_distinct = None
        last_score = current["score"]
        for previous in reversed(history):
            if previous["direction"] != current["direction"]:
                break
            if previous["score"] == last_score:
                continue
            count += 1
            previous_distinct = previous
            last_score = previous["score"]
            if count >= 2:
                break
        history.append(current)
        if count >= 2 and passes_score_filters(current["direction"], current["score"]):
            chosen.append((pair, previous_distinct, current))
            break

chosen_pairs = {pair for pair, _, _ in chosen}
future_quotes: dict[str, list[dict]] = defaultdict(list)
terminal_quotes: dict[str, dict] = {}
for path in PATHS:
    for line in filtered_lines(path, "Strategy evaluate scheduled: pair_id="):
        ts = timestamp(line)
        if ts is None or ts < START:
            continue
        match = re.search(r"pair_id=(.+?), sport=.*?, event=([^,]+), order_books=(\[.*\])$", line)
        if match is None or match.group(1) not in chosen_pairs:
            continue
        pair = match.group(1)
        snapshot = {"ts": ts, "books": parse_books(match.group(3))}
        quotes = pm_quotes(snapshot)
        if quotes:
            future_quotes[pair].append({"ts": ts, "quotes": quotes})
            terminal_quotes[pair] = quotes
for pair in future_quotes:
    unique = {}
    for item in future_quotes[pair]:
        key = (
            item["ts"],
            item["quotes"].get("yes", {}).get("best_bid"),
            item["quotes"].get("yes", {}).get("best_ask"),
            item["quotes"].get("no", {}).get("best_bid"),
            item["quotes"].get("no", {}).get("best_ask"),
        )
        unique[key] = item
    future_quotes[pair] = sorted(unique.values(), key=lambda value: value["ts"])


def resolved_winner(pair: str) -> str | None:
    quotes = terminal_quotes.get(pair, {})
    yes = quotes.get("yes", {})
    no = quotes.get("no", {})
    if yes.get("best_bid") is not None and yes["best_bid"] >= 0.99:
        return "yes"
    if no.get("best_bid") is not None and no["best_bid"] >= 0.99:
        return "no"
    tail = scores[pair][-1] if scores[pair] else {}
    if tail.get("status") == "finished":
        winner = winner_from_score(tail.get("score", ""))
        if winner:
            return winner
    return KNOWN_WINNERS.get(pair)


def strategy_direction(pair: str, event: dict, quotes: dict) -> tuple[str, str, str | None]:
    source = event["direction"]
    final = source
    action = "买原方向"
    venue = source_venue(event)
    exempt = str(event.get("period") or "").upper() == "S3"
    start = prior(starts[pair], event["ts"]) or {}
    if venue == "PM" and not exempt:
        final = "no" if source == "yes" else "yes"
        action = "convert反买"
    elif venue == "OE" and not exempt:
        asks = [quotes.get(role, {}).get("best_ask") for role in ("yes", "no")]
        bid = quotes.get(source, {}).get("best_bid")
        start_price = start.get(source)
        commission_ok = all(value is not None for value in asks) and 0.98 <= sum(asks) <= 1.02
        if commission_ok and bid is not None and start_price is not None and bid <= start_price:
            final = "no" if source == "yes" else "yes"
            action = "attitude反买"
        elif commission_ok and bid is not None and start_price is not None and bid >= 1.2 * start_price:
            final = "no" if source == "yes" else "yes"
            action = "deviate_convert反买"
    return final, action, venue


def maker_result(pair: str, ts: datetime, direction: str, price: float | None, winner: str | None):
    fill = None
    if price is not None:
        for item in future_quotes[pair]:
            if item["ts"] < ts:
                continue
            ask = item["quotes"].get(direction, {}).get("best_ask")
            if ask is not None and ask <= price:
                fill = item["ts"]
                break
    gross = None if fill is None or winner is None else 10 * ((1 - price) if winner == direction else -price)
    return fill, gross, gross


def taker_result(direction: str, quotes: dict, winner: str | None):
    yes_ask = quotes.get("yes", {}).get("best_ask")
    no_ask = quotes.get("no", {}).get("best_ask")
    commission = yes_ask + no_ask if yes_ask is not None and no_ask is not None else None
    qualified = commission is not None and 0.98 <= commission <= 1.02
    price = quotes.get(direction, {}).get("best_ask") if qualified else None
    gross = None if price is None or winner is None else 10 * ((1 - price) if winner == direction else -price)
    fee = None if gross is None else 10 * price * (1 - price) * 0.05
    net = None if gross is None else gross - fee
    return commission, qualified, price, gross, fee, net


rows = []
for pair, previous, current in sorted(chosen, key=lambda value: value[2]["ts"]):
    quotes = pm_quotes(current.get("snapshot"))
    source = current["direction"]
    final, action, venue = strategy_direction(pair, current, quotes)
    winner = resolved_winner(pair)
    start = prior(starts[pair], current["ts"]) or {}
    original_bid = quotes.get(source, {}).get("best_bid")
    strategy_bid = quotes.get(final, {}).get("best_bid")
    original_bid_fill, original_bid_gross, original_bid_net = maker_result(
        pair, current["ts"], source, original_bid, winner,
    )
    strategy_bid_fill, strategy_bid_gross, strategy_bid_net = maker_result(
        pair, current["ts"], final, strategy_bid, winner,
    )
    commission, commission_ok, original_ask, original_ask_gross, original_ask_fee, original_ask_net = taker_result(
        source, quotes, winner,
    )
    _, _, strategy_ask, strategy_ask_gross, strategy_ask_fee, strategy_ask_net = taker_result(
        final, quotes, winner,
    )
    rows.append(
        {
            "pair": pair,
            "tier": "低级别" if pair in low_pairs else "非低级别",
            "previous_ts": previous["ts"] if previous else None,
            "previous_score": previous["score"] if previous else None,
            "confirm_ts": current["ts"],
            "confirm_score": current["score"],
            "period": current["period"],
            "source": source,
            "source_venue": venue,
            "action": action,
            "final": final,
            "start_yes": start.get("yes"),
            "start_no": start.get("no"),
            "yes_bid": quotes.get("yes", {}).get("best_bid"),
            "yes_ask": quotes.get("yes", {}).get("best_ask"),
            "no_bid": quotes.get("no", {}).get("best_bid"),
            "no_ask": quotes.get("no", {}).get("best_ask"),
            "commission": commission,
            "commission_ok": commission_ok,
            "winner": winner,
            "original_bid": original_bid,
            "original_bid_fill": original_bid_fill,
            "original_bid_gross": original_bid_gross,
            "original_bid_net": original_bid_net,
            "original_ask": original_ask,
            "original_ask_gross": original_ask_gross,
            "original_ask_fee": original_ask_fee,
            "original_ask_net": original_ask_net,
            "strategy_bid": strategy_bid,
            "strategy_bid_fill": strategy_bid_fill,
            "strategy_bid_gross": strategy_bid_gross,
            "strategy_bid_net": strategy_bid_net,
            "strategy_ask": strategy_ask,
            "strategy_ask_gross": strategy_ask_gross,
            "strategy_ask_fee": strategy_ask_fee,
            "strategy_ask_net": strategy_ask_net,
        },
    )

wb = load_workbook(SOURCE_BOOK)
sheet_name = "连续前置重算"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name)
headers = [
    "序号", "pair_id", "等级分组", "前一次时间(北京时间)", "前一次比分", "确认时间(北京时间)", "确认比分", "盘数",
    "原单方向", "触发源venue", "策略动作", "策略方向", "start_price_YES", "start_price_NO",
    "PM_YES_bid", "PM_YES_ask", "PM_NO_bid", "PM_NO_ask", "commission_ask_sum", "commission合格", "胜者方向",
    "原单_bid价格", "原单_bid成交时间", "原单_bid毛利润", "原单_bid净利润",
    "原单_ask价格", "原单_ask毛利润", "原单_ask手续费", "原单_ask净利润",
    "策略_bid价格", "策略_bid成交时间", "策略_bid毛利润", "策略_bid净利润",
    "策略_ask价格", "策略_ask毛利润", "策略_ask手续费", "策略_ask净利润",
]
ws.append(headers)
for index, row in enumerate(rows, 1):
    ws.append(
        [
            index, row["pair"], row["tier"],
            row["previous_ts"].astimezone(LOCAL_TZ).replace(tzinfo=None) if row["previous_ts"] else None,
            row["previous_score"], row["confirm_ts"].astimezone(LOCAL_TZ).replace(tzinfo=None), row["confirm_score"], row["period"],
            row["source"], row["source_venue"], row["action"], row["final"], row["start_yes"], row["start_no"],
            row["yes_bid"], row["yes_ask"], row["no_bid"], row["no_ask"], row["commission"], "是" if row["commission_ok"] else "否", row["winner"],
            row["original_bid"], row["original_bid_fill"].astimezone(LOCAL_TZ).replace(tzinfo=None) if row["original_bid_fill"] else None,
            row["original_bid_gross"], row["original_bid_net"], row["original_ask"], row["original_ask_gross"], row["original_ask_fee"], row["original_ask_net"],
            row["strategy_bid"], row["strategy_bid_fill"].astimezone(LOCAL_TZ).replace(tzinfo=None) if row["strategy_bid_fill"] else None,
            row["strategy_bid_gross"], row["strategy_bid_net"], row["strategy_ask"], row["strategy_ask_gross"], row["strategy_ask_fee"], row["strategy_ask_net"],
        ],
    )

summary_start = ws.max_row + 2
summary_headers = ["方向口径", "价格口径", "分组", "候选", "成交/合格", "已结算", "胜", "负", "毛利润", "手续费", "净利润"]
ws.append([])
ws.append(summary_headers)
summaries = []
for direction_name, prefix in (("原单方向", "original"), ("策略方向", "strategy")):
    for price_name, suffix in (("bid挂单", "bid"), ("ask吃单", "ask")):
        for tier in ("非低级别", "低级别", "全部"):
            subset = [row for row in rows if tier == "全部" or row["tier"] == tier]
            if suffix == "bid":
                eligible = [row for row in subset if row[f"{prefix}_bid_fill"] is not None]
                settled = [row for row in eligible if row[f"{prefix}_bid_gross"] is not None]
                fees = [0.0 for _ in settled]
            else:
                eligible = [row for row in subset if row["commission_ok"] and row[f"{prefix}_ask"] is not None]
                settled = [row for row in eligible if row[f"{prefix}_ask_gross"] is not None]
                fees = [row[f"{prefix}_ask_fee"] for row in settled]
            gross = sum(row[f"{prefix}_{suffix}_gross"] for row in settled)
            fee = sum(fees)
            summaries.append(
                [
                    direction_name, price_name, tier, len(subset), len(eligible), len(settled),
                    sum(row["winner"] == (row["source"] if prefix == "original" else row["final"]) for row in settled),
                    sum(row["winner"] != (row["source"] if prefix == "original" else row["final"]) for row in settled),
                    gross, fee, gross - fee,
                ],
            )
for summary in summaries:
    ws.append(summary)

fill = PatternFill("solid", fgColor="D9EAF7")
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = fill
for cell in ws[summary_start]:
    cell.font = Font(bold=True)
    cell.fill = fill
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{1 + len(rows)}"
for column in range(1, len(headers) + 1):
    max_length = max(len(str(ws.cell(row=row, column=column).value or "")) for row in range(1, min(ws.max_row, 150) + 1))
    ws.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 10), 34)
for row in ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=False)
for column in range(13, 38):
    for row_index in range(2, ws.max_row + 1):
        ws.cell(row=row_index, column=column).number_format = "0.0000"
for column in (4, 6, 23, 31):
    for row_index in range(2, 2 + len(rows)):
        ws.cell(row=row_index, column=column).number_format = "yyyy-mm-dd hh:mm:ss.000"

wb.save(OUTPUT_BOOK)
print(OUTPUT_BOOK)
print("detail_rows", len(rows))
for summary in summaries:
    print(*summary)
