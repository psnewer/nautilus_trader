"""复现 2026-09-24 五笔成交明细追加操作的一次性脚本。"""

from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/miller/Desktop/成交订单比分_start_price_venue_stale_window_20260921-22.xlsx")
TARGET = Path("/private/tmp/成交订单比分_start_price_venue_stale_window_20260921-22.updated.xlsx")


ROWS = [
    [
        datetime(2026, 9, 24, 11, 22, 4), "WTA", "Anna Bondar", "Alina Charaeva",
        "no", 0.46, 0.46, "1-2", "S1", 0.63, 0.38, "captured",
        "ORBITEXCH", "POLYMARKET",
        0.49, 0.52, 0.54945055, 0.46296296, None, None,
        0.49, 0.53, 0.54945055, 0.46296296, None, None,
        63.810, "ORBITEXCH",
        0.52, 0.49, 0.49504950, 0.51020408, None, None,
    ],
    [
        datetime(2026, 9, 24, 13, 16, 47), "WTA", "Yeon-Woo Ku", "Gabriela Ruse",
        "yes", 0.22, 0.22, "1-1", "S1", 0.26, 0.75, "captured",
        "ORBITEXCH", "POLYMARKET",
        0.29, 0.72, 0.24390244, 0.76335878, None, None,
        0.29, 0.73, 0.24390244, 0.76335878, None, None,
        84.495, "ORBITEXCH",
        0.30, 0.73, 0.31746032, 0.72463768, None, None,
    ],
    [
        datetime(2026, 9, 24, 13, 18, 21), "ATP", "Sho Shimabukuro", "Hugo Gaston",
        "yes", 0.44, 0.44, "1-1", "S1", 0.42, 0.59, "captured",
        "ORBITEXCH", "POLYMARKET",
        0.50, 0.52, 0.45045045, 0.55555556, None, None,
        0.50, 0.51, 0.45045045, 0.55555556, None, None,
        124.973, "ORBITEXCH",
        0.44, 0.57, 0.44642857, 0.55555556, None, None,
    ],
    [
        datetime(2026, 9, 24, 15, 14, 46), "ATP", "Tallon Griekspoor", "Denis Shapovalov",
        "yes", 0.44, 0.44, "1-0", "S1", 0.43, 0.58, "captured",
        "ORBITEXCH", "POLYMARKET",
        0.50, 0.51, 0.43103448, 0.58139535, None, None,
        0.51, 0.51, 0.43103448, 0.58139535, None, None,
        50.190, "ORBITEXCH",
        0.52, 0.52, 0.50, 0.51020408, None, None,
    ],
    [
        datetime(2026, 9, 24, 15, 15, 35), "ATP", "Jaime Faria", "Terence Atmane",
        "no", 0.56, 0.56, "3-6, 6-6(9-9)", "TB2", 0.44, 0.57, "captured",
        "ORBITEXCH", "POLYMARKET",
        0.39, 0.62, 0.43103448, 0.57471264, None, None,
        0.39, 0.63, 0.43103448, 0.57471264, None, None,
        53.508, "ORBITEXCH",
        0.33, 0.69, 0.34246575, 0.67567568, None, None,
    ],
]


workbook = load_workbook(SOURCE)
sheet = workbook["成交明细"]
if sheet.max_column != 34:
    raise RuntimeError(f"成交明细列数异常: {sheet.max_column}")

existing = {
    (sheet.cell(row, 3).value, sheet.cell(row, 4).value, sheet.cell(row, 6).value)
    for row in range(2, sheet.max_row + 1)
}
template_row = sheet.max_row
added = []
for values in ROWS:
    key = (values[2], values[3], values[5])
    if key in existing:
        continue
    sheet.append(values)
    row = sheet.max_row
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(template_row, column)
        target = sheet.cell(row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
    sheet.cell(row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    sheet.cell(row, 12)._style = copy(sheet.cell(template_row, 12)._style)
    added.append(key)

sheet.auto_filter.ref = f"A1:AH{sheet.max_row}"
workbook.save(TARGET)

check = load_workbook(TARGET, read_only=True, data_only=True)
check_sheet = check["成交明细"]
if check_sheet.max_row != 50 + len(added):
    raise RuntimeError(f"保存后行数异常: {check_sheet.max_row}")
print(f"target={TARGET}")
print(f"added={len(added)} rows={check_sheet.max_row} columns={check_sheet.max_column}")
for row in check_sheet.iter_rows(min_row=check_sheet.max_row - len(added) + 1, values_only=True):
    print(row[:14])
