"""按下单时 PM bid 重算低级别赛事的反向购买利润。"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


LOW_LEVEL_LEAGUES = {
    "Adana",
    "Bari",
    "Columbus",
    "Curitiba",
    "Jingshan",
    "Mouilleron-Le-Captif",
    "Plovdiv 4",
    "Porto 2",
    "San Diego 2",
    "St. Tropez",
    "Tolentino",
}


def opposite(side: str) -> str:
    return "no" if side == "yes" else "yes"


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = text.rsplit(":", 1)[-1]
    text = re.sub(r"\bluis guto miguel\b", "luis miguel", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def quote_for_target(order: dict, target: str) -> tuple[float, str]:
    """直接 bid 优先；目标 bid 不存在时才使用 1 - 对向 ask。"""
    bids: dict[str, tuple[float, str]] = {}
    asks: dict[str, tuple[float, str]] = {}
    if order["original_bid"] is not None:
        bids[order["original_side"]] = (float(order["original_bid"]), "original_bid")
    if order["actual_bid"] is not None:
        bids[order["actual_side"]] = (float(order["actual_bid"]), "actual_bid")
    if order["original_ask"] is not None:
        asks[order["original_side"]] = (float(order["original_ask"]), "original_ask")
    if order["actual_ask"] is not None:
        asks[order["actual_side"]] = (float(order["actual_ask"]), "actual_ask")

    direct = bids.get(target)
    if direct is not None:
        return direct
    other = asks.get(opposite(target))
    if other is None:
        raise ValueError(f"{order['id']} 缺少 {target} bid，且缺少对向 ask")
    return 1.0 - other[0], f"1-{other[1]}"


def load_orders(workbook: Path) -> list[dict]:
    sheet = load_workbook(workbook, read_only=True, data_only=True)["nohup成交_统一规则"]
    orders = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[4] != "已成交" or row[16] is None or row[18] is None:
            continue
        orders.append(
            {
                "time": row[0],
                "id": str(row[2]),
                "match": str(row[3]),
                "original_side": str(row[5]).lower(),
                "original_bid": row[10],
                "original_ask": row[11],
                "actual_side": str(row[16]).lower(),
                "actual_bid": row[30],
                "actual_ask": row[31],
                "qty": float(row[18]),
                "result": row[23],
            },
        )
    return orders


def settled_low_level_orders(workbook: Path, snapshot: Path) -> list[tuple[dict, str, str]]:
    positions = json.loads(snapshot.read_text())
    by_match: dict[str, dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    for position in positions:
        date_match = re.search(r"(20\d\d-\d\d-\d\d)$", str(position.get("slug", "")))
        event_date = date_match.group(1) if date_match else ""
        by_match[normalize(position.get("title", ""))][event_date].append(position)

    result = []
    for order in load_orders(workbook):
        candidates = by_match.get(normalize(order["match"]), {})
        if not candidates:
            continue
        order_date = (
            order["time"].date()
            if isinstance(order["time"], datetime)
            else datetime.fromisoformat(str(order["time"])).date()
        )
        _, rows = min(
            candidates.items(),
            key=lambda item: abs((datetime.fromisoformat(item[0]).date() - order_date).days)
            if item[0]
            else 99999,
        )
        league = str(rows[0].get("title", "")).split(":", 1)[0]
        if league not in LOW_LEVEL_LEAGUES:
            continue
        wanted_index = 0 if order["actual_side"] == "yes" else 1
        actual_rows = [row for row in rows if int(row.get("outcomeIndex", -1)) == wanted_index]
        settled_price = float(actual_rows[0].get("curPrice", -1)) if actual_rows else -1.0
        if settled_price in {0.0, 1.0}:
            winner = order["actual_side"] if settled_price == 1.0 else opposite(order["actual_side"])
        elif order["result"] in {"赢", "输"}:
            winner = order["actual_side"] if order["result"] == "赢" else opposite(order["actual_side"])
        else:
            continue
        result.append((order, winner, league))
    return result


def calculate(rows: list[tuple[dict, str, str]], mode: str) -> list[dict]:
    details = []
    for order, winner, league in rows:
        base = order["original_side"] if mode == "original_reverse" else order["actual_side"]
        target = opposite(base)
        price, source = quote_for_target(order, target)
        won = target == winner
        pnl = order["qty"] * ((1.0 if won else 0.0) - price)
        details.append(
            {
                "id": order["id"],
                "match": order["match"],
                "league": league,
                "target": target,
                "price": price,
                "price_source": source,
                "qty": order["qty"],
                "won": won,
                "pnl": pnl,
            },
        )
    return details


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("snapshot", type=Path)
    parser.add_argument("--details", action="store_true")
    args = parser.parse_args()

    rows = settled_low_level_orders(args.workbook, args.snapshot)
    for mode in ("original_reverse", "actual_reverse"):
        details = calculate(rows, mode)
        direct = sum(not item["price_source"].startswith("1-") for item in details)
        print(
            f"{mode}: count={len(details)} wins={sum(item['won'] for item in details)} "
            f"losses={sum(not item['won'] for item in details)} direct_bid={direct} "
            f"fallback={len(details) - direct} pnl={sum(item['pnl'] for item in details):.4f}",
        )
        if args.details:
            for item in details:
                print(json.dumps({"mode": mode, **item}, ensure_ascii=False))


if __name__ == "__main__":
    main()
