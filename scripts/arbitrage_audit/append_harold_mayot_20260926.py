"""把 Harold Mayot 54c 订单诊断行追加到既有比分报表。"""

from __future__ import annotations

import argparse
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment


ORDER_ID = "ARB-426443da"
SHEET_NAME = "明细"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook = load_workbook(args.workbook)
    sheet = workbook[SHEET_NAME]
    existing = {sheet.cell(row, 2).value for row in range(2, sheet.max_row + 1)}

    if ORDER_ID not in existing:
        style_row = sheet.max_row
        sheet.append([
            datetime(2026, 9, 26, 21, 55, 10, 132000),
            ORDER_ID,
            "Harold Mayot vs Murphy Cassone",
            "已成交",
            "5-7, 6-2, 0-1",
            "5-7, 6-2, 0-1",
            "S3",
            58.650395,
            0.64,
            0.59,
            0.60,
            0.59 / 0.64,
            "venue_replace convert反买（不计汇总）",
            "yes",
            0.54,
            "未出",
            5.05,
            2.727,
            None,
            None,
            None,
        ])
        target_row = sheet.max_row
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(style_row, column)
            target = sheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = source.number_format
            target.alignment = copy(source.alignment)
        sheet.cell(target_row, 13).comment = Comment(
            "不是比分左右交叉：比分门控先保留第三盘领先的 Murphy Cassone（PM NO），"
            "随后 venue_replace convert=true 在比分门控之后把 PM NO 反转为 Harold Mayot YES。",
            "Codex",
        )

    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(args.output)
    print(f"{args.output} rows={sheet.max_row - 1}")


if __name__ == "__main__":
    main()
