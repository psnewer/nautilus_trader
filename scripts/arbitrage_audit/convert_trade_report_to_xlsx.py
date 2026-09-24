"""将成交审计的管道文本报告转换为带样式的 Excel 工作簿。"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path("/Users/miller/Desktop/成交订单比分_start_price_venue_stale_window_20260921-22.txt")
TARGET = Path("/Users/miller/Desktop/成交订单比分_start_price_venue_stale_window_20260921-22.xlsx")

PRICE_VENUES = ("PM", "OE", "SE")
PRICE_ROLES = ("YES", "NO")


def number(value: str):
    return None if value == "-" else float(value)


def price_vector(value: str) -> list[float | None]:
    return [number(item) for item in value.split("/")]


lines = SOURCE.read_text(encoding="utf-8").splitlines()
rows = []
for line in lines[3:]:
    fields = line.split("|")
    if len(fields) != 13:
        raise ValueError(f"Unexpected field count {len(fields)}: {line}")
    (
        time_text,
        sport,
        player_1,
        player_2,
        role,
        fill_text,
        score_text,
        start_text,
        native_venue,
        trigger_venue,
        before_at_text,
        delay_text,
        after_text,
    ) = fields

    fill_price, effective_price = (float(item) for item in fill_text.split("/"))
    score_match = re.fullmatch(r"(.*)\(([^()]*)\)", score_text)
    score = score_match.group(1) if score_match else score_text
    period = score_match.group(2) if score_match else None
    start_prices, start_source = start_text.rsplit(";", 1)
    start_yes, start_no = (float(item) for item in start_prices.split("/"))
    before_text, at_text = before_at_text.split("=>", 1)

    if delay_text == "-":
        delay_seconds = None
        later_changed_venue = None
    else:
        seconds_text, later_changed_venue = delay_text.split("/", 1)
        delay_seconds = float(seconds_text.removesuffix("s"))

    rows.append([
        datetime.strptime(f"2026-{time_text}", "%Y-%m-%d %H:%M:%S"),
        sport,
        player_1,
        player_2,
        role,
        fill_price,
        effective_price,
        score,
        period,
        start_yes,
        start_no,
        start_source,
        native_venue,
        trigger_venue,
        *price_vector(before_text),
        *price_vector(at_text),
        delay_seconds,
        later_changed_venue,
        *price_vector(after_text),
    ])

headers = [
    "成交时间(北京时间)",
    "赛事",
    "选手1",
    "选手2",
    "成交方向",
    "成交价",
    "含佣价格",
    "成交比分",
    "盘/局",
    "start_yes",
    "start_no",
    "start来源",
    "原生腿venue",
    "触发venue",
]
for stage in ("变化前", "触发时"):
    headers.extend(f"{stage}_{venue}_{role}" for venue in PRICE_VENUES for role in PRICE_ROLES)
headers.extend(["另一venue变化间隔(秒)", "后来变化venue"])
headers.extend(f"变化后_{venue}_{role}" for venue in PRICE_VENUES for role in PRICE_ROLES)

wb = Workbook()
ws = wb.active
ws.title = "成交明细"
ws.append(headers)
for row in rows:
    ws.append(row)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

source_column = headers.index("start来源") + 1
fills = {
    "default": PatternFill("solid", fgColor="F4CCCC"),
    "possible_captured": PatternFill("solid", fgColor="FFF2CC"),
    "captured": PatternFill("solid", fgColor="D9EAD3"),
}
for row_index in range(2, ws.max_row + 1):
    source_cell = ws.cell(row_index, source_column)
    source_cell.fill = fills.get(source_cell.value, PatternFill())
    ws.cell(row_index, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    for column in range(6, ws.max_column + 1):
        if isinstance(ws.cell(row_index, column).value, float):
            ws.cell(row_index, column).number_format = "0.000000"

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 32
for index, header in enumerate(headers, 1):
    width = 14
    if header in {"选手1", "选手2", "成交比分"}:
        width = 23
    elif header == "成交时间(北京时间)":
        width = 21
    ws.column_dimensions[get_column_letter(index)].width = width

notes = wb.create_sheet("说明")
notes.append(["项目", "说明"])
notes.append(["captured", "日志可确认 PRE 期间 PM 顶价发生变化，start_price 为日志重建值"])
notes.append(["possible_captured", "存在合格 PRE 快照，但旧日志未记录 delta 来源 venue"])
notes.append(["default", "旧版本未采集到 start_price；0.6/0.6 是当时的默认占位值，不是实际价格"])
notes.append(["价格列", "PM=POLYMARKET，OE=ORBITEXCH，SE=SHARPEXCH；YES/NO 为对应 outcome ask"])
for cell in notes[1]:
    cell.fill = header_fill
    cell.font = header_font
notes.column_dimensions["A"].width = 24
notes.column_dimensions["B"].width = 90
notes.freeze_panes = "A2"

wb.save(TARGET)
print(f"saved={TARGET} rows={len(rows)} columns={len(headers)}")
