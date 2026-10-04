"""把单份 nohup.out 的全部下单按 venue_replace 优先级写入 Excel。"""

from __future__ import annotations

import argparse
import ast
import json
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ANSI = re.compile(r"\x1b\[[0-9;]*m")
ISO_TS = re.compile(r"(\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d\.\d+Z)")
LOCAL_TS = re.compile(r"^(\d{4}-\d\d-\d\d \d\d:\d\d:\d\d,\d{3})")
BEIJING = timezone(timedelta(hours=8))
SHEET_NAME = "nohup成交_统一规则"
HEADER_ALIASES = {
    "策略动作": "pre口径策略动作",
    "规则买入方向": "pre口径买入方向",
    "规则买入价(PM ask)": "历史规则买入价(PM ask)",
    "实际成交方向": "历史实际成交方向",
    "实际成交价": "历史实际成交价",
    "成交数量": "历史成交数量",
    "最终输赢": "历史实际方向最终输赢",
    "OrderInitialized最终方向": "历史OrderInitialized最终方向",
    "最终方向start_price": "历史最终方向start_price",
    "最终方向PM_bid(下单时)": "历史最终方向PM_bid(下单时)",
    "最终方向PM_ask(下单时)": "历史最终方向PM_ask(下单时)",
}
# PMS 比分左右顺序与本地 pair 角色不一致的已核实场次，以实际结算方向为准。
WINNER_OVERRIDES = {
    "Tennis|Julia Grabher|Oksana Selekhmeteva": "no",
}
SCORE_TAIL_OVERRIDES = {
    "Tennis|Luis Miguel|Facundo Mena": "Luis Guto Miguel|Facundo Mena",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("audit_json", type=Path)
    parser.add_argument("nohup", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--append",
        action="store_true",
        help="保留目标 sheet 的历史订单；同一 client_order_id 以本次审计结果更新",
    )
    parser.add_argument(
        "--venue-replace-convert",
        action="store_true",
        help="源日志对应的 venue_replace 配置启用了 convert=true",
    )
    return parser.parse_args()


def parse_iso_ts(line: str) -> datetime | None:
    match = ISO_TS.search(line)
    if match:
        return datetime.fromisoformat(match.group(1).replace("Z", "+00:00"))
    match = LOCAL_TS.search(line)
    if not match:
        return None
    return datetime.strptime(match.group(1), "%Y-%m-%d %H:%M:%S,%f").replace(
        tzinfo=BEIJING,
    ).astimezone(timezone.utc)


def parse_beijing(value: str | None) -> datetime | None:
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d %H:%M:%S.%f").replace(tzinfo=BEIJING)


def prior(items: list[dict], timestamp: datetime) -> dict | None:
    found = None
    for item in items:
        if item["ts"] <= timestamp:
            found = item
        else:
            break
    return found


def read_log(
    path: Path,
) -> tuple[dict[str, list[dict]], dict[str, list[dict]], dict[str, list[dict]]]:
    snapshots: dict[str, list[dict]] = defaultdict(list)
    scores: dict[str, list[dict]] = defaultdict(list)
    tier_events: dict[str, list[dict]] = defaultdict(list)
    with path.open(errors="replace") as stream:
        for raw in stream:
            line = ANSI.sub("", raw.rstrip("\n"))
            timestamp = parse_iso_ts(line)
            if timestamp is None:
                continue
            if "Strategy evaluate scheduled: pair_id=" in line and "order_books=" in line:
                match = re.search(r"pair_id=(.+?), sport=.*?order_books=(\[.*\])$", line)
                if not match:
                    continue
                try:
                    books = ast.literal_eval(match.group(2))
                except (SyntaxError, ValueError):
                    continue
                snapshots[match.group(1)].append({"ts": timestamp, "books": books})
            elif "Sports score updated: game=" in line:
                match = re.search(
                    r"game=\d+ (.+?) vs (.+?) score=(.*?) period=([^ ]+) elapsed=.*? status=([^ ]+)$",
                    line,
                )
                if match:
                    scores[f"{match.group(1)}|{match.group(2)}"].append({
                        "ts": timestamp,
                        "score": match.group(3).split("->")[-1],
                        "period": match.group(4),
                        "status": match.group(5),
                    })
            elif (
                "VenueReplace: pair=" in line
                and " tier_convert=" in line
                and " not hit outcome=" not in line
            ):
                match = re.search(
                    r"VenueReplace: pair=(.+?) tier_convert=(pre|post) "
                    r"(?:hit|eligible) oe_competition=(.+)$",
                    line,
                )
                if match:
                    tier_events[match.group(1)].append({
                        "ts": timestamp,
                        "mode": match.group(2),
                        "kind": "eligible",
                    })
            elif "VenueReplace: pair=" in line and " post ignored outcome=" in line:
                match = re.search(
                    r"VenueReplace: pair=(.+?) leg=.*? post ignored outcome=(yes|no) "
                    r"start_price=([^ ]+) threshold=([^ ]+)$",
                    line,
                )
                if match:
                    tier_events[match.group(1)].append({
                        "ts": timestamp,
                        "mode": "post",
                        "kind": "ignored",
                        "outcome": match.group(2),
                        "start_price": match.group(3),
                        "threshold": match.group(4),
                    })
            elif "VenueReplace: pair=" in line and " not hit outcome=" in line:
                match = re.search(
                    r"VenueReplace: pair=(.+?) leg=.*? tier_convert=(pre|post) "
                    r"not hit outcome=(yes|no) start_price=([^ ]+) threshold=([^ ]+)$",
                    line,
                )
                if match:
                    tier_events[match.group(1)].append({
                        "ts": timestamp,
                        "mode": match.group(2),
                        "kind": "ignored",
                        "outcome": match.group(3),
                        "start_price": match.group(4),
                        "threshold": match.group(5),
                    })
    return snapshots, scores, tier_events


def recent_tier_state(
    events: list[dict],
    order_time: datetime,
    source_role: str,
    *,
    max_age_seconds: float = 1.0,
) -> tuple[str | None, dict | None]:
    recent = [
        (index, event)
        for index, event in enumerate(events)
        if event["ts"] <= order_time
        and (order_time - event["ts"]).total_seconds() <= max_age_seconds
    ]
    eligible = [item for item in recent if item[1]["kind"] == "eligible"]
    if not eligible:
        return None, None
    eligible_index, eligible_event = eligible[-1]
    ignored = next(
        (
            event
            for index, event in recent
            if index > eligible_index
            and event["kind"] == "ignored"
            and event["outcome"] == source_role
        ),
        None,
    )
    return (None, ignored) if ignored else (str(eligible_event["mode"]), None)


def action_description(
    *,
    source_role: str,
    actual_role: str,
    native_venue: str,
    tier_mode: str | None,
    tier_ignored: dict | None,
    convert_enabled: bool,
    fallback_action: str,
) -> str:
    """按真实 action 顺序描述反转，避免双反转因首尾同向而被漏记。"""
    if tier_ignored is not None:
        return (
            f"{fallback_action}（tier_convert {tier_ignored['mode']}豁免："
            f"start_price={tier_ignored['start_price']} < {tier_ignored['threshold']}）"
        )
    if tier_mode == "pre":
        return "tier_convert pre反买"
    if tier_mode != "post":
        return fallback_action

    before_post = "no" if actual_role == "yes" else "yes"
    prior_flipped = before_post != source_role
    if prior_flipped:
        prior_action = (
            "venue_replace convert反买"
            if convert_enabled and native_venue == "POLYMARKET"
            else "attitude/deviate_convert反买"
        )
        suffix = "（最终原方向）" if actual_role == source_role else ""
        return f"{prior_action} + tier_convert post再反转{suffix}"
    return "tier_convert post反买"


def quote_map(books: list[dict]) -> dict[str, dict]:
    return {
        str(item["outcome"]).lower(): {
            "bid": item.get("best_bid"),
            "ask": item.get("best_ask"),
        }
        for item in books
        if str(item.get("venue", "")).upper() == "POLYMARKET"
    }


def tennis_winner(score: str | None) -> tuple[str | None, int, int]:
    """从逗号分隔的盘分推断胜方；当前未结束盘不计入盘胜负。"""
    left_wins = 0
    right_wins = 0
    for raw_set in str(score or "").split(","):
        match = re.match(r"\s*(\d+)\s*-\s*(\d+)", raw_set)
        if not match:
            continue
        left, right = map(int, match.groups())
        high, low = max(left, right), min(left, right)
        if not (high >= 6 and (high - low >= 2 or high == 7)):
            continue
        left_wins += left > right
        right_wins += left < right
    winner = "yes" if left_wins > right_wins else "no" if right_wins > left_wins else None
    return winner, left_wins, right_wins


def main() -> None:
    args = parse_args()
    audit = json.loads(args.audit_json.read_text(encoding="utf-8"))
    source_rows = audit["rows"]
    snapshots, scores, tier_events = read_log(args.nohup)

    headers = [
        "下单时间(北京时间)", "首次成交时间(北京时间)", "client_order_id", "比赛", "订单状态",
        "规则输入方向(venue_replace前)", "下单时盘/局", "start_price",
        "start_yes+start_no", "commission区间合格",
        "current_PM_bid(规则判断价)", "current_PM_ask(参考)", "bid/start", "pre口径策略动作",
        "pre口径买入方向", "历史规则买入价(PM ask)", "历史实际成交方向", "历史实际成交价", "历史成交数量",
        "下单时比分", "首次成交时比分", "日志末比分", "日志末状态", "历史实际方向最终输赢",
        "本金", "毛利润", "手续费", "净利润",
        "历史OrderInitialized最终方向", "历史最终方向start_price", "历史最终方向PM_bid(下单时)",
        "历史最终方向PM_ask(下单时)", "规则方向数据源", "原始下单数量",
        "post口径策略动作", "post口径买入方向", "官方最终胜方", "历史实际方向官方输赢",
    ]
    rows = []
    for source in source_rows:
        pair = f"{source['赛事']}|{source['选手1']}|{source['选手2']}"
        tail = SCORE_TAIL_OVERRIDES.get(
            pair,
            f"{source['选手1']}|{source['选手2']}",
        )
        order_time = parse_beijing(source["下单时间(北京时间)"])
        if order_time is None:
            continue
        order_time_utc = order_time.astimezone(timezone.utc)
        snapshot = prior(snapshots[pair], order_time_utc)
        quotes = quote_map(snapshot["books"]) if snapshot else {}
        actual_role = str(source["下单方向"]).lower()
        native_venue = str(source.get("原生腿venue") or "").upper()
        role = str(source.get("venue_replace前方向") or "").lower()
        if role not in {"yes", "no"}:
            role = ""
        role_quote = quotes.get(role, {})
        current_bid = role_quote.get("bid")
        current_ask = role_quote.get("ask")
        start_price = source.get(f"start_{role}")
        start_sum = None
        if source.get("start_yes") is not None and source.get("start_no") is not None:
            start_sum = float(source["start_yes"]) + float(source["start_no"])
        commission_ok = start_sum is not None and 0.98 <= start_sum <= 1.02

        explicit_convert = native_venue == "POLYMARKET" and actual_role != role
        if not role:
            action = "无法确认（缺少venue_replace前方向）"
        elif explicit_convert:
            flip = True
            action = "venue_replace convert反买"
        elif start_price is None:
            flip = False
            action = "买原方向（缺少start_price，规则不触发）"
        elif not commission_ok:
            flip = False
            action = "买原方向（commission区间不合格）"
        elif current_bid is None:
            flip = False
            action = "买原方向（缺少PM bid，规则不触发）"
        else:
            start_value = float(start_price)
            bid_value = float(current_bid)
            if bid_value <= start_value:
                action = "attitude反买"
            elif bid_value >= 1.2 * start_value:
                action = "deviate_convert反买"
            else:
                action = "买原方向"

        tier_mode, tier_ignored = recent_tier_state(
            tier_events[pair],
            order_time_utc,
            role,
        )
        action = action_description(
            source_role=role,
            actual_role=actual_role,
            native_venue=native_venue,
            tier_mode=tier_mode,
            tier_ignored=tier_ignored,
            convert_enabled=args.venue_replace_convert,
            fallback_action=action,
        )

        # 实际报表以 OrderInitialized 的最终方向为准；策略推演只用于解释动作。
        buy_role = actual_role
        buy_quote = quotes.get(buy_role, {})
        buy_price = buy_quote.get("ask")
        buy_start_price = source.get(f"start_{buy_role}")
        latest_score = scores[tail][-1] if scores[tail] else None
        winner_role, left_wins, right_wins = tennis_winner(latest_score["score"] if latest_score else None)
        terminal = latest_score is not None and (
            str(latest_score["status"]).lower() in {"finished", "final", "ended"}
            or max(left_wins, right_wins) >= 2
        )
        if not terminal:
            winner_role = None
        winner_role = WINNER_OVERRIDES.get(pair, winner_role)
        won = winner_role == buy_role if winner_role else None
        quantity = float(source["已成交量"])
        has_fill = quantity > 0
        fill_time = parse_beijing(source["首次成交时间(北京时间)"])
        fill_score = (
            prior(scores[tail], fill_time.astimezone(timezone.utc))
            if fill_time else None
        )
        principal = quantity * float(buy_price) if has_fill and buy_price is not None else None
        gross = quantity * ((1.0 if won else 0.0) - float(buy_price)) if has_fill and won is not None and buy_price is not None else None
        commission = quantity * float(buy_price) * (1.0 - float(buy_price)) * 0.05 if gross is not None else None

        rows.append([
            order_time.replace(tzinfo=None),
            fill_time.replace(tzinfo=None) if fill_time else None,
            source["client_order_id"], f"{source['选手1']} vs {source['选手2']}", source["状态"],
            role or None, source["盘/局"], start_price, start_sum,
            "是" if commission_ok else "否（缺值）" if start_sum is None else "否",
            current_bid, current_ask,
            float(current_bid) / float(start_price) if current_bid is not None and start_price not in {None, 0} else None,
            action, buy_role, buy_price,
            actual_role if has_fill else None,
            source["成交均价"] if has_fill else None,
            quantity if has_fill else None,
            source["下单比分"],
            fill_score["score"] if fill_score else None,
            latest_score["score"] if latest_score else None,
            latest_score["status"] if latest_score else None,
            "赢" if won else "输" if won is not None else "未出",
            principal, gross, commission, gross - commission if gross is not None else None,
            buy_role, buy_start_price, buy_quote.get("bid"), buy_price,
            "OrderInitialized + 下单前最近PM帧", source["下单数量"],
            None, None, None, None,
        ])

    workbook = load_workbook(args.workbook)
    previous_count = 0
    if args.append and SHEET_NAME in workbook.sheetnames:
        old_sheet = workbook[SHEET_NAME]
        old_headers = [HEADER_ALIASES.get(cell.value, cell.value) for cell in old_sheet[1]]
        missing_headers = set(headers) - set(old_headers)
        if set(old_headers) - set(headers) or missing_headers - {
            "首次成交时比分", "post口径策略动作", "post口径买入方向",
            "官方最终胜方", "历史实际方向官方输赢",
        }:
            raise ValueError(f"{SHEET_NAME} 表头与当前脚本不一致，拒绝追加")
        old_rows = [
            row
            for row in old_sheet.iter_rows(min_row=2, values_only=True)
            if row[2] and row[3] != "合计"
        ]
        previous_rows = [
            [dict(zip(old_headers, row)).get(header) for header in headers]
            for row in old_rows
        ]
        previous_count = len(previous_rows)
        fresh_by_id = {row[2]: row for row in rows}
        merged_rows = []
        for row in previous_rows:
            merged_rows.append(fresh_by_id.pop(row[2], row))
        merged_rows.extend(fresh_by_id.values())
        rows = merged_rows
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME, 0)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 4, "合计")
    sheet.cell(total_row, 15, f"反买{sum(row[13] in {'买对手盘', 'venue_replace convert反买'} for row in rows)}笔")
    sheet.cell(total_row, 24, f"{sum(row[23] == '赢' for row in rows)}赢/{sum(row[23] == '输' for row in rows)}输/{sum(row[23] == '未出' for row in rows)}未出")
    for column in (25, 26, 27, 28):
        letter = get_column_letter(column)
        sheet.cell(total_row, column, f"=SUM({letter}2:{letter}{total_row - 1})")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[total_row]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)
    for row_number in range(2, total_row):
        sheet.cell(row_number, 1).number_format = "yyyy-mm-dd hh:mm:ss.000"
        sheet.cell(row_number, 2).number_format = "yyyy-mm-dd hh:mm:ss.000"
        for column in list(range(9, 15)) + list(range(17, 21)) + list(range(25, 29)):
            sheet.cell(row_number, column).number_format = "0.0000"
    for column in range(25, 29):
        sheet.cell(total_row, column).number_format = "0.0000"
    widths = [23, 23, 17, 43, 13, 24, 13, 13, 18, 19, 27, 21, 12, 38, 13, 21, 15, 13, 12, 24, 24, 24, 14, 12, 13, 13, 13, 13, 24, 20, 24, 24, 34, 16, 46, 18, 18, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:AB{max(1, total_row - 1)}"
    workbook.save(args.output)
    print(json.dumps({
        "orders": len(rows),
        "previous_orders": previous_count,
        "audit_orders": len(source_rows),
        "sheet": SHEET_NAME,
        "output": str(args.output),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
