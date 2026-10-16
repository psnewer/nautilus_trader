"""PM paper 服务的本地账本和 USER WS 协议测试。"""

from __future__ import annotations

import asyncio
import json
import sqlite3
from pathlib import Path
from types import SimpleNamespace

import httpx
import msgspec
from aiohttp import ClientSession
from aiohttp import web
from openpyxl import load_workbook
from py_clob_client_v2 import ClobClient
from py_clob_client_v2.clob_types import ApiCreds
from py_clob_client_v2.clob_types import OrderPayload
from py_clob_client_v2.http_helpers import helpers as clob_http_helpers
from py_clob_client_v2.order_utils.model.order_data_v2 import SignedOrderV2
from py_clob_client_v2.order_utils.model.side import Side
from py_clob_client_v2.order_utils.model.signature_type_v2 import SignatureTypeV2

from nautilus_trader.adapters.polymarket.schemas.trade import PolymarketTradeReport
from nautilus_trader.adapters.polymarket.schemas.user import PolymarketUserOrder
from nautilus_trader.adapters.polymarket.schemas.user import PolymarketUserTrade
from nautilus_trader.adapters.polymarket.websocket.types import USER_WS_MESSAGE
from scripts.arbitrage_audit.export_paper_execution import export_workbook
from src.arbitrage.paper.polymarket import PaperServerConfig
from src.arbitrage.paper.polymarket import PaperStore
from src.arbitrage.paper.polymarket import PolymarketPaperExchange
from src.arbitrage.paper.polymarket import create_app


class _Exchange(PolymarketPaperExchange):
    def __init__(self, config, store, *, bid=0.39, ask=0.40):
        super().__init__(config, store)
        self.bid = bid
        self.ask = ask
        self.events = []

    async def resolve_market(self, token_id: str) -> str:
        return "0xcondition"

    async def ensure_market_subscription(self, token_id: str) -> None:
        return None

    async def best_prices(self, token_id: str):
        return self.bid, self.ask

    async def broadcast(self, payload):
        self.events.append(payload)


def _payload(*, side="BUY", price=0.5, size=5.0, order_type="GTC"):
    if side == "BUY":
        maker_amount = round(price * size * 1_000_000)
        taker_amount = round(size * 1_000_000)
    else:
        maker_amount = round(size * 1_000_000)
        taker_amount = round(price * size * 1_000_000)
    return {
        "order": {
            "salt": 1,
            "maker": "0x0000000000000000000000000000000000000001",
            "signer": "0x0000000000000000000000000000000000000001",
            "taker": "0x0000000000000000000000000000000000000000",
            "tokenId": "123",
            "makerAmount": str(maker_amount),
            "takerAmount": str(taker_amount),
            "expiration": "0",
            "nonce": "0",
            "feeRateBps": "0",
            "side": side,
            "signatureType": 0,
            "signature": "0x00",
        },
        "owner": "paper-api-key",
        "orderType": order_type,
        "postOnly": False,
    }


def test_marketable_order_persists_fill_position_balance_and_decodable_ws(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
        )
        response = await exchange.place(_payload(price=0.5, size=5.0))
        await asyncio.sleep(0.03)

        order = store.order(response["orderID"])
        assert order["status"] == "MATCHED"
        assert store.positions() == [{
            "conditionId": "0xcondition",
            "asset": "123",
            "size": 5.0,
            "avgPrice": 0.4,
            "redeemable": False,
            "negRisk": False,
        }]
        assert store.balance_units() == 98_000_000
        first_event = msgspec.json.decode(
            json.dumps(exchange.events[0]).encode(),
            type=USER_WS_MESSAGE,
        )
        second_event = msgspec.json.decode(
            json.dumps(exchange.events[1]).encode(),
            type=USER_WS_MESSAGE,
        )
        assert isinstance(first_event, PolymarketUserOrder)
        assert isinstance(second_event, PolymarketUserTrade)
        assert isinstance(
            msgspec.json.decode(json.dumps(store.trades()[0]).encode(), type=PolymarketTradeReport),
            PolymarketTradeReport,
        )
        store.close()

    asyncio.run(scenario())


def test_resting_gtc_fills_only_after_later_price_cross(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
            ask=0.60,
        )
        response = await exchange.place(_payload(price=0.5))
        await asyncio.sleep(0.03)
        assert store.order(response["orderID"])["status"] == "LIVE"
        await exchange._handle_market_ws_payload(json.dumps({
            "event_type": "price_change",
            "market": "0xcondition",
            "price_changes": [{
                "asset_id": "123",
                "price": "0.49",
                "side": "SELL",
                "size": "1",
                "best_bid": "0.39",
                "best_ask": "0.49",
            }],
            "timestamp": "1",
        }))
        assert store.order(response["orderID"])["status"] == "MATCHED"
        assert exchange.events[-1]["trader_side"] == "MAKER"
        assert exchange.events[-1]["price"] == "0.5"
        # 下一帧立即回撤也不能抹掉上一帧已经发生的成交。
        await exchange._on_market_prices("123", 0.39, 0.60)
        assert len(store.trades()) == 1
        store.close()

    asyncio.run(scenario())


def test_crossed_market_frame_is_not_published_or_matched(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
            ask=0.60,
        )
        response = await exchange.place(_payload(price=0.5))
        await asyncio.sleep(0.03)

        await exchange._on_market_prices("123", 0.51, 0.49)

        assert store.order(response["orderID"])["status"] == "LIVE"
        assert "123" not in exchange._market_books
        store.close()

    asyncio.run(scenario())


def test_unmatched_ioc_is_canceled(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
            ask=0.60,
        )
        response = await exchange.place(_payload(price=0.5, order_type="FAK"))
        await asyncio.sleep(0.03)
        assert store.order(response["orderID"])["status"] == "CANCELED"
        assert exchange.events[-1]["type"] == "CANCELLATION"
        store.close()

    asyncio.run(scenario())


def test_open_orders_reserve_balance_and_position(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 2.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
            ask=0.90,
        )
        first = await exchange.place(_payload(price=0.5, size=4.0))
        assert first["success"]
        second = await exchange.place(_payload(price=0.4, size=1.0))
        assert second == {"success": False, "errorMsg": "insufficient paper balance"}

        exchange.ask = 0.4
        await exchange._on_market_prices("123", exchange.bid, exchange.ask)
        assert store.order(first["orderID"])["status"] == "MATCHED"
        sell = await exchange.place(_payload(side="SELL", price=0.8, size=4.0))
        assert sell["success"]
        oversell = await exchange.place(_payload(side="SELL", price=0.9, size=1.0))
        assert oversell == {"success": False, "errorMsg": "insufficient paper position"}
        store.close()

    asyncio.run(scenario())


def test_duplicate_submission_is_idempotent(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
            ask=0.90,
        )
        payload = _payload(price=0.5)
        first = await exchange.place(payload)
        second = await exchange.place(payload)
        assert second == first
        assert len(store.orders()) == 1
        store.close()

    asyncio.run(scenario())


def test_existing_database_is_migrated_and_order_metadata_is_persisted(tmp_path):
    database_path = tmp_path / "paper.sqlite3"
    db = sqlite3.connect(database_path)
    db.execute("CREATE TABLE orders (id TEXT PRIMARY KEY)")
    db.commit()
    db.close()

    store = PaperStore(str(database_path), 100.0)
    columns = {
        row["name"]
        for row in store._db.execute("PRAGMA table_info(orders)")
    }
    assert {"client_order_id", "pair_id", "origin", "intent", "order_score"} <= columns
    store.close()


def test_reset_recreates_empty_ledger_with_starting_balance(tmp_path):
    database_path = tmp_path / "paper.sqlite3"
    store = PaperStore(str(database_path), 100.0)
    store.insert_order(
        {
            "id": "old-order",
            "market": "0xcondition",
            "asset_id": "123",
            "side": "BUY",
            "price": 0.4,
            "original_size": 5.0,
            "order_type": "GTC",
            "post_only": False,
            "maker_address": "paper",
            "owner": "paper",
            "expiration": "0",
            "created_at": 1_700_000_000,
        },
    )
    store.close()
    Path(f"{database_path}-wal").write_bytes(b"stale wal")
    Path(f"{database_path}-shm").write_bytes(b"stale shm")

    reset_store = PaperStore(str(database_path), 250.0, reset=True)

    assert reset_store.orders() == []
    assert reset_store.available_balance() == 250.0
    assert Path(f"{database_path}-wal").read_bytes() != b"stale wal"
    assert Path(f"{database_path}-shm").read_bytes() != b"stale shm"
    reset_store.close()


def test_export_workbook_writes_detail_pair_and_type_sheets(tmp_path):
    database_path = tmp_path / "paper.sqlite3"
    output_path = tmp_path / "paper.xlsx"
    store = PaperStore(str(database_path), 100.0)
    store.insert_order({
        "id": "paper-order-1",
        "market": "0xcondition",
        "asset_id": "123",
        "side": "BUY",
        "price": 0.4,
        "original_size": 5.0,
        "order_type": "GTC",
        "post_only": False,
        "maker_address": "paper",
        "owner": "paper",
        "expiration": "0",
        "created_at": 1_700_000_000,
    })
    store.update_order_metadata("paper-order-1", {
        "client_order_id": "ARB-1",
        "pair_id": "Tennis|Player A|Player B",
        "origin": "in_game",
        "intent": "arbitrage",
        "instrument_id": "0xcondition-123.POLYMARKET",
        "outcome": "Player A",
        "selection_role": "yes",
        "order_score": "1-0, 2-1",
    })
    trade = {
        "id": "trade-1",
        "taker_order_id": "paper-order-1",
        "trader_side": "TAKER",
        "match_time": "1700000001",
        "price": "0.39",
        "size": "5",
        "maker_orders": [],
    }
    store.fill("paper-order-1", 0.39, trade)
    store.close()

    assert export_workbook(database_path, output_path) == (1, 1)
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["订单明细", "Pair汇总", "类型汇总", "说明"]
    detail = list(workbook["订单明细"].values)
    assert detail[1][4] == "Tennis|Player A|Player B"
    assert detail[1][8] == "in_game/arbitrage"
    pair_summary = list(workbook["Pair汇总"].values)
    assert pair_summary[1][4:6] == (1, 1)


def test_export_workbook_merges_cancel_score_and_quotes_from_log_audit(tmp_path):
    database_path = tmp_path / "paper.sqlite3"
    audit_path = tmp_path / "audit.json"
    output_path = tmp_path / "paper.xlsx"
    store = PaperStore(str(database_path), 100.0)
    store.insert_order({
        "id": "paper-order-1",
        "market": "0xcondition",
        "asset_id": "123",
        "side": "BUY",
        "price": 0.4,
        "original_size": 5.0,
        "order_type": "GTC",
        "post_only": False,
        "maker_address": "paper",
        "owner": "paper",
        "expiration": "0",
        "created_at": 1_700_000_000,
    })
    store.update_order_metadata("paper-order-1", {
        "client_order_id": "ARB-1",
        "pair_id": "Tennis|Player A|Player B",
        "origin": "pre_game",
        "intent": "arbitrage",
    })
    store.cancel("paper-order-1")
    store.close()
    audit_path.write_text(json.dumps({"rows": [{
        "client_order_id": "ARB-1",
        "撤单时间(北京时间)": "2026-10-15 12:00:00.123",
        "撤单比分": "3-2",
        "撤单盘/局": "S1",
        "撤单方向PM_bid": 0.38,
        "撤单方向PM_ask": 0.40,
        "撤单时_PM_YES_bid": 0.38,
        "撤单时_PM_YES_ask": 0.40,
        "首次成交比分": "3-2",
        "首次成交盘/局": "S1",
        "下单方向": "no",
        "venue_replace前方向": "yes",
    }]}), encoding="utf-8")

    export_workbook(database_path, output_path, audit_path)
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    values = list(workbook["订单明细"].values)
    headers = {value: index for index, value in enumerate(values[0])}
    row = values[1]
    assert row[headers["首次成交比分"]] == "3-2"
    assert row[headers["首次成交盘/局"]] == "S1"
    assert row[headers["撤单时间(北京时间)"]] == "2026-10-15 12:00:00.123"
    assert row[headers["撤单比分"]] == "3-2"
    assert row[headers["撤单方向PM_bid"]] == 0.38
    assert row[headers["撤单时_PM_YES_ask"]] == 0.40
    assert row[headers["venue_replace前方向"]] == "yes"
    assert row[headers["原单/反买单"]] == "反买单"


def test_concurrent_orders_cannot_overreserve_balance(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 2.0)
        exchange = _Exchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
            ask=0.90,
        )
        first, second = await asyncio.gather(
            exchange.place(_payload(price=0.5, size=4.0)),
            exchange.place({**_payload(price=0.5, size=4.0), "owner": "second"}),
        )

        assert sorted([first["success"], second["success"]]) == [False, True]
        assert len(store.orders(open_only=True)) == 1
        assert store.available_balance() == 0.0
        store.close()

    asyncio.run(scenario())


def test_http_and_user_ws_protocol_round_trip(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = PolymarketPaperExchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
        )

        async def resolve_market(_token_id):
            return "0xcondition"

        async def best_prices(_token_id):
            return 0.39, 0.40

        async def ensure_market_subscription(_token_id):
            return None

        exchange.resolve_market = resolve_market
        exchange.best_prices = best_prices
        exchange.ensure_market_subscription = ensure_market_subscription
        runner = web.AppRunner(create_app(exchange))
        await runner.setup()
        site = web.TCPSite(runner, "127.0.0.1", 0)
        await site.start()
        port = site._server.sockets[0].getsockname()[1]
        try:
            async with ClientSession() as session:
                async with session.ws_connect(f"http://127.0.0.1:{port}/ws/user") as ws:
                    async with session.post(
                        f"http://127.0.0.1:{port}/order",
                        json=_payload(price=0.5),
                    ) as response:
                        posted = await response.json()
                    placement = await ws.receive_json(timeout=1)
                    trade = await ws.receive_json(timeout=1)
                    assert placement["id"] == posted["orderID"]
                    assert trade["taker_order_id"] == posted["orderID"]

                async with session.get(
                    f"http://127.0.0.1:{port}/data/order/{posted['orderID']}",
                ) as response:
                    report = await response.json()
                assert report["status"] == "MATCHED"

                async with session.post(
                    f"http://127.0.0.1:{port}/paper/order-metadata",
                    json={
                        "order_id": posted["orderID"],
                        "client_order_id": "ARB-1",
                        "pair_id": "Tennis|Player A|Player B",
                        "origin": "in_game",
                        "intent": "arbitrage",
                        "order_score": "1-0, 2-1",
                    },
                ) as response:
                    assert response.status == 200
                enriched = store.order(posted["orderID"])
                assert enriched["pair_id"] == "Tennis|Player A|Player B"
                assert enriched["origin"] == "in_game"
                assert enriched["intent"] == "arbitrage"

                async with session.get(f"http://127.0.0.1:{port}/positions") as response:
                    positions = await response.json()
                assert positions[0]["asset"] == "123"
        finally:
            await runner.cleanup()

    asyncio.run(scenario())


def test_real_clob_sdk_submit_and_reports_round_trip(tmp_path):
    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = PolymarketPaperExchange(
            PaperServerConfig(database_path=str(tmp_path / "paper.sqlite3")),
            store,
        )

        async def resolve_market(_token_id):
            return "0xcondition"

        async def best_prices(_token_id):
            return 0.39, 0.40

        async def ensure_market_subscription(_token_id):
            return None

        exchange.resolve_market = resolve_market
        exchange.best_prices = best_prices
        exchange.ensure_market_subscription = ensure_market_subscription
        runner = web.AppRunner(create_app(exchange))
        await runner.setup()
        site = web.TCPSite(runner, "127.0.0.1", 0)
        await site.start()
        port = site._server.sockets[0].getsockname()[1]
        original_http_client = clob_http_helpers._http_client
        direct_http_client = httpx.Client(trust_env=False)
        clob_http_helpers._http_client = direct_http_client
        try:
            client = ClobClient(
                f"http://127.0.0.1:{port}",
                chain_id=137,
                key="0x" + "1" * 64,
                creds=ApiCreds(
                    api_key="paper-key",
                    api_secret="AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA=",
                    api_passphrase="paper-passphrase",
                ),
            )
            signed = SignedOrderV2(
                salt="1",
                maker="0x0000000000000000000000000000000000000001",
                signer="0x0000000000000000000000000000000000000001",
                tokenId="123",
                makerAmount="2500000",
                takerAmount="5000000",
                side=Side.BUY,
                signatureType=SignatureTypeV2.EOA,
                timestamp="0",
                metadata="0x",
                builder="0x0000000000000000000000000000000000000000",
                signature="0x00",
            )

            posted = await asyncio.to_thread(client.post_order, signed)
            await asyncio.sleep(0.03)
            report = await asyncio.to_thread(client.get_order, posted["orderID"])
            trades = await asyncio.to_thread(client.get_trades)

            assert report["status"] == "MATCHED"
            assert trades[0]["taker_order_id"] == posted["orderID"]

            resting = SignedOrderV2(
                salt="2",
                maker=signed.maker,
                signer=signed.signer,
                tokenId=signed.tokenId,
                makerAmount="1500000",
                takerAmount="5000000",
                side=Side.BUY,
                signatureType=SignatureTypeV2.EOA,
                timestamp="0",
                metadata="0x",
                builder=signed.builder,
                signature="0x00",
            )
            posted_resting = await asyncio.to_thread(client.post_order, resting)
            await asyncio.sleep(0.03)
            open_orders = await asyncio.to_thread(client.get_open_orders)
            canceled = await asyncio.to_thread(
                client.cancel_order,
                OrderPayload(orderID=posted_resting["orderID"]),
            )
            canceled_report = await asyncio.to_thread(
                client.get_order,
                posted_resting["orderID"],
            )

            assert [item["id"] for item in open_orders] == [posted_resting["orderID"]]
            assert canceled["canceled"] == [posted_resting["orderID"]]
            assert canceled_report["status"] == "CANCELED"
        finally:
            clob_http_helpers._http_client = original_http_client
            direct_http_client.close()
            await runner.cleanup()

    asyncio.run(scenario())


def test_paper_server_launcher_passes_public_pm_proxy(monkeypatch):
    from launchers import polymarket_paper_server

    captured = []
    cfg = SimpleNamespace(
        execution=SimpleNamespace(
            mode="paper",
            paper=SimpleNamespace(
                host="127.0.0.1",
                port=8765,
                database_path="paper.sqlite3",
                starting_balance=100.0,
                reset=True,
            ),
        ),
        venues=SimpleNamespace(
            polymarket=SimpleNamespace(
                clob_url="https://clob.polymarket.com",
                ws_url="wss://ws-subscriptions-clob.polymarket.com/ws/",
                proxy_url="http://pm-proxy:7890",
            ),
        ),
    )
    monkeypatch.setattr(polymarket_paper_server, "load_arb_config", lambda _path: cfg)
    monkeypatch.setattr(polymarket_paper_server, "validate_execution_config", lambda _cfg: None)
    monkeypatch.setattr(polymarket_paper_server, "run_server", captured.append)

    assert polymarket_paper_server.main(["--config", "unused.json"]) == 0
    assert captured[0].market_proxy_url == "http://pm-proxy:7890"
    assert captured[0].market_ws_url == "wss://ws-subscriptions-clob.polymarket.com/ws/market"
    assert captured[0].reset is True


def test_public_market_query_uses_configured_proxy(tmp_path):
    class _Response:
        status = 200
        headers = {"Content-Type": "application/json"}

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return False

        async def read(self):
            return b"{}"

    class _Session:
        def __init__(self):
            self.calls = []

        def get(self, url, **kwargs):
            self.calls.append((url, kwargs))
            return _Response()

    async def scenario():
        store = PaperStore(str(tmp_path / "paper.sqlite3"), 100.0)
        exchange = PolymarketPaperExchange(
            PaperServerConfig(
                database_path=str(tmp_path / "paper.sqlite3"),
                market_proxy_url="http://pm-proxy:7890",
            ),
            store,
        )
        session = _Session()
        exchange._session = session

        assert await exchange.public_get("/book?token_id=123") == (
            200,
            b"{}",
            "application/json",
        )
        assert session.calls == [(
            "https://clob.polymarket.com/book?token_id=123",
            {"proxy": "http://pm-proxy:7890"},
        )]
        store.close()

    asyncio.run(scenario())
